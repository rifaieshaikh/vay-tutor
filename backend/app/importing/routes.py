from __future__ import annotations

import hashlib
import json
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

from bson import ObjectId
from fastapi import APIRouter, File, Request, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import Workbook
from pydantic import BaseModel, ConfigDict, Field
from pymongo.errors import DuplicateKeyError

from app.acl import allows
from app.catalog_pages import calendar_date, place_changes, place_text
from app.errors import AppError
from app.http import current_user, db, write_audit
from app.importing.commit import commit_import, known_sessions
from app.importing.corrections import parse_correction_file
from app.importing.counts import annotate_plan
from app.importing.parse import parse_workbook, workbook_rows
from app.importing.plan import apply_decisions, variant_suggestions
from app.importing.publish import (
    correct_marksheet,
    reject_marksheet,
    require_current_version,
    submit_marksheet,
    withdraw_marksheet,
)
from app.security import name_key
from app.students import (
    card_visible,
    directory_facts,
    enrollment_context,
    enrollment_resource,
    enrollment_visible,
    history_item,
    history_permitted,
    history_visible,
    identity_editable,
)

router = APIRouter(prefix="/api/v1")
_MAX_BYTES = 25 * 1024 * 1024


class DecisionBody(BaseModel):
    sheets: dict = Field(default_factory=dict)
    defaults: dict = Field(default_factory=dict)


class ReasonBody(BaseModel):
    reason: str = ""
    changes: list[dict] = Field(default_factory=list)
    edit_version: int | None = None


class ResultPatch(BaseModel):
    score: float | None = None
    status: str | None = None
    edit_version: int | None = None


class CatalogPatch(BaseModel):
    name: str | None = None
    archived: bool | None = None
    notes: str | None = None
    street: str | None = None
    place: str | None = None
    district: str | None = None
    state: str | None = None
    pin: str | None = None
    phone: str | None = None
    email: str | None = None
    started_on: str | None = None
    ended_on: str | None = None
    timings: str | None = None


class MergeBody(BaseModel):
    source_id: str
    target_id: str
    acknowledge: bool = False


class StudentName(BaseModel):
    model_config = ConfigDict(extra="forbid")
    display_name: str = ""
    edit_version: int | None = None
    student_code: str | None = None
    date_of_birth: str | None = None
    blood_group: str | None = None
    phone: str | None = None
    guardian_name: str | None = None
    guardian_phone: str | None = None
    notes: str | None = None
    street: str | None = None
    place: str | None = None
    district: str | None = None
    state: str | None = None
    pin: str | None = None
    qualification: str | None = None
    institution: str | None = None
    board: str | None = None
    passing_year: str | None = None


_PROFILE_LIMITS = {
    "phone": 30,
    "guardian_name": 120,
    "guardian_phone": 30,
    "notes": 500,
    "street": 200,
    "place": 80,
    "district": 80,
    "state": 80,
    "qualification": 120,
    "institution": 200,
    "board": 160,
}
_BLOOD_GROUPS = {"A+", "A-", "B+", "B-", "O+", "O-", "AB+", "AB-"}
_PROFILE_FIELDS = (
    "date_of_birth", "blood_group", "phone", "guardian_name", "guardian_phone", "notes",
    "street", "place", "district", "state", "pin",
    "qualification", "institution", "board", "passing_year",
)


class StudentAlias(BaseModel):
    model_config = ConfigDict(extra="forbid")
    display_name: str = ""


class EnrollmentDates(BaseModel):
    model_config = ConfigDict(extra="forbid")
    started_on: str = ""
    ended_on: str = ""


class EnrollmentCreate(BaseModel):
    model_config = ConfigDict(extra="forbid")
    branch_id: str
    course_id: str
    batch_id: str
    started_on: str = ""


_MAX_PHOTO = 2 * 1024 * 1024


class BandBody(BaseModel):
    name: str = ""
    through: float | None = None


class MarksheetPatch(BaseModel):
    exam_date: str | None = None
    bands: list[BandBody] | None = None


class MarksheetCreate(BaseModel):
    title: str = Field(min_length=1)
    branch_id: str
    course_id: str
    batch_id: str
    subject_id: str
    paper_id: str
    exam_type: str
    maximum: float = Field(gt=0)
    exam_date: str | None = None
    attempt: str = "original"


class ResultCreate(BaseModel):
    student_id: str
    status: str = "missing"
    score: float | None = None


def _now() -> datetime:
    return datetime.now(timezone.utc)


@router.post("/imports")
async def upload_import(request: Request, file: UploadFile = File(...)) -> dict:
    user = await current_user(request)
    if not any("marksheet.upload" in grant["actions"] for grant in user["grants"]):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")
    filename = file.filename or "marklist.xlsx"
    if not filename.lower().endswith(".xlsx"):
        raise AppError(422, "import.file_type", "Upload an .xlsx workbook.")
    data = await file.read()
    if len(data) > _MAX_BYTES:
        raise AppError(422, "import.too_large", "The workbook is larger than 25 MB.")
    digest = hashlib.sha256(data).hexdigest()
    existing = await db(request).import_jobs.find_one(
        {"institute_id": user["institute_id"], "sha256": digest, "state": {"$nin": ["committed", "cancelled"]}}
    )
    if existing:
        return {"id": existing["_id"], "state": existing["state"], "duplicate": True}
    try:
        parsed = parse_workbook(data, filename)
    except ValueError as exc:
        raise AppError(422, "import.unreadable", str(exc)) from exc
    file_id = str(ObjectId())
    path = request.app.state.settings.data_dir / "files" / f"{file_id}.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    job_id = str(ObjectId())
    await db(request).files.insert_one(
        {
            "_id": file_id,
            "institute_id": user["institute_id"],
            "kind": "source_xlsx",
            "sha256": digest,
            "path": str(path),
            "filename": filename,
            "created_at": _now(),
        }
    )
    state = "needs_resolution" if _needs_resolution(parsed) else "ready"
    await db(request).import_jobs.insert_one(
        {
            "_id": job_id,
            "institute_id": user["institute_id"],
            "actor_id": user["_id"],
            "filename": filename,
            "sha256": digest,
            "file_id": file_id,
            "state": state,
            "parsed": parsed,
            "decisions": {"sheets": {}},
            "committed_sheets": [],
            "outcomes": [],
            "created_at": _now(),
        }
    )
    await db(request).audit_events.insert_one(
        {
            "_id": str(ObjectId()),
            "institute_id": user["institute_id"],
            "action": "import.upload",
            "actor_id": user["_id"],
            "at": _now(),
            "context": {"import_id": job_id, "file": filename},
        }
    )
    return {"id": job_id, "state": state, "duplicate": False, "sheets": len(parsed["sheets"])}


def _needs_resolution(parsed: dict) -> bool:
    for sheet in parsed["sheets"]:
        if sheet["heading_conflict"]:
            return True
        if any(not group["confirmed"] for group in sheet["groups"]):
            return True
        if len(sheet["recommended_batches"]) > 1:
            return True
    return False


@router.get("/imports/{import_id}")
async def get_import(request: Request, import_id: str) -> dict:
    user = await current_user(request)
    job = await _job(request, user, import_id)
    known = await known_sessions(db(request), user["institute_id"])
    plan = apply_decisions(job["parsed"], job.get("decisions") or {}, known)
    plan["totals"].update(await _plan_totals(db(request), user["institute_id"], plan, known))
    await annotate_plan(db(request), user["institute_id"], plan)
    stored_rows = _stored_rows(request, job)
    return {
        "id": job["_id"],
        "state": job["state"],
        "filename": job["filename"],
        "sheets": [_public_sheet(sheet, stored_rows.get(sheet["name"], [])) for sheet in plan["sheets"]],
        "blockers": plan["blockers"],
        "totals": plan["totals"],
        "variants": variant_suggestions(job["parsed"]),
        "outcomes": job.get("outcomes") or [],
        "summary": job.get("summary"),
        "ready": plan["ready"],
    }


@router.patch("/imports/{import_id}")
async def patch_import(request: Request, import_id: str, body: DecisionBody) -> dict:
    user = await current_user(request)
    job = await _job(request, user, import_id)
    if job["state"] == "committed":
        raise AppError(422, "import.committed", "This workbook is already committed.")
    decisions = job.get("decisions") or {"sheets": {}}
    if body.defaults:
        decisions["defaults"] = {**(decisions.get("defaults") or {}), **body.defaults}
    for sheet_id, choice in body.sheets.items():
        current = decisions["sheets"].get(sheet_id, {})
        groups = {**(current.get("groups") or {}), **(choice.get("groups") or {})}
        levels = {**(current.get("levels") or {}), **(choice.get("levels") or {})}
        row_batches = {**(current.get("row_batches") or {}), **(choice.get("row_batches") or {})}
        corrections = {**(current.get("corrections") or {}), **(choice.get("corrections") or {})}
        names = {**(current.get("names") or {}), **(choice.get("names") or {})}
        added = {**(current.get("added") or {}), **(choice.get("added") or {})}
        drop_rows = list(dict.fromkeys([*(current.get("drop_rows") or []), *(choice.get("drop_rows") or [])]))
        current.update(choice)
        if groups:
            current["groups"] = groups
        if levels:
            current["levels"] = levels
        if row_batches:
            current["row_batches"] = row_batches
        if corrections:
            current["corrections"] = corrections
        if names:
            current["names"] = names
        if added:
            current["added"] = added
        if drop_rows:
            current["drop_rows"] = drop_rows
        decisions["sheets"][sheet_id] = current
    known = await known_sessions(db(request), user["institute_id"])
    plan = apply_decisions(job["parsed"], decisions, known)
    state = "ready" if plan["ready"] else "needs_resolution"
    await db(request).import_jobs.update_one({"_id": job["_id"]}, {"$set": {"decisions": decisions, "state": state}})
    return {"id": job["_id"], "state": state, "ready": plan["ready"], "blockers": plan["blockers"]}


@router.post("/imports/{import_id}/corrections")
async def upload_corrections(request: Request, import_id: str, file: UploadFile = File(...)) -> dict:
    user = await current_user(request)
    job = await _job(request, user, import_id)
    if job["state"] == "committed":
        raise AppError(422, "import.committed", "This workbook is already committed.")
    filename = file.filename or "corrections.xlsx"
    if not filename.lower().endswith(".xlsx"):
        raise AppError(422, "import.file_type", "Upload an .xlsx correction file.")
    try:
        rows = parse_correction_file(await file.read())
    except ValueError as exc:
        raise AppError(422, "import.unreadable", str(exc)) from exc
    known = await known_sessions(db(request), user["institute_id"])
    plan = apply_decisions(job["parsed"], job.get("decisions") or {}, known)
    decisions = job.get("decisions") or {"sheets": {}}
    decisions.setdefault("sheets", {})
    applied = []
    unmatched = []
    for row in rows:
        sheet = next((item for item in plan["sheets"] if item["name"].strip().casefold() == row["sheet"].casefold()), None)
        if sheet is None:
            unmatched.append({"sheet": row["sheet"], "student": row["student"], "message": "That sheet is not in this workbook."})
            continue
        wanted_row = row["row"]
        matches = []
        for group in sheet["groups"]:
            for student in group["students"]:
                if wanted_row not in (None, "") and int(wanted_row) != int(student["row"]):
                    continue
                if name_key(row["student"]) == student["name_key"]:
                    matches.append(student)
        if len(matches) != 1:
            unmatched.append({"sheet": row["sheet"], "student": row["student"], "message": "Match one student. Add the Row column when a name appears twice."})
            continue
        score = row["score"]
        status = None
        if isinstance(score, str) and score.strip().upper() in {"A", "AB"}:
            status = "absent"
            score = None
        elif score in (None, ""):
            status = "missing"
            score = None
        choice = decisions["sheets"].setdefault(sheet["id"], {})
        corrections = choice.setdefault("corrections", {})
        corrections[str(matches[0]["row"])] = {"action": "replace", "score": score, "status": status, "reason": row["reason"]}
        applied.append({"sheet": sheet["name"], "student": matches[0]["display_name"], "row": matches[0]["row"]})
    await db(request).import_jobs.update_one({"_id": job["_id"]}, {"$set": {"decisions": decisions}})
    return {"applied": applied, "unmatched": unmatched}


@router.get("/imports/{import_id}/correction-template")
async def correction_template(request: Request, import_id: str):
    payload = await get_import(request, import_id)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Corrections"
    sheet.append(["Sheet", "Row", "Student", "File score", "Saved score", "Score", "Reason"])
    for item in payload["sheets"]:
        for group in item["groups"]:
            for student in group["students"]:
                if student.get("change") not in {"reject", "update"}:
                    continue
                sheet.append([
                    item["name"],
                    student["row"],
                    student["display_name"],
                    student.get("score"),
                    student.get("before_score"),
                    "",
                    "",
                ])
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="corrections.xlsx"'},
    )


@router.get("/imports/{import_id}/preview")
async def preview_import(request: Request, import_id: str) -> dict:
    return await get_import(request, import_id)


@router.get("/imports/{import_id}/validation-report")
async def validation_report(request: Request, import_id: str):
    user = await current_user(request)
    if not any("export.xlsx" in grant["actions"] for grant in user["grants"]):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")
    payload = await get_import(request, import_id)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Validation"
    sheet.append(["Sheet", "Row", "Name", "Status", "Score", "Before", "Change", "Maximum", "Percentage", "Rank", "Source rank", "Blocker"])
    blockers = {(item.get("sheet_id"), item.get("row")): item["message"] for item in payload["blockers"]}
    for item in payload["sheets"]:
        for group in item["groups"]:
            for student in group["students"]:
                sheet.append(
                    [
                        item["name"],
                        student["row"],
                        student["display_name"],
                        student["status"],
                        student["score"],
                        student.get("before_score"),
                        student.get("change"),
                        group["maximum"],
                        None if student["percentage"] is None else round(student["percentage"], 2),
                        student["rank"],
                        student["source_rank"],
                        blockers.get((item["id"], student["row"]), ""),
                    ]
                )
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="validation.xlsx"'},
    )


@router.get("/imports/{import_id}/outcome-report")
async def outcome_report(request: Request, import_id: str):
    user = await current_user(request)
    if not any("export.xlsx" in grant["actions"] for grant in user["grants"]):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")
    job = await _job(request, user, import_id)
    summary = job.get("summary")
    if not summary:
        raise AppError(422, "import.not_committed", "Commit the workbook before downloading the outcome.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Outcome"
    sheet.append(["Kind", "Inserted", "Updated", "Unchanged", "Skipped", "Rejected"])
    for kind, buckets in (summary.get("counts") or {}).items():
        sheet.append([kind, buckets.get("insert", 0), buckets.get("update", 0), buckets.get("unchanged", 0), buckets.get("skip", 0), buckets.get("reject", 0)])
    detail = workbook.create_sheet("Sheets")
    detail.append(["Sheet", "Status", "Kind", "Inserted", "Updated", "Unchanged", "Skipped", "Rejected"])
    for item in summary.get("by_sheet") or []:
        for kind, buckets in (item.get("counts") or {}).items():
            detail.append([item.get("name"), item.get("status"), kind, buckets.get("insert", 0), buckets.get("update", 0), buckets.get("unchanged", 0), buckets.get("skip", 0), buckets.get("reject", 0)])
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="import-outcome.xlsx"'},
    )


@router.post("/imports/{import_id}/commit")
async def commit_import_route(request: Request, import_id: str) -> dict:
    user = await current_user(request)
    job = await _job(request, user, import_id)
    if not any("marksheet.upload" in grant["actions"] for grant in user["grants"]):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")
    if job["state"] == "committed":
        return {"id": job["_id"], "state": "committed", "outcomes": job.get("outcomes") or []}
    if job["state"] == "cancelled":
        raise AppError(422, "import.cancelled", "This import was cancelled.")
    raw = await request.body()
    confirm_updates = bool(json.loads(raw).get("confirm_updates")) if raw else False
    result = await commit_import(db(request), request.app.state.client, job, user, confirm_updates=confirm_updates)
    return {"id": job["_id"], **result}


@router.post("/imports/{import_id}/cancel")
async def cancel_import(request: Request, import_id: str) -> dict:
    user = await current_user(request)
    job = await _job(request, user, import_id)
    if job["state"] == "committed":
        raise AppError(422, "import.committed", "A committed workbook cannot be cancelled.")
    await db(request).import_jobs.update_one({"_id": job["_id"]}, {"$set": {"state": "cancelled"}})
    return {"id": job["_id"], "state": "cancelled"}


@router.get("/imports")
async def list_imports(request: Request) -> dict:
    user = await current_user(request)
    if not any("marksheet.upload" in grant["actions"] or "marksheet.view" in grant["actions"] for grant in user["grants"]):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")
    items = []
    cursor = db(request).import_jobs.find({"institute_id": user["institute_id"]}).sort("created_at", -1)
    async for job in cursor:
        items.append(
            {
                "id": job["_id"],
                "filename": job.get("filename"),
                "state": job.get("state"),
                "created_at": job.get("created_at").isoformat() if job.get("created_at") else None,
                "outcomes": job.get("outcomes") or [],
                "summary": job.get("summary"),
            }
        )
    return {"items": items}


@router.post("/marksheets")
async def create_marksheet(request: Request, body: MarksheetCreate) -> dict:
    user = await current_user(request)
    if body.exam_type not in {"unit", "part", "chapter"}:
        raise AppError(422, "marksheet.exam_type", "Choose unit, part, or chapter.")
    if body.attempt not in {"original", "retest"}:
        raise AppError(422, "marksheet.attempt", "Choose an original attempt or a retest.")
    database = db(request)
    institute_id = user["institute_id"]
    branch = await database.branches.find_one({"_id": body.branch_id, "institute_id": institute_id})
    course = await database.courses.find_one({"_id": body.course_id, "institute_id": institute_id})
    batch = await database.batches.find_one({"_id": body.batch_id, "institute_id": institute_id})
    subject = await database.subjects.find_one({"_id": body.subject_id, "institute_id": institute_id})
    paper = await database.papers.find_one({"_id": body.paper_id, "institute_id": institute_id})
    if not all((branch, course, batch, subject, paper)):
        raise AppError(404, "not_found", "That academic record was not found.")
    if batch.get("course_id") != course["_id"] or subject.get("course_id") != course["_id"] or paper.get("subject_id") != subject["_id"]:
        raise AppError(422, "marksheet.context", "That batch, subject, and paper do not belong to the same course.")
    if batch.get("branch_id") and batch.get("branch_id") != branch["_id"]:
        raise AppError(422, "marksheet.context", "That batch does not belong to the chosen branch.")
    offering = await database.offerings.find_one({"institute_id": institute_id, "branch_id": branch["_id"], "course_id": course["_id"]})
    if offering is None:
        raise AppError(422, "marksheet.context", "That branch does not offer the chosen course.")
    resource = {
        "institute_id": institute_id,
        "branch_id": branch["_id"],
        "course_id": course["_id"],
        "batch_id": batch["_id"],
        "subject_id": subject["_id"],
        "paper_id": paper["_id"],
    }
    if not allows(user["grants"], "marksheet.upload", resource):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")
    title = body.title.strip()
    identity = "|".join(("manual", name_key(title), body.exam_type, body.attempt, batch["_id"], paper["_id"]))
    existing = await database.assessments.find_one({"institute_id": institute_id, "identity_key": identity})
    if existing:
        found = await database.marksheets.find_one({"assessment_id": existing["_id"], "institute_id": institute_id})
        raise AppError(409, "marksheet.exists", "A marksheet with that title already exists for this batch.", id=found["_id"] if found else None)
    sheet_id = str(ObjectId())
    assessment_id = str(ObjectId())
    attempt_id = str(ObjectId())
    client = request.app.state.client
    async with await client.start_session() as session:
        async with session.start_transaction():
            await database.assessments.insert_one(
                {
                    "_id": assessment_id,
                    "institute_id": institute_id,
                    "identity_key": identity,
                    "title": title,
                    "title_key": name_key(title),
                    "exam_type": body.exam_type,
                    "format": "manual",
                    "subject_id": subject["_id"],
                    "paper_id": paper["_id"],
                    "branch_id": branch["_id"],
                    "course_id": course["_id"],
                    "batch_ids": [batch["_id"]],
                    "exam_date": body.exam_date,
                    "maximum": body.maximum,
                },
                session=session,
            )
            await database.attempts.insert_one(
                {
                    "_id": attempt_id,
                    "institute_id": institute_id,
                    "assessment_id": assessment_id,
                    "kind": body.attempt,
                    "maximum": body.maximum,
                },
                session=session,
            )
            await database.marksheets.insert_one(
                {
                    "_id": sheet_id,
                    "institute_id": institute_id,
                    "assessment_id": assessment_id,
                    "attempt_id": attempt_id,
                    "cohort_key": batch["_id"],
                    "branch_id": branch["_id"],
                    "course_id": course["_id"],
                    "subject_id": subject["_id"],
                    "paper_id": paper["_id"],
                    "batch_ids": [batch["_id"]],
                    "title": title,
                    "maximum": body.maximum,
                    "exam_type": body.exam_type,
                    "exam_date": body.exam_date,
                    "status": "draft",
                    "revision": 1,
                    "active_revision": None,
                    "source_sheet": "Entered here",
                    "uploader_id": user["_id"],
                    "edit_version": 0,
                    "attempt_kind": body.attempt,
                },
                session=session,
            )
    await database.audit_events.insert_one(
        {
            "_id": str(ObjectId()),
            "institute_id": institute_id,
            "action": "marksheet.create",
            "actor_id": user["_id"],
            "at": _now(),
            "context": {"marksheet_id": sheet_id, "title": title},
        }
    )
    return {"id": sheet_id, "status": "draft"}


@router.post("/marksheets/{marksheet_id}/results")
async def add_result(request: Request, marksheet_id: str, body: ResultCreate) -> dict:
    user, sheet = await _editable(request, marksheet_id, "marksheet.edit_draft")
    if sheet.get("status") not in {"draft", "submitted"}:
        raise AppError(422, "marksheet.state", "Add students on a draft. A published sheet uses a correction.")
    if body.status not in {"scored", "absent", "missing"}:
        raise AppError(422, "import.invalid_mark", "Choose scored, absent, or missing.")
    score = body.score
    if body.status == "scored":
        if score is None or score < 0 or (sheet.get("maximum") is not None and score > sheet["maximum"]):
            raise AppError(422, "import.invalid_mark", "The score must be from 0 through the maximum.")
    else:
        score = None
    enrollment = await db(request).enrollments.find_one(
        {"student_id": body.student_id, "batch_id": {"$in": sheet.get("batch_ids") or []}}
    )
    if enrollment is None:
        raise AppError(422, "marksheet.not_enrolled", "That student is not enrolled in this batch.")
    revision = sheet.get("revision") or 1
    duplicate = await db(request).results.find_one(
        {"marksheet_id": sheet["_id"], "revision": revision, "student_id": body.student_id}
    )
    if duplicate:
        raise AppError(409, "marksheet.duplicate", "That student is already on this marksheet.")
    result_id = str(ObjectId())
    await db(request).results.insert_one(
        {
            "_id": result_id,
            "institute_id": user["institute_id"],
            "marksheet_id": sheet["_id"],
            "enrollment_id": enrollment["_id"],
            "student_id": body.student_id,
            "revision": revision,
            "active": False,
            "status": body.status,
            "score": score,
            "maximum": sheet.get("maximum"),
            "exam_date": sheet.get("exam_date"),
            "title": sheet.get("title"),
            "assessment_id": sheet.get("assessment_id"),
            "attempt_kind": sheet.get("attempt_kind") or "original",
            "exam_type": sheet.get("exam_type"),
            "counts_in_aggregate": True,
            "branch_id": sheet.get("branch_id"),
            "course_id": sheet.get("course_id"),
            "batch_id": enrollment["batch_id"],
            "subject_id": sheet.get("subject_id"),
            "paper_id": sheet.get("paper_id"),
            "source": {"sheet": "Entered here"},
        }
    )
    version = sheet.get("edit_version", 0) + 1
    await db(request).marksheets.update_one({"_id": sheet["_id"]}, {"$set": {"edit_version": version}})
    await _refresh_sheet_marks(request, sheet)
    return {"id": result_id, "edit_version": version}


@router.get("/marksheets")
async def list_marksheets(
    request: Request,
    status: str = "",
    exam_type: str = "",
    branch_id: str = "",
    course_id: str = "",
    batch_id: str = "",
    subject_id: str = "",
    paper_id: str = "",
    exam_date: str = "",
    exam_date_from: str = "",
    exam_date_to: str = "",
    q: str = "",
    sort: str = "title",
    page: int = 0,
    page_size: int = 8,
) -> dict:
    user = await current_user(request)
    if not any("marksheet.view" in grant["actions"] for grant in user["grants"]):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")
    query: dict = {"institute_id": user["institute_id"]}
    for field, value in (
        ("status", status),
        ("exam_type", exam_type),
        ("branch_id", branch_id),
        ("course_id", course_id),
        ("subject_id", subject_id),
        ("paper_id", paper_id),
        ("exam_date", exam_date),
    ):
        if value:
            query[field] = value
    if not exam_date and (exam_date_from or exam_date_to):
        window = {}
        if exam_date_from:
            window["$gte"] = exam_date_from
        if exam_date_to:
            window["$lte"] = exam_date_to
        query["exam_date"] = window
    if batch_id:
        query["batch_ids"] = batch_id
    named_sheets: set[str] = set()
    if q:
        matched_students = await db(request).students.find(
            {"institute_id": user["institute_id"], "display_name": {"$regex": q, "$options": "i"}},
            {"_id": 1},
        ).to_list(length=200)
        student_ids = [item["_id"] for item in matched_students]
        if student_ids:
            named_sheets = {
                item["marksheet_id"]
                for item in await db(request).results.find(
                    {"student_id": {"$in": student_ids}},
                    {"marksheet_id": 1},
                ).to_list(length=2000)
            }
    items = []
    names = {}
    cursor = db(request).marksheets.find(query)
    async for sheet in cursor:
        resource = {
            "institute_id": sheet["institute_id"],
            "branch_id": sheet.get("branch_id"),
            "course_id": sheet.get("course_id"),
            "subject_id": sheet.get("subject_id"),
        }
        if not allows(user["grants"], "marksheet.view", resource):
            continue
        summary = _marksheet_summary(sheet)
        haystack = " ".join(str(summary.get(field) or "") for field in ("title", "source_file", "source_sheet")).casefold()
        if q and q.casefold() not in haystack and sheet["_id"] not in named_sheets:
            continue
        for field, key in (("uploader_id", "uploader_name"), ("reviewer_id", "reviewer_name")):
            person_id = sheet.get(field)
            if person_id and person_id not in names:
                person = await db(request).users.find_one({"_id": person_id})
                names[person_id] = person["name"] if person else ""
            summary[key] = names.get(person_id, "")
        items.append(summary)
    items.sort(key=lambda item: str(item.get(sort if sort in {"exam_date", "status", "title"} else "title") or "").casefold())
    total = len(items)
    current = page or 1
    if page:
        start = (page - 1) * page_size
        items = items[start : start + page_size]
    pages = max(1, (total + page_size - 1) // page_size)
    return {"items": items, "total": total, "page": current, "pages": pages}


@router.post("/marksheets/{marksheet_id}/submit")
async def submit_route(request: Request, marksheet_id: str) -> dict:
    user, sheet = await _editable(request, marksheet_id, "marksheet.submit")
    return await submit_marksheet(db(request), user, sheet)


@router.post("/marksheets/{marksheet_id}/reject")
async def reject_route(request: Request, marksheet_id: str) -> dict:
    user, sheet = await _editable(request, marksheet_id, "marksheet.publish")
    return await reject_marksheet(db(request), user, sheet)


@router.post("/marksheets/{marksheet_id}/correct")
async def correct_route(request: Request, marksheet_id: str, body: ReasonBody) -> dict:
    user, sheet = await _editable(request, marksheet_id, "marksheet.correct")
    return await correct_marksheet(
        db(request), request.app.state.client, user, sheet, body.reason, body.changes, body.edit_version
    )


@router.post("/marksheets/{marksheet_id}/withdraw")
async def withdraw_route(request: Request, marksheet_id: str, body: ReasonBody) -> dict:
    user, sheet = await _editable(request, marksheet_id, "marksheet.withdraw")
    return await withdraw_marksheet(db(request), request.app.state.client, user, sheet, body.reason, body.edit_version)


@router.patch("/marksheets/{marksheet_id}")
async def patch_marksheet(request: Request, marksheet_id: str, body: MarksheetPatch) -> dict:
    wants_bands = "bands" in body.model_fields_set
    if body.exam_date and wants_bands:
        _user, sheet = await _editable(request, marksheet_id, "marksheet.edit_draft")
    elif wants_bands:
        sheet = await db(request).marksheets.find_one({"_id": marksheet_id})
        action = "marksheet.edit_draft" if sheet and sheet.get("status") == "draft" else "marksheet.correct"
        _user, sheet = await _editable(request, marksheet_id, action)
    else:
        _user, sheet = await _editable(request, marksheet_id, "marksheet.edit_draft")
    updates = {}
    unset = {}
    if body.exam_date:
        if sheet.get("status") != "draft":
            raise AppError(422, "marksheet.state", "Only a draft can be edited.")
        updates["exam_date"] = body.exam_date
    if wants_bands:
        if body.bands is None:
            unset["bands"] = ""
        else:
            from app.calculating import prepare_bands

            try:
                updates["bands"] = prepare_bands([item.model_dump() for item in body.bands])
            except ValueError as error:
                raise AppError(422, "policy.bands", str(error)) from None
    if updates or unset:
        if body.exam_date:
            updates["edit_version"] = sheet.get("edit_version", 0) + 1
        change = {}
        if updates:
            change["$set"] = updates
        if unset:
            change["$unset"] = unset
        await db(request).marksheets.update_one({"_id": sheet["_id"]}, change)
    from app.calculating import public_bands

    saved = updates.get("bands")
    if "bands" in unset:
        shown = None
    elif saved:
        shown = public_bands(saved)
    elif sheet.get("bands"):
        shown = public_bands(sheet.get("bands"))
    else:
        shown = None
    return {"id": sheet["_id"], **{key: value for key, value in updates.items() if key != "bands"}, "bands": shown}


@router.patch("/marksheets/{marksheet_id}/results/{result_id}")
async def patch_result(request: Request, marksheet_id: str, result_id: str, body: ResultPatch) -> dict:
    _user, sheet = await _editable(request, marksheet_id, "marksheet.edit_draft")
    require_current_version(sheet, body.edit_version)
    if sheet.get("status") != "draft":
        raise AppError(422, "marksheet.state", "Only a draft result can be corrected here.")
    result = await db(request).results.find_one({"_id": result_id, "marksheet_id": sheet["_id"]})
    if result is None:
        raise AppError(404, "not_found", "That record was not found.")
    updates = {"previous_score": result.get("score")}
    if body.status in {"absent", "missing"}:
        updates["status"] = body.status
        updates["score"] = None
    elif body.status == "scored" or body.score is not None:
        if body.score is None:
            raise AppError(422, "import.invalid_mark", "Enter a score, or mark the row missing or absent.")
        if body.score < 0 or (sheet.get("maximum") is not None and body.score > sheet["maximum"]):
            raise AppError(422, "import.invalid_mark", "The score must be from 0 through the maximum.")
        updates["score"] = body.score
        updates["status"] = "scored"
    elif body.status:
        raise AppError(422, "import.invalid_mark", "Choose scored, absent, or missing.")
    version = sheet.get("edit_version", 0) + 1
    await db(request).results.update_one({"_id": result_id}, {"$set": updates})
    await db(request).marksheets.update_one({"_id": sheet["_id"]}, {"$set": {"edit_version": version}})
    await _refresh_sheet_marks(request, sheet)
    return {"id": result_id, "edit_version": version, **updates}


async def _refresh_sheet_marks(request: Request, sheet: dict) -> None:
    from app.calculating import band, dense_ranks, display_percentage, percentage

    revision = sheet.get("revision") or 1
    rows = await db(request).results.find({"marksheet_id": sheet["_id"], "revision": revision}).to_list(length=500)
    policy = {"bands": sheet["bands"]} if sheet.get("bands") else await db(request).policies.find_one(
        {"institute_id": sheet["institute_id"]}, sort=[("version", -1)]
    )
    ranked = dense_ranks([
        {"id": row["_id"], "status": row.get("status"), "score": row.get("score")}
        for row in rows
    ])
    for row in rows:
        exact = percentage(row.get("score"), row.get("maximum") if row.get("maximum") is not None else sheet.get("maximum"))
        exact = exact if row.get("status") == "scored" else None
        await db(request).results.update_one(
            {"_id": row["_id"]},
            {"$set": {
                "percentage": display_percentage(exact),
                "band": band(exact, policy) if row.get("status") == "scored" else None,
                "rank": ranked.get(row["_id"]),
            }},
        )


@router.get("/catalog")
async def catalog(request: Request) -> dict:
    user = await current_user(request)
    if not any(action in grant["actions"] for grant in user["grants"] for action in ("student.lookup", "catalog.correct", "catalog.manage")):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")
    database = db(request)
    branches = await database.branches.find({"institute_id": user["institute_id"]}).to_list(length=200)
    courses = await database.courses.find({"institute_id": user["institute_id"]}).to_list(length=200)
    batches = await database.batches.find({"institute_id": user["institute_id"]}).to_list(length=200)
    subjects = await database.subjects.find({"institute_id": user["institute_id"]}).to_list(length=200)
    students = await database.students.find({"institute_id": user["institute_id"]}).to_list(length=500)
    papers = await database.papers.find({"institute_id": user["institute_id"]}).to_list(length=200)
    offerings = await database.offerings.find({"institute_id": user["institute_id"]}).to_list(length=200)
    return {
        "branches": [_catalog_row(item) for item in branches],
        "courses": [_catalog_row(item) for item in courses],
        "batches": [_catalog_row(item) for item in batches],
        "subjects": [_catalog_row(item) for item in subjects],
        "papers": [_catalog_row(item) for item in papers],
        "students": [_catalog_row(item) for item in students],
        "offerings": [
            {"id": item["_id"], "branch_id": item.get("branch_id"), "course_id": item.get("course_id")}
            for item in offerings
        ],
    }


@router.patch("/catalog/{kind}/{record_id}")
async def patch_catalog(request: Request, kind: str, record_id: str, body: CatalogPatch) -> dict:
    user = await current_user(request)
    if not any("catalog.correct" in grant["actions"] or "catalog.manage" in grant["actions"] for grant in user["grants"]):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")
    collection = {
        "branch": "branches",
        "course": "courses",
        "batch": "batches",
        "subject": "subjects",
        "paper": "papers",
        "student": "students",
    }.get(kind)
    if collection is None:
        raise AppError(404, "not_found", "That record was not found.")
    record = await db(request)[collection].find_one({"_id": record_id, "institute_id": user["institute_id"]})
    if record is None:
        raise AppError(404, "not_found", "That record was not found.")
    updates = {}
    unsets = {}
    sent = body.model_fields_set
    if "name" in sent:
        text = " ".join(str(body.name or "").split())
        if not text:
            raise AppError(422, "catalog.empty", "Enter a name.")
        current_name = record.get("name") or record.get("display_name") or ""
        if text != current_name:
            updates["name"] = text
            updates["previous_name"] = current_name
            if kind == "student":
                updates["display_name"] = text
            else:
                updates["name_key"] = name_key(text)
    if body.archived is not None:
        updates["archived"] = body.archived
    if kind == "branch":
        detail_sets, detail_unsets = place_changes(body, record)
        updates.update(detail_sets)
        unsets.update(detail_unsets)
    if kind == "course" and "notes" in sent:
        notes = place_text(body.notes, 500, "notes")
        if notes != str(record.get("notes") or ""):
            if notes:
                updates["notes"] = notes
            else:
                unsets["notes"] = ""
    if kind == "batch":
        started = calendar_date(body.started_on or "", "started_on") if "started_on" in sent else str(record.get("started_on") or "")
        ended = calendar_date(body.ended_on or "", "ended_on") if "ended_on" in sent else str(record.get("ended_on") or "")
        if started and ended and ended < started:
            raise AppError(422, "catalog.date", "The until date is before the from date.", field="ended_on")
        for field, value in (("started_on", started if "started_on" in sent else None), ("ended_on", ended if "ended_on" in sent else None)):
            if value is None:
                continue
            if value != str(record.get(field) or ""):
                if value:
                    updates[field] = value
                else:
                    unsets[field] = ""
        if "timings" in sent:
            timings = place_text(body.timings, 120, "timings")
            if timings != str(record.get("timings") or ""):
                if timings:
                    updates["timings"] = timings
                else:
                    unsets["timings"] = ""
    if not updates and not unsets:
        raise AppError(422, "catalog.empty", "Choose a new name or an archive state.")
    change = {}
    if updates:
        change["$set"] = updates
    if unsets:
        change["$unset"] = unsets
    try:
        await db(request)[collection].update_one({"_id": record_id}, change)
    except DuplicateKeyError:
        raise AppError(409, "catalog.exists", "That name is already used.") from None
    await db(request).audit_events.insert_one(
        {
            "_id": str(ObjectId()),
            "institute_id": user["institute_id"],
            "action": "catalog.correct",
            "actor_id": user["_id"],
            "at": _now(),
            "context": {"kind": kind, "id": record_id, **updates},
        }
    )
    return {"id": record_id, **updates}


@router.post("/students/merge-preview")
async def merge_preview(request: Request, body: MergeBody) -> dict:
    user = await current_user(request)
    if not any("catalog.manage" in grant["actions"] for grant in user["grants"]):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")
    source, target = await _merge_pair(request, user["institute_id"], body)
    source_rows = await db(request).enrollments.find({"student_id": source["_id"]}).to_list(length=100)
    target_rows = await db(request).enrollments.find({"student_id": target["_id"]}).to_list(length=100)
    target_batches = {item["batch_id"] for item in target_rows}
    moving = [item for item in source_rows if item["batch_id"] not in target_batches]
    kept = [item for item in source_rows if item["batch_id"] in target_batches]
    facts = await directory_facts(db(request), user["institute_id"])
    return {
        "source": source.get("display_name"),
        "target": target.get("display_name"),
        "source_code": source.get("student_code"),
        "target_code": target.get("student_code"),
        "move": [item["_id"] for item in moving],
        "kept_separate": [item["_id"] for item in kept],
        "moves": enrollment_context(moving, facts),
        "kept": enrollment_context(kept, facts),
        "alias": source.get("name_key"),
    }


@router.post("/students/merge")
async def merge_students(request: Request, body: MergeBody) -> dict:
    user = await current_user(request)
    if not any("catalog.manage" in grant["actions"] for grant in user["grants"]):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")
    if not body.acknowledge:
        raise AppError(422, "catalog.acknowledge", "Review the merge impact and confirm it.")
    preview = await merge_preview(request, body)
    source, target = await _merge_pair(request, user["institute_id"], body)
    for enrollment_id in preview["move"]:
        await db(request).enrollments.update_one({"_id": enrollment_id}, {"$set": {"student_id": target["_id"]}})
    await db(request).aliases.insert_one(
        {
            "_id": str(ObjectId()),
            "institute_id": user["institute_id"],
            "student_id": target["_id"],
            "name_key": source.get("name_key"),
            "display_name": source.get("display_name"),
        }
    )
    await db(request).students.update_one({"_id": source["_id"]}, {"$set": {"archived": True, "merged_into": target["_id"]}})
    await db(request).audit_events.insert_one(
        {
            "_id": str(ObjectId()),
            "institute_id": user["institute_id"],
            "action": "student.merge",
            "actor_id": user["_id"],
            "at": _now(),
            "context": {"source_id": source["_id"], "target_id": target["_id"], "moved": preview["move"]},
        }
    )
    return {"id": target["_id"], "moved": len(preview["move"]), "archived_id": source["_id"]}


@router.get("/students/{student_id}/enrollments")
async def enrollment_history(request: Request, student_id: str) -> dict:
    user = await current_user(request)
    student, visible, owned, facts = await _student_directory(request, user, student_id)
    aliases = await db(request).aliases.find({"student_id": student_id}).to_list(length=20)
    contexts = enrollment_context(visible, facts)
    visible_by_id = {row["_id"]: row for row in visible}
    for item in contexts:
        source = visible_by_id.get(item["id"]) or {}
        item["manage"] = bool(source) and allows(user["grants"], "student.manage", enrollment_resource(source))
    return {
        "id": student["_id"],
        "student_code": student.get("student_code"),
        "display_name": student.get("display_name"),
        "archived": bool(student.get("archived")),
        "edit_version": student.get("edit_version", 0),
        "photo_id": student.get("photo_file_id") or None,
        **_profile_view(student),
        "enrollments": contexts,
        "aliases": [
            {"id": item["_id"], "display_name": item.get("display_name")}
            for item in aliases
            if item.get("display_name")
        ],
        "actions": {
            "edit": identity_editable(user["grants"], owned),
            "card": card_visible(user["grants"], visible),
            "enroll": any("student.manage" in (grant.get("actions") or []) for grant in user["grants"]),
        },
    }


@router.patch("/students/{student_id}")
async def rename_student(request: Request, student_id: str, body: StudentName) -> dict:
    user = await current_user(request)
    if body.student_code is not None:
        raise AppError(422, "student.identity", "The student id stays as it is.")
    student, _visible, owned, _facts = await _student_directory(request, user, student_id)
    if not identity_editable(user["grants"], owned):
        raise AppError(403, "auth.forbidden", "This name is shared across enrollments you do not manage.")
    current = student.get("edit_version", 0)
    if body.edit_version is None or body.edit_version != current:
        raise AppError(
            409,
            "student.conflict",
            "Someone else saved this student. Compare the details and try again.",
            **_conflict_details(student, current),
        )
    display_name = " ".join(body.display_name.split())
    if not display_name:
        raise AppError(422, "student.name", "Enter the student's name.")
    if len(display_name) > 120:
        raise AppError(422, "student.name", "Use a name of 120 characters or fewer.")
    key = name_key(display_name)
    await _reject_shared_name(request, user["institute_id"], student_id, key, owned)
    profile_set, profile_unset, profile_before, profile_after = _profile_changes(body, student)
    version_query: dict = {"_id": student_id, "institute_id": user["institute_id"]}
    if current == 0:
        version_query["$or"] = [{"edit_version": 0}, {"edit_version": {"$exists": False}}]
    else:
        version_query["edit_version"] = current
    changes: dict = {"$set": {
        "display_name": display_name,
        "name_key": key,
        "previous_name": student.get("display_name"),
        "edit_version": current + 1,
        **profile_set,
    }}
    if profile_unset:
        changes["$unset"] = profile_unset
    updated = await db(request).students.update_one(version_query, changes)
    if updated.matched_count == 0:
        fresh = await db(request).students.find_one({"_id": student_id, "institute_id": user["institute_id"]})
        current_student = fresh or student
        raise AppError(
            409,
            "student.conflict",
            "Someone else saved this student. Compare the details and try again.",
            **_conflict_details(current_student, current_student.get("edit_version", 0)),
        )
    if display_name != student.get("display_name") and student.get("display_name"):
        await db(request).aliases.insert_one(
            {
                "_id": str(ObjectId()),
                "institute_id": user["institute_id"],
                "student_id": student_id,
                "name_key": student.get("name_key"),
                "display_name": student.get("display_name"),
            }
        )
    await write_audit(
        request,
        user,
        "student.update",
        {"display_name": student.get("display_name"), "student_code": student.get("student_code"), **profile_before},
        {"display_name": display_name, "student_code": student.get("student_code"), **profile_after},
        {"student_id": student_id},
    )
    return {
        "id": student_id,
        "display_name": display_name,
        "student_code": student.get("student_code"),
        "edit_version": current + 1,
    }


def _profile_view(student: dict) -> dict:
    return {field: student.get(field) or None for field in _PROFILE_FIELDS}


def _conflict_details(student: dict, version: int) -> dict:
    details = {
        "current_display_name": student.get("display_name") or "",
        "edit_version": version,
    }
    for field in _PROFILE_FIELDS:
        details[f"current_{field}"] = student.get(field) or ""
    return details


def _profile_text(value: str | None, limit: int, field: str) -> str:
    text = " ".join(str(value or "").split())
    if field in {"street", "notes"}:
        text = str(value or "").strip()
    if len(text) > limit:
        raise AppError(422, "student.profile", f"Use {limit} characters or fewer.", field=field)
    return text


def _birth_date(value: str | None) -> str | None:
    text = str(value or "").strip()
    if not text:
        return None
    try:
        parsed = datetime.strptime(text, "%Y-%m-%d").date()
    except ValueError as exc:
        raise AppError(422, "student.profile", "Use a date of birth as YYYY-MM-DD.", field="date_of_birth") from exc
    if parsed > _now().date():
        raise AppError(422, "student.profile", "The date of birth is in the future.", field="date_of_birth")
    return text


def _profile_changes(body: StudentName, student: dict) -> tuple[dict, dict, dict, dict]:
    sent = body.model_fields_set
    sets: dict = {}
    unsets: dict = {}
    before: dict = {}
    after: dict = {}

    def keep(field: str, text: str) -> None:
        previous = str(student.get(field) or "")
        if text == previous:
            return
        before[field] = previous or "Not recorded"
        after[field] = text or "Not recorded"
        if text:
            sets[field] = text
        else:
            unsets[field] = ""

    for field, limit in _PROFILE_LIMITS.items():
        if field not in sent:
            continue
        keep(field, _profile_text(getattr(body, field), limit, field))
    if "date_of_birth" in sent:
        born = _birth_date(body.date_of_birth)
        keep("date_of_birth", born or "")
    if "blood_group" in sent:
        group = str(body.blood_group or "").strip()
        if group and group not in _BLOOD_GROUPS:
            raise AppError(422, "student.profile", "Choose a blood group.", field="blood_group")
        keep("blood_group", group)
    if "pin" in sent:
        keep("pin", _pin(body.pin))
    if "passing_year" in sent:
        keep("passing_year", _passing_year(body.passing_year))
    return sets, unsets, before, after


def _pin(value: str | None) -> str:
    text = "".join(str(value or "").split())
    if text and (len(text) != 6 or not text.isdigit()):
        raise AppError(422, "student.profile", "Use a 6-digit PIN.", field="pin")
    return text


def _passing_year(value: str | None) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    year = int(text) if len(text) == 4 and text.isdigit() else 0
    if year < 1950 or year > _now().year:
        raise AppError(422, "student.profile", "Use a year of passing as YYYY.", field="passing_year")
    return text


def _photo_type(data: bytes) -> tuple[str, str] | None:
    if data.startswith(b"\xff\xd8\xff"):
        return "image/jpeg", "jpg"
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png", "png"
    if len(data) >= 12 and data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp", "webp"
    return None


def _calendar_date(value: str, field: str) -> str | None:
    text = value.strip()
    if not text:
        return None
    if len(text) != 10 or text[4] != "-" or text[7] != "-":
        raise AppError(422, "enrollment.date", "Use a date as YYYY-MM-DD.", field=field)
    try:
        datetime.strptime(text, "%Y-%m-%d")
    except ValueError as exc:
        raise AppError(422, "enrollment.date", "Use a date as YYYY-MM-DD.", field=field) from exc
    return text


def _require_identity_edit(user: dict, owned: list[dict]) -> None:
    if not identity_editable(user["grants"], owned):
        raise AppError(403, "auth.forbidden", "This record is shared across enrollments you do not manage.")


async def _managed_enrollment(request: Request, user: dict, student_id: str, enrollment_id: str) -> tuple[dict, dict]:
    _student, visible, _owned, facts = await _student_directory(request, user, student_id)
    row = next((item for item in visible if item["_id"] == enrollment_id), None)
    if row is None:
        raise AppError(404, "not_found", "That record was not found.")
    if not allows(user["grants"], "student.manage", enrollment_resource(row)):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")
    return row, facts


@router.get("/students/{student_id}/photo")
async def student_photo(request: Request, student_id: str):
    user = await current_user(request)
    student, _visible, _owned, _facts = await _student_directory(request, user, student_id)
    file_id = student.get("photo_file_id")
    if not file_id:
        raise AppError(404, "not_found", "That record was not found.")
    file_doc = await db(request).files.find_one(
        {"_id": file_id, "institute_id": user["institute_id"], "kind": "student_photo"}
    )
    path = Path(file_doc["path"]) if file_doc else None
    if file_doc is None or path is None or not path.is_file():
        raise AppError(404, "not_found", "That record was not found.")
    return FileResponse(path, media_type=file_doc.get("media_type") or "image/jpeg")


@router.post("/students/{student_id}/photo")
async def upload_student_photo(request: Request, student_id: str, photo: UploadFile = File(...)) -> dict:
    user = await current_user(request)
    student, _visible, owned, _facts = await _student_directory(request, user, student_id)
    _require_identity_edit(user, owned)
    data = await photo.read()
    if not data or len(data) > _MAX_PHOTO:
        raise AppError(422, "student.photo", "Use a JPEG, PNG, or WebP photo of 2 MB or less.")
    kind = _photo_type(data)
    if kind is None:
        raise AppError(422, "student.photo", "Use a JPEG, PNG, or WebP photo.")
    media_type, extension = kind
    file_id = str(ObjectId())
    path = request.app.state.settings.data_dir / "files" / f"{file_id}.{extension}"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    await db(request).files.insert_one(
        {
            "_id": file_id,
            "institute_id": user["institute_id"],
            "kind": "student_photo",
            "media_type": media_type,
            "sha256": hashlib.sha256(data).hexdigest(),
            "path": str(path),
            "filename": f"photo.{extension}",
            "created_at": _now(),
        }
    )
    previous = student.get("photo_file_id")
    updated = await db(request).students.update_one(
        {"_id": student_id, "institute_id": user["institute_id"]},
        {"$set": {"photo_file_id": file_id}},
    )
    if updated.matched_count == 0:
        path.unlink(missing_ok=True)
        await db(request).files.delete_one({"_id": file_id})
        raise AppError(404, "not_found", "That record was not found.")
    await write_audit(
        request,
        user,
        "student.photo",
        {"photo": "Recorded"} if previous else None,
        {"photo": "Recorded"},
        {"student_id": student_id},
    )
    return {"id": student_id, "photo_id": file_id}


@router.delete("/students/{student_id}/photo")
async def remove_student_photo(request: Request, student_id: str) -> dict:
    user = await current_user(request)
    student, _visible, owned, _facts = await _student_directory(request, user, student_id)
    _require_identity_edit(user, owned)
    previous = student.get("photo_file_id")
    if previous:
        await db(request).students.update_one(
            {"_id": student_id, "institute_id": user["institute_id"]},
            {"$unset": {"photo_file_id": ""}},
        )
        await write_audit(
            request,
            user,
            "student.photo",
            {"photo": "Recorded"},
            None,
            {"student_id": student_id},
        )
    return {"id": student_id, "photo_id": None}


@router.post("/students/{student_id}/aliases")
async def add_student_alias(request: Request, student_id: str, body: StudentAlias) -> dict:
    user = await current_user(request)
    student, _visible, owned, _facts = await _student_directory(request, user, student_id)
    _require_identity_edit(user, owned)
    display_name = " ".join(body.display_name.split())
    if not display_name or len(display_name) > 120:
        raise AppError(422, "student.name", "Enter a previous name of 120 characters or fewer.")
    key = name_key(display_name)
    if key == student.get("name_key"):
        raise AppError(422, "student.alias", "That is the student's current name.")
    existing = await db(request).aliases.find_one({"student_id": student_id, "name_key": key})
    if existing:
        raise AppError(409, "student.alias", "That previous name is already recorded.")
    alias_id = str(ObjectId())
    await db(request).aliases.insert_one(
        {
            "_id": alias_id,
            "institute_id": user["institute_id"],
            "student_id": student_id,
            "name_key": key,
            "display_name": display_name,
        }
    )
    await write_audit(
        request,
        user,
        "student.alias",
        None,
        {"alias": display_name},
        {"student_id": student_id},
    )
    return {"id": alias_id, "display_name": display_name}


@router.delete("/students/{student_id}/aliases/{alias_id}")
async def remove_student_alias(request: Request, student_id: str, alias_id: str) -> dict:
    user = await current_user(request)
    _student, _visible, owned, _facts = await _student_directory(request, user, student_id)
    _require_identity_edit(user, owned)
    alias = await db(request).aliases.find_one(
        {"_id": alias_id, "student_id": student_id, "institute_id": user["institute_id"]}
    )
    if alias is None:
        raise AppError(404, "not_found", "That record was not found.")
    await db(request).aliases.delete_one({"_id": alias_id})
    await write_audit(
        request,
        user,
        "student.alias",
        {"alias": alias.get("display_name")},
        None,
        {"student_id": student_id},
    )
    return {"id": alias_id}


@router.patch("/students/{student_id}/enrollments/{enrollment_id}")
async def update_enrollment_dates(request: Request, student_id: str, enrollment_id: str, body: EnrollmentDates) -> dict:
    user = await current_user(request)
    row, facts = await _managed_enrollment(request, user, student_id, enrollment_id)
    started = _calendar_date(body.started_on, "started_on")
    ended = _calendar_date(body.ended_on, "ended_on")
    if started and ended and ended < started:
        raise AppError(422, "enrollment.date", "The until date is before the from date.")
    update: dict = {"$set": {}, "$unset": {}}
    if started:
        update["$set"]["started_on"] = started
    else:
        update["$unset"]["started_on"] = ""
    if ended:
        update["$set"]["ended_on"] = ended
        update["$set"]["status"] = "historical"
    else:
        update["$unset"]["ended_on"] = ""
        update["$unset"]["status"] = ""
    if not update["$set"]:
        del update["$set"]
    await db(request).enrollments.update_one({"_id": enrollment_id, "student_id": student_id}, update)
    batch_name = facts["batch_name"].get(row.get("batch_id"), "")
    await write_audit(
        request,
        user,
        "enrollment.update",
        {
            "batch_name": batch_name,
            "started_on": row.get("started_on") or "Not recorded",
            "ended_on": row.get("ended_on") or "Not recorded",
        },
        {
            "batch_name": batch_name,
            "started_on": started or "Not recorded",
            "ended_on": ended or "Not recorded",
        },
        {
            "student_id": student_id,
            "branch_id": row.get("branch_id"),
            "course_id": row.get("course_id"),
            "batch_id": row.get("batch_id"),
        },
    )
    return {"id": enrollment_id, "started_on": started, "ended_on": ended}


@router.post("/students/{student_id}/enrollments")
async def add_student_enrollment(request: Request, student_id: str, body: EnrollmentCreate) -> dict:
    user = await current_user(request)
    student, _visible, _owned, _facts = await _student_directory(request, user, student_id)
    started = _calendar_date(body.started_on, "started_on")
    batch = await db(request).batches.find_one({"_id": body.batch_id, "institute_id": user["institute_id"]})
    if batch is None or batch.get("branch_id") != body.branch_id or batch.get("course_id") != body.course_id:
        raise AppError(422, "enrollment.context", "Choose a branch, course, and batch that belong together.")
    resource = {
        "institute_id": user["institute_id"],
        "branch_id": body.branch_id,
        "course_id": body.course_id,
        "batch_id": body.batch_id,
    }
    if not allows(user["grants"], "student.manage", resource):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")
    already = await db(request).enrollments.find_one({"student_id": student_id, "batch_id": body.batch_id})
    if already:
        raise AppError(409, "enrollment.exists", "This student is already in that batch.")
    aliases = await db(request).aliases.find({"student_id": student_id}).to_list(length=20)
    keys = {student.get("name_key")}
    keys.update(item.get("name_key") for item in aliases if item.get("name_key"))
    keys.discard(None)
    taken = await db(request).enrollment_names.find(
        {"batch_id": body.batch_id, "name_key": {"$in": list(keys)}, "student_id": {"$ne": student_id}}
    ).to_list(length=5)
    if taken:
        raise AppError(
            409,
            "student.identity",
            "Another student in that batch already uses this name. Merge them only from the merge action.",
        )
    enrollment_id = str(ObjectId())
    enrollment = {
        "_id": enrollment_id,
        "institute_id": user["institute_id"],
        "student_id": student_id,
        "branch_id": body.branch_id,
        "course_id": body.course_id,
        "batch_id": body.batch_id,
    }
    if started:
        enrollment["started_on"] = started
    names = [
        {
            "_id": str(ObjectId()),
            "institute_id": user["institute_id"],
            "batch_id": body.batch_id,
            "student_id": student_id,
            "name_key": key,
        }
        for key in keys
    ]
    try:
        async with await request.app.state.client.start_session() as session:
            async with session.start_transaction():
                await db(request).enrollments.insert_one(enrollment, session=session)
                if names:
                    await db(request).enrollment_names.insert_many(names, session=session)
    except DuplicateKeyError as exc:
        raise AppError(
            409,
            "student.identity",
            "Another student in that batch already uses this name. Merge them only from the merge action.",
        ) from exc
    await write_audit(
        request,
        user,
        "enrollment.create",
        None,
        {"batch_name": batch.get("name") or "", "started_on": started or "Not recorded"},
        resource | {"student_id": student_id},
    )
    return {"id": enrollment_id, "batch_id": body.batch_id, "started_on": started}


@router.get("/students/{student_id}/history")
async def student_history(request: Request, student_id: str) -> dict:
    user = await current_user(request)
    _student, visible, _owned, _facts = await _student_directory(request, user, student_id)
    if not history_permitted(user["grants"], visible):
        return {"permitted": False, "items": []}
    events = (
        await db(request)
        .audit_events.find(
            {
                "institute_id": user["institute_id"],
                "$or": [
                    {"scope.student_id": student_id},
                    {"context.student_id": student_id},
                    {"context.source_id": student_id},
                    {"context.target_id": student_id},
                    {"context.kind": "student", "context.id": student_id},
                ],
            }
        )
        .sort("at", -1)
        .to_list(length=100)
    )
    chosen = [event for event in events if history_visible(user["grants"], event, student_id, visible)]
    actor_names: dict[str, str] = {}
    for event in chosen:
        actor_id = event.get("actor_id")
        if actor_id and actor_id not in actor_names:
            person = await db(request).users.find_one({"_id": actor_id})
            actor_names[actor_id] = person["name"] if person else ""
    return {
        "permitted": True,
        "items": [history_item(event, actor_names.get(event.get("actor_id"), "")) for event in chosen],
    }


async def _student_directory(request: Request, user: dict, student_id: str) -> tuple[dict, list[dict], list[dict], dict]:
    student = await db(request).students.find_one({"_id": student_id, "institute_id": user["institute_id"]})
    if student is None:
        raise AppError(404, "not_found", "That record was not found.")
    owned = await db(request).enrollments.find({"student_id": student_id}).to_list(length=100)
    facts = await directory_facts(db(request), user["institute_id"])
    visible = [
        row for row in owned
        if enrollment_visible(user["grants"], row, facts["subject_by_id"], facts["paper_subject"], facts["rosters"])
    ]
    if not visible:
        raise AppError(404, "not_found", "That record was not found.")
    return student, visible, owned, facts


async def _reject_shared_name(request: Request, institute_id: str, student_id: str, key: str, owned: list[dict]) -> None:
    others = await db(request).students.find(
        {"institute_id": institute_id, "name_key": key, "_id": {"$ne": student_id}, "archived": {"$ne": True}}
    ).to_list(length=20)
    if not others:
        return
    batches = {row.get("batch_id") for row in owned if row.get("batch_id")}
    other_ids = [item["_id"] for item in others]
    overlap = await db(request).enrollments.find(
        {"student_id": {"$in": other_ids}, "batch_id": {"$in": list(batches)}}
    ).to_list(length=20)
    if overlap:
        raise AppError(
            409,
            "student.identity",
            "Another student in the same batch already uses that name. Merge them only from the merge action.",
        )


async def _plan_totals(database, institute_id: str, plan: dict, known: dict) -> dict:
    seen = {key: list(value) for key, value in known.items()}
    create = 0
    reuse = 0
    for sheet in plan["sheets"]:
        if not sheet["included"]:
            continue
        for group in sheet["groups"]:
            if group["interpretation"] == "skip_duplicate":
                continue
            for student in group["students"]:
                if not str(student.get("display_name") or "").strip():
                    continue
                sessions = seen.get(student["name_key"], [])
                if student.get("batch_session") in sessions:
                    reuse += 1
                else:
                    create += 1
                    if student.get("batch_session"):
                        seen.setdefault(student["name_key"], []).append(student["batch_session"])
    branch = await database.branches.find_one({"institute_id": institute_id})
    return {"students_to_create": create, "students_to_reuse": reuse, "catalog_started": branch is not None}


async def _merge_pair(request: Request, institute_id: str, body: MergeBody) -> tuple[dict, dict]:
    source = await db(request).students.find_one({"_id": body.source_id, "institute_id": institute_id})
    target = await db(request).students.find_one({"_id": body.target_id, "institute_id": institute_id})
    if source is None or target is None or source["_id"] == target["_id"]:
        raise AppError(404, "not_found", "That record was not found.")
    return source, target


def _catalog_row(item: dict) -> dict:
    row = {
        "id": item["_id"],
        "name": item.get("name") or item.get("display_name"),
        "archived": bool(item.get("archived")),
        "student_code": item.get("student_code"),
    }
    for key in ("branch_id", "course_id", "subject_id", "offering_id", "number", "started_on", "ended_on", "timings", "notes"):
        if item.get(key) is not None:
            row[key] = item[key]
    if item.get("units"):
        row["units"] = [
            {"id": unit.get("id"), "name": unit.get("name") or "", "kind": unit.get("kind") or "chapter"}
            for unit in item["units"]
            if unit.get("id")
        ]
    return row


async def _job(request: Request, user: dict, import_id: str) -> dict:
    job = await db(request).import_jobs.find_one({"_id": import_id, "institute_id": user["institute_id"]})
    if job is None:
        raise AppError(404, "not_found", "That record was not found.")
    if not any("marksheet.upload" in grant["actions"] or "marksheet.view" in grant["actions"] for grant in user["grants"]):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")
    return job


async def _editable(request: Request, marksheet_id: str, action: str) -> tuple[dict, dict]:
    user = await current_user(request)
    sheet = await db(request).marksheets.find_one({"_id": marksheet_id, "institute_id": user["institute_id"]})
    if sheet is None:
        raise AppError(404, "not_found", "That record was not found.")
    resource = {
        "institute_id": sheet["institute_id"],
        "branch_id": sheet.get("branch_id"),
        "course_id": sheet.get("course_id"),
        "subject_id": sheet.get("subject_id"),
    }
    if not allows(user["grants"], "marksheet.view", resource):
        raise AppError(404, "not_found", "That record was not found.")
    if not allows(user["grants"], action, resource):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")
    return user, sheet


def _stored_rows(request: Request, job: dict) -> dict:
    file_id = job.get("file_id")
    if not file_id:
        return {}
    path = request.app.state.settings.data_dir / "files" / f"{file_id}.xlsx"
    if not path.exists():
        return {}
    try:
        return workbook_rows(path.read_bytes())
    except ValueError:
        return {}


def _public_sheet(sheet: dict, rows: list | None = None) -> dict:
    return {
        "id": sheet["id"],
        "name": sheet["name"],
        "included": sheet["included"],
        "title": sheet["title"],
        "suggested_title": sheet["suggested_title"],
        "branch_name": sheet["branch_name"],
        "course_name": sheet["course_name"],
        "subject_name": sheet["subject_name"],
        "paper_number": sheet["paper_number"],
        "batches": sheet["batches"],
        "heading_conflict": sheet["heading_conflict"],
        "heading_acknowledged": sheet["heading_acknowledged"],
        "exam_date": sheet["exam_date"],
        "rows": rows or [],
        "attempt": sheet["attempt"],
        "exam_type": sheet["exam_type"],
        "planned": sheet.get("planned"),
        "groups": [
            {
                "id": group["id"],
                "maximum": group["maximum"],
                "interpretation": group["interpretation"],
                "confirmed": group["confirmed"],
                "duplicate_of": group["duplicate_of"],
                "students": group["students"],
            }
            for group in sheet["groups"]
        ],
        "blockers": sheet["blockers"],
    }


def _marksheet_summary(sheet: dict) -> dict:
    return {
        "id": sheet["_id"],
        "title": sheet.get("title") or sheet.get("status"),
        "status": sheet.get("status"),
        "branch_id": sheet.get("branch_id"),
        "course_id": sheet.get("course_id"),
        "subject_id": sheet.get("subject_id"),
        "paper_id": sheet.get("paper_id"),
        "batch_ids": sheet.get("batch_ids") or [],
        "exam_date": sheet.get("exam_date"),
        "exam_type": sheet.get("exam_type"),
        "maximum": sheet.get("maximum"),
        "source_file": sheet.get("source_file"),
        "source_sheet": sheet.get("source_sheet"),
        "file_id": sheet.get("file_id"),
        "revision": sheet.get("active_revision") or sheet.get("revision"),
        "edit_version": sheet.get("edit_version", 0),
    }

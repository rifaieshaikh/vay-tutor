from __future__ import annotations

import re
from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook

from app.security import name_key

_MONTHS = {
    "january": 1, "jan": 1,
    "february": 2, "feb": 2,
    "march": 3, "mar": 3,
    "april": 4, "apr": 4,
    "may": 5,
    "june": 6, "jun": 6,
    "july": 7, "jul": 7,
    "august": 8, "aug": 8,
    "september": 9, "sept": 9, "sep": 9,
    "october": 10, "oct": 10,
    "november": 11, "nov": 11,
    "december": 12, "dec": 12,
}
_MONTH_LABEL = {
    1: "January", 2: "February", 3: "March", 4: "April",
    5: "May", 6: "June", 7: "July", 8: "August",
    9: "September", 10: "October", 11: "November", 12: "December",
}
_BATCH_RE = re.compile(
    r"\b(january|jan|february|feb|march|mar|april|apr|may|june|jun|july|jul|august|aug|september|sept|sep|october|oct|november|nov|december|dec)\b[\s.\-/]*(\d{2,4})",
    re.IGNORECASE,
)
_MARK_RE = re.compile(r"MARK\s*\(\s*(\d+)\s*\)", re.IGNORECASE)
_PAPER_RE = re.compile(r"PAPER\s*-?\s*(\d+)\s*-\s*(.+)", re.IGNORECASE)
_LEGEND = ("more than 60", "between 40", "less than 40", "safe zone", "fifty-fifty", "danger zone", "learn, grow")


def parse_workbook(data: bytes, filename: str) -> dict:
    try:
        workbook = load_workbook(BytesIO(data), data_only=False, read_only=False)
    except Exception as exc:
        raise ValueError(f"The workbook could not be read: {exc}") from exc
    sheets = []
    for position, sheet_name in enumerate(workbook.sheetnames):
        sheets.append(_parse_sheet(workbook[sheet_name], position))
    workbook.close()
    _mark_duplicate_groups(sheets)
    return {"filename": filename, "sheets": sheets}


def workbook_rows(data: bytes) -> dict[str, list[dict]]:
    """Every non-empty row, with the role the opening lines play."""
    try:
        workbook = load_workbook(BytesIO(data), data_only=False, read_only=False)
    except Exception as exc:
        raise ValueError(f"The workbook could not be read: {exc}") from exc
    found = {}
    for sheet_name in workbook.sheetnames:
        found[sheet_name] = _source_rows(_grid(workbook[sheet_name]))
    workbook.close()
    return found


def parse_batches(text: str) -> list[dict]:
    found = []
    seen = set()
    cleaned = re.sub(r"\(\s*\d+\s*\)", " ", text or "")
    for match in _BATCH_RE.finditer(cleaned):
        month = _MONTHS[match.group(1).lower()]
        year = int(match.group(2))
        if year < 100:
            year += 2000
        key = f"{year}-{month:02d}"
        if key in seen:
            continue
        seen.add(key)
        found.append({"session_key": key, "label": f"{_MONTH_LABEL[month]} {year}"})
    return found


def session_key_for_label(text: str) -> str | None:
    found = parse_batches(text)
    if len(found) == 1:
        return found[0]["session_key"]
    return None


def _parse_sheet(worksheet, position: int) -> dict:
    rows = _grid(worksheet)
    header_index = _header_row(rows)
    labels = []
    for row in rows[: header_index or 0]:
        text = _first_text(row)
        if text and text.upper() not in {"SL.NO", "NAME"}:
            labels.append(text)
    branch = next((item for item in labels if "IAM" in item.upper()), "")
    course_line = next((item for item in labels if "FOUNDATION" in item.upper() or "CA " in item.upper()), "")
    paper_line = next((item for item in labels if "PAPER" in item.upper()), "")
    series = next((item for item in labels if "SERIES" in item.upper()), "")
    title = labels[-1] if labels else worksheet.title
    if title in {branch, course_line, paper_line, series}:
        title = worksheet.title
    paper_number, subject = _paper(paper_line)
    heading_batches = parse_batches(course_line)
    sheet_batches = parse_batches(worksheet.title)
    conflict = bool(sheet_batches) and {item["session_key"] for item in sheet_batches} != {
        item["session_key"] for item in heading_batches
    }
    batches = sheet_batches if conflict else heading_batches
    groups = _score_groups(rows, header_index)
    exam_date = _exam_date(worksheet)
    return {
        "id": f"sheet-{position}",
        "position": position,
        "name": worksheet.title,
        "branch_name": _display_branch(branch),
        "course_name": "CA Foundation" if course_line else "",
        "subject_name": subject,
        "paper_number": paper_number,
        "series": "Test series" if series else "",
        "title": title,
        "title_key": _title_key(title),
        "suggested_title": title.replace("EAXAM", "EXAM"),
        "exam_type": _exam_type(title),
        "format": "mcq" if "mcq" in title.casefold() else "descriptive",
        "attempt": "retest" if re.search(r"re\s*test", title, re.IGNORECASE) else "original",
        "heading_batches": heading_batches,
        "recommended_batches": batches,
        "heading_conflict": conflict,
        "exam_date": exam_date,
        "groups": groups,
        "included": True,
    }


def _source_rows(rows: list[list]) -> list[dict]:
    header_index = _header_row(rows)
    found = []
    for index, row in enumerate(rows):
        cells = [_cell_text(cell) for cell in row]
        while cells and not cells[-1]:
            cells.pop()
        if not any(cells):
            continue
        found.append({"row": index + 1, "cells": cells, "role": _row_role(row, cells, index, header_index)})
    return found


def _row_role(row: list, cells: list[str], index: int, header_index: int | None) -> str:
    if header_index is not None and index == header_index:
        return "header"
    if header_index is not None and index > header_index:
        name = row[1] if len(row) > 1 else None
        if isinstance(name, str) and name.strip() and not _is_legend(name):
            return "student"
        return "legend"
    text = _first_text(row).upper()
    if "IAM" in text:
        return "branch"
    if "FOUNDATION" in text or text.startswith("CA "):
        return "course"
    if "PAPER" in text:
        return "paper"
    if "SERIES" in text:
        return "series"
    if text:
        return "title"
    return "note"


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _grid(worksheet) -> list[list]:
    grid = []
    for row in worksheet.iter_rows(max_row=worksheet.max_row or 1, max_col=max(worksheet.max_column or 1, 1)):
        grid.append([cell.value for cell in row])
    return grid


def _header_row(rows: list[list]) -> int | None:
    for index, row in enumerate(rows):
        texts = [str(cell).strip().upper() for cell in row if isinstance(cell, str)]
        if "NAME" in texts and any(_MARK_RE.search(text) for text in texts):
            return index
    return None


def _first_text(row: list) -> str:
    for cell in row[:1]:
        if isinstance(cell, str) and cell.strip():
            return cell.strip()
    return ""


def _paper(line: str) -> tuple[int | None, str]:
    match = _PAPER_RE.search(line or "")
    if not match:
        return None, ""
    subject = match.group(2).strip(" -")
    names = {"ACCOUNTING": "Accounting", "BUSINESS LAW": "Business Law", "QT": "QT", "ECONOMICS": "Economics"}
    return int(match.group(1)), names.get(subject.upper(), subject.title())


def _display_branch(text: str) -> str:
    if "TIRUR" in text.upper():
        return "IAM Tirur"
    return " ".join(text.split()).title()


def _title_key(title: str) -> str:
    cleaned = re.sub(r"\(\s*re\s*test\s*\)", " ", title, flags=re.IGNORECASE)
    cleaned = cleaned.replace("EAXAM", "EXAM")
    return name_key(re.sub(r"[^a-z0-9]+", " ", cleaned, flags=re.IGNORECASE))


def _exam_type(title: str) -> str:
    lowered = title.casefold()
    if "part exam" in lowered or lowered.startswith("part"):
        return "part"
    if "chapter" in lowered:
        return "chapter"
    return "unit"


def _exam_date(worksheet) -> str | None:
    for row in worksheet.iter_rows(max_row=min(worksheet.max_row or 1, 12), max_col=worksheet.max_column or 1):
        for cell in row:
            if isinstance(cell.value, datetime):
                return cell.value.date().isoformat()
            if isinstance(cell.value, date):
                return cell.value.isoformat()
    return None


def _score_groups(rows: list[list], header_index: int | None) -> list[dict]:
    if header_index is None:
        return []
    header = rows[header_index]
    groups = []
    for column, cell in enumerate(header):
        if not isinstance(cell, str):
            continue
        match = _MARK_RE.search(cell)
        if not match:
            continue
        maximum = int(match.group(1))
        rank_column = column + 2 if column + 2 < len(header) else None
        students = []
        for offset, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
            name = row[1] if len(row) > 1 else None
            if not isinstance(name, str) or not name.strip():
                continue
            if _is_legend(name):
                continue
            raw = row[column] if column < len(row) else None
            source_rank = row[rank_column] if rank_column is not None and rank_column < len(row) else None
            students.append(_student(name, raw, source_rank, offset, column + 1, maximum))
        groups.append(
            {
                "id": f"group-{column}",
                "column": column + 1,
                "maximum": maximum,
                "interpretation": "assessment" if not groups else "separate_assessment",
                "confirmed": not groups,
                "duplicate_of": None,
                "students": students,
            }
        )
    return groups


def _is_legend(name: str) -> bool:
    lowered = name.casefold()
    return any(phrase in lowered for phrase in _LEGEND)


def _student(name: str, raw, source_rank, row_number: int, column: int, maximum: int) -> dict:
    status, score = _mark(raw, maximum)
    percentage = (score / maximum * 100) if status == "scored" else None
    return {
        "row": row_number,
        "column": column,
        "display_name": " ".join(name.split()),
        "name_key": name_key(name),
        "status": status,
        "score": score,
        "maximum": maximum,
        "percentage": percentage,
        "band": _band(percentage),
        "source_rank": source_rank if isinstance(source_rank, (int, float)) else None,
        "batch_session": None,
    }


def _mark(raw, maximum: int) -> tuple[str, float | None]:
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return "missing", None
    if isinstance(raw, str) and raw.strip().upper() in {"A", "AB"}:
        return "absent", None
    if isinstance(raw, str):
        return "invalid", None
    if isinstance(raw, (int, float)):
        if raw < 0 or raw > maximum:
            return "invalid", float(raw)
        return "scored", float(raw)
    return "invalid", None


def _band(percentage: float | None) -> str | None:
    if percentage is None:
        return None
    if percentage < 40:
        return "danger"
    if percentage <= 60:
        return "fifty-fifty"
    return "safe"


def _mark_duplicate_groups(sheets: list[dict]) -> None:
    seen: dict[tuple, dict] = {}
    for sheet in sheets:
        for group in sheet["groups"]:
            key = (sheet["title_key"], group["maximum"], tuple(item["session_key"] for item in sheet["recommended_batches"]))
            previous = seen.get(key)
            if previous is None:
                seen[key] = {"sheet": sheet["name"], "scores": _scores(group)}
                continue
            overlap = [
                name
                for name, score in _scores(group).items()
                if name in previous["scores"] and previous["scores"][name] == score
            ]
            if overlap and overlap == [name for name in _scores(group) if name in previous["scores"]]:
                group["duplicate_of"] = previous["sheet"]
                group["interpretation"] = "skip_duplicate"
                group["confirmed"] = False


def _scores(group: dict) -> dict[str, float]:
    return {
        student["name_key"]: student["score"]
        for student in group["students"]
        if student["status"] == "scored" and student["score"] is not None
    }

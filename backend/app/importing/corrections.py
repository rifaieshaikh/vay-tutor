from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook


def parse_correction_file(data: bytes) -> list[dict]:
    workbook = load_workbook(BytesIO(data), data_only=False, read_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("The correction file is empty.")
    header = [str(cell or "").strip().casefold() for cell in rows[0]]

    def column(*names: str) -> int | None:
        for name in names:
            if name in header:
                return header.index(name)
        return None

    sheet_index = column("sheet", "sheet name")
    name_index = column("student", "name")
    score_index = column("score", "correct score", "new score")
    reason_index = column("reason")
    row_index = column("row")
    if sheet_index is None or name_index is None or score_index is None:
        raise ValueError("Use columns Sheet, Student, Score, and Reason.")
    parsed = []
    for line in rows[1:]:
        if not line or all(cell in (None, "") for cell in line):
            continue
        parsed.append(
            {
                "sheet": str(line[sheet_index] or "").strip(),
                "student": str(line[name_index] or "").strip(),
                "score": line[score_index] if score_index < len(line) else None,
                "reason": str(line[reason_index] or "").strip() if reason_index is not None and reason_index < len(line) else "",
                "row": line[row_index] if row_index is not None and row_index < len(line) else None,
            }
        )
    return parsed

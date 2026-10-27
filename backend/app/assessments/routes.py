"""Workbook import and marksheet drafts."""

from __future__ import annotations

import hashlib
import json
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

from bson import ObjectId
from fastapi import APIRouter, File, Request, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from pydantic import BaseModel, Field

from app.acl import allows
from app.deps import current_user, db, write_audit
from app.errors import AppError
from app.importing.commit import commit_import, known_sessions
from app.importing.corrections import parse_correction_file
from app.importing.counts import annotate_plan
from app.importing.parse import parse_workbook, workbook_rows
from app.importing.plan import apply_decisions, variant_suggestions
from app.importing.publish import (
    correct_marksheet,
    reject_marksheet,
    require_current_version,
    submit_marksheet,
    withdraw_marksheet,
)
from app.security import name_key

router = APIRouter(prefix="/api/v1")
_MAX_BYTES = 25 * 1024 * 1024

class DecisionBody(BaseModel):
    sheets: dict = Field(default_factory=dict)
    defaults: dict = Field(default_factory=dict)


class ReasonBody(BaseModel):
    reason: str = ""
    changes: list[dict] = Field(default_factory=list)
    edit_version: int | None = None


class ResultPatch(BaseModel):
    score: float | None = None
    status: str | None = None
    edit_version: int | None = None


class BandBody(BaseModel):
    name: str = ""
    through: float | None = None
    color: str | None = None


class MarksheetPatch(BaseModel):
    exam_date: str | None = None
    bands: list[BandBody] | None = None


class MarksheetCreate(BaseModel):
    title: str = Field(min_length=1)
    branch_id: str
    course_id: str
    batch_id: str
    subject_id: str
    paper_id: str
    exam_type: str
    maximum: float = Field(gt=0)
    exam_date: str | None = None
    attempt: str = "original"


class ResultCreate(BaseModel):
    student_id: str
    status: str = "absent"
    score: float | None = None


class GridMark(BaseModel):
    enrollment_code: str
    score: str = ""


class MarksheetGrid(BaseModel):
    title: str = Field(min_length=1)
    branch_id: str
    course_id: str
    batch_id: str
    subject_id: str
    paper_id: str
    exam_type: str = "unit"
    maximum: float = Field(gt=0)
    exam_date: str | None = None
    rows: list[GridMark]


def _now() -> datetime:
    return datetime.now(timezone.utc)


@router.post("/imports")
async def upload_import(request: Request, file: UploadFile = File(...)) -> dict:
    user = await current_user(request)
    if not any("marksheet.upload" in grant["actions"] for grant in user["grants"]):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")
    filename = file.filename or "marklist.xlsx"
    if not filename.lower().endswith(".xlsx"):
        raise AppError(422, "import.file_type", "Upload an .xlsx workbook.")
    data = await file.read()
    if len(data) > _MAX_BYTES:
        raise AppError(422, "import.too_large", "The workbook is larger than 25 MB.")
    digest = hashlib.sha256(data).hexdigest()
    existing = await db(request).import_jobs.find_one(
        {
            "institute_id": user["institute_id"],
            "sha256": digest,
            "state": {"$nin": ["committed", "cancelled"]},
            "exam_id": {"$exists": False},
            "bundle": {"$ne": True},
        }
    )
    if existing:
        return {"id": existing["_id"], "state": existing["state"], "duplicate": True}
    try:
        parsed = parse_workbook(data, filename)
    except ValueError as exc:
        raise AppError(422, "import.unreadable", str(exc)) from exc
    file_id = str(ObjectId())
    path = request.app.state.settings.data_dir / "files" / f"{file_id}.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    job_id = str(ObjectId())
    await db(request).files.insert_one(
        {
            "_id": file_id,
            "institute_id": user["institute_id"],
            "kind": "source_xlsx",
            "sha256": digest,
            "path": str(path),
            "filename": filename,
            "created_at": _now(),
        }
    )
    state = "needs_resolution" if _needs_resolution(parsed) else "ready"
    await db(request).import_jobs.insert_one(
        {
            "_id": job_id,
            "institute_id": user["institute_id"],
            "actor_id": user["_id"],
            "filename": filename,
            "sha256": digest,
            "file_id": file_id,
            "state": state,
            "parsed": parsed,
            "decisions": {"sheets": {}},
            "committed_sheets": [],
            "outcomes": [],
            "created_at": _now(),
        }
    )
    await db(request).audit_events.insert_one(
        {
            "_id": str(ObjectId()),
            "institute_id": user["institute_id"],
            "action": "import.upload",
            "actor_id": user["_id"],
            "at": _now(),
            "context": {"import_id": job_id, "file": filename},
        }
    )
    return {"id": job_id, "state": state, "duplicate": False, "sheets": len(parsed["sheets"])}


def _needs_resolution(parsed: dict) -> bool:
    for sheet in parsed["sheets"]:
        if sheet["heading_conflict"]:
            return True
        if any(not group["confirmed"] for group in sheet["groups"]):
            return True
        if len(sheet["recommended_batches"]) > 1:
            return True
    return False


@router.get("/imports/{import_id}")
async def get_import(request: Request, import_id: str) -> dict:
    user = await current_user(request)
    job = await _job(request, user, import_id)
    known = await known_sessions(db(request), user["institute_id"])
    plan = apply_decisions(job["parsed"], job.get("decisions") or {}, known)
    context = None
    if job.get("bundle"):
        from app.importing.bundle_plan import annotate_bundle_plan, bundle_context

        await annotate_bundle_plan(db(request), plan, job)
        context = await bundle_context(db(request), job)
    elif job.get("exam_id"):
        from app.importing.exam_plan import annotate_exam_plan

        exam = await db(request).exams.find_one({"_id": job["exam_id"], "institute_id": user["institute_id"]})
        if exam:
            await annotate_exam_plan(db(request), plan, exam)
        else:
            plan["totals"].update(await _plan_totals(db(request), user["institute_id"], plan, known))
            await annotate_plan(db(request), user["institute_id"], plan)
    else:
        plan["totals"].update(await _plan_totals(db(request), user["institute_id"], plan, known))
        await annotate_plan(db(request), user["institute_id"], plan)
    stored_rows = _stored_rows(request, job)
    return {
        "id": job["_id"],
        "state": job["state"],
        "filename": job["filename"],
        "sheets": [_public_sheet(sheet, stored_rows.get(sheet["name"], [])) for sheet in plan["sheets"]],
        "blockers": plan["blockers"],
        "totals": plan["totals"],
        "variants": variant_suggestions(job["parsed"]),
        "outcomes": job.get("outcomes") or [],
        "summary": job.get("summary"),
        "ready": plan["ready"],
        "exam_id": job.get("exam_id") or "",
        "bundle": bool(job.get("bundle")),
        "context": context,
    }


@router.patch("/imports/{import_id}")
async def patch_import(request: Request, import_id: str, body: DecisionBody) -> dict:
    user = await current_user(request)
    job = await _job(request, user, import_id)
    if job["state"] == "committed":
        raise AppError(422, "import.committed", "This workbook is already committed.")
    decisions = job.get("decisions") or {"sheets": {}}
    if body.defaults:
        decisions["defaults"] = {**(decisions.get("defaults") or {}), **body.defaults}
    for sheet_id, choice in body.sheets.items():
        current = decisions["sheets"].get(sheet_id, {})
        groups = {**(current.get("groups") or {}), **(choice.get("groups") or {})}
        levels = {**(current.get("levels") or {}), **(choice.get("levels") or {})}
        row_batches = {**(current.get("row_batches") or {}), **(choice.get("row_batches") or {})}
        corrections = {**(current.get("corrections") or {}), **(choice.get("corrections") or {})}
        names = {**(current.get("names") or {}), **(choice.get("names") or {})}
        added = {**(current.get("added") or {}), **(choice.get("added") or {})}
        row_batch_ids = {**(current.get("row_batch_ids") or {}), **(choice.get("row_batch_ids") or {})}
        roster_marks = {**(current.get("roster_marks") or {}), **(choice.get("roster_marks") or {})}
        drop_rows = list(dict.fromkeys([*(current.get("drop_rows") or []), *(choice.get("drop_rows") or [])]))
        previous_exam = current.get("exam") or {}
        current.update(choice)
        if isinstance(choice.get("exam"), dict):
            current["exam"] = {**previous_exam, **choice["exam"]}
            if isinstance(choice["exam"].get("scope"), dict):
                current["exam"]["scope"] = choice["exam"]["scope"]
        if groups:
            current["groups"] = groups
        if levels:
            current["levels"] = levels
        if row_batches:
            current["row_batches"] = row_batches
        if row_batch_ids:
            current["row_batch_ids"] = row_batch_ids
        if "roster_marks" in choice:
            current["roster_marks"] = roster_marks
        if corrections:
            current["corrections"] = corrections
        if names:
            current["names"] = names
        if added:
            current["added"] = added
        if drop_rows:
            current["drop_rows"] = drop_rows
        decisions["sheets"][sheet_id] = current
    known = await known_sessions(db(request), user["institute_id"])
    plan = apply_decisions(job["parsed"], decisions, known)
    job["decisions"] = decisions
    if job.get("bundle"):
        from app.importing.bundle_plan import annotate_bundle_plan

        await annotate_bundle_plan(db(request), plan, job)
    state = "ready" if plan["ready"] else "needs_resolution"
    await db(request).import_jobs.update_one({"_id": job["_id"]}, {"$set": {"decisions": decisions, "state": state}})
    return {"id": job["_id"], "state": state, "ready": plan["ready"], "blockers": plan["blockers"]}


@router.post("/imports/{import_id}/corrections")
async def upload_corrections(request: Request, import_id: str, file: UploadFile = File(...)) -> dict:
    user = await current_user(request)
    job = await _job(request, user, import_id)
    if job["state"] == "committed":
        raise AppError(422, "import.committed", "This workbook is already committed.")
    filename = file.filename or "corrections.xlsx"
    if not filename.lower().endswith(".xlsx"):
        raise AppError(422, "import.file_type", "Upload an .xlsx correction file.")
    try:
        rows = parse_correction_file(await file.read())
    except ValueError as exc:
        raise AppError(422, "import.unreadable", str(exc)) from exc
    known = await known_sessions(db(request), user["institute_id"])
    plan = apply_decisions(job["parsed"], job.get("decisions") or {}, known)
    decisions = job.get("decisions") or {"sheets": {}}
    decisions.setdefault("sheets", {})
    applied = []
    unmatched = []
    for row in rows:
        sheet = next(
            (item for item in plan["sheets"] if item["name"].strip().casefold() == row["sheet"].casefold()), None
        )
        if sheet is None:
            unmatched.append(
                {"sheet": row["sheet"], "student": row["student"], "message": "That sheet is not in this workbook."}
            )
            continue
        wanted_row = row["row"]
        matches = []
        for group in sheet["groups"]:
            for student in group["students"]:
                if wanted_row not in (None, "") and int(wanted_row) != int(student["row"]):
                    continue
                if name_key(row["student"]) == student["name_key"]:
                    matches.append(student)
        if len(matches) != 1:
            unmatched.append(
                {
                    "sheet": row["sheet"],
                    "student": row["student"],
                    "message": "Match one student. Add the Row column when a name appears twice.",
                }
            )
            continue
        score = row["score"]
        status = None
        if isinstance(score, str) and score.strip().upper() in {"A", "AB"}:
            status = "absent"
            score = None
        elif score in (None, ""):
            status = "absent"
            score = None
        choice = decisions["sheets"].setdefault(sheet["id"], {})
        corrections = choice.setdefault("corrections", {})
        corrections[str(matches[0]["row"])] = {
            "action": "replace",
            "score": score,
            "status": status,
            "reason": row["reason"],
        }
        applied.append({"sheet": sheet["name"], "student": matches[0]["display_name"], "row": matches[0]["row"]})
    await db(request).import_jobs.update_one({"_id": job["_id"]}, {"$set": {"decisions": decisions}})
    return {"applied": applied, "unmatched": unmatched}


@router.get("/imports/{import_id}/correction-template")
async def correction_template(request: Request, import_id: str):
    payload = await get_import(request, import_id)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Corrections"
    sheet.append(["Sheet", "Row", "Student", "File score", "Saved score", "Score", "Reason"])
    for item in payload["sheets"]:
        for group in item["groups"]:
            for student in group["students"]:
                if student.get("change") not in {"reject", "update"}:
                    continue
                sheet.append(
                    [
                        item["name"],
                        student["row"],
                        student["display_name"],
                        student.get("score"),
                        student.get("before_score"),
                        "",
                        "",
                    ]
                )
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="corrections.xlsx"'},
    )


@router.get("/imports/{import_id}/preview")
async def preview_import(request: Request, import_id: str) -> dict:
    return await get_import(request, import_id)


@router.get("/imports/{import_id}/validation-report")
async def validation_report(request: Request, import_id: str):
    user = await current_user(request)
    if not any("export.xlsx" in grant["actions"] for grant in user["grants"]):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")
    payload = await get_import(request, import_id)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Validation"
    sheet.append(
        [
            "Sheet",
            "Row",
            "Name",
            "Status",
            "Score",
            "Before",
            "Change",
            "Maximum",
            "Percentage",
            "Rank",
            "Source rank",
            "Blocker",
        ]
    )
    blockers = {(item.get("sheet_id"), item.get("row")): item["message"] for item in payload["blockers"]}
    for item in payload["sheets"]:
        for group in item["groups"]:
            for student in group["students"]:
                sheet.append(
                    [
                        item["name"],
                        student["row"],
                        student["display_name"],
                        student["status"],
                        student["score"],
                        student.get("before_score"),
                        student.get("change"),
                        group["maximum"],
                        None if student["percentage"] is None else round(student["percentage"], 2),
                        student["rank"],
                        student["source_rank"],
                        blockers.get((item["id"], student["row"]), ""),
                    ]
                )
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="validation.xlsx"'},
    )


@router.get("/imports/{import_id}/outcome-report")
async def outcome_report(request: Request, import_id: str):
    user = await current_user(request)
    if not any("export.xlsx" in grant["actions"] for grant in user["grants"]):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")
    job = await _job(request, user, import_id)
    summary = job.get("summary")
    if not summary:
        raise AppError(422, "import.not_committed", "Commit the workbook before downloading the outcome.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Outcome"
    sheet.append(["Kind", "Inserted", "Updated", "Unchanged", "Skipped", "Rejected"])
    for kind, buckets in (summary.get("counts") or {}).items():
        sheet.append(
            [
                kind,
                buckets.get("insert", 0),
                buckets.get("update", 0),
                buckets.get("unchanged", 0),
                buckets.get("skip", 0),
                buckets.get("reject", 0),
            ]
        )
    detail = workbook.create_sheet("Sheets")
    detail.append(["Sheet", "Status", "Kind", "Inserted", "Updated", "Unchanged", "Skipped", "Rejected"])
    for item in summary.get("by_sheet") or []:
        for kind, buckets in (item.get("counts") or {}).items():
            detail.append(
                [
                    item.get("name"),
                    item.get("status"),
                    kind,
                    buckets.get("insert", 0),
                    buckets.get("update", 0),
                    buckets.get("unchanged", 0),
                    buckets.get("skip", 0),
                    buckets.get("reject", 0),
                ]
            )
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="import-outcome.xlsx"'},
    )


@router.post("/imports/{import_id}/commit")
async def commit_import_route(request: Request, import_id: str) -> dict:
    user = await current_user(request)
    job = await _job(request, user, import_id)
    if not any("marksheet.upload" in grant["actions"] for grant in user["grants"]):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")
    if job["state"] == "committed":
        return {"id": job["_id"], "state": "committed", "outcomes": job.get("outcomes") or []}
    if job["state"] == "cancelled":
        raise AppError(422, "import.cancelled", "This import was cancelled.")
    raw = await request.body()
    confirm_updates = bool(json.loads(raw).get("confirm_updates")) if raw else False
    result = await commit_import(db(request), request.app.state.client, job, user, confirm_updates=confirm_updates)
    return {"id": job["_id"], **result}


@router.post("/imports/{import_id}/cancel")
async def cancel_import(request: Request, import_id: str) -> dict:
    user = await current_user(request)
    job = await _job(request, user, import_id)
    if job["state"] == "committed":
        raise AppError(422, "import.committed", "A committed workbook cannot be cancelled.")
    await db(request).import_jobs.update_one({"_id": job["_id"]}, {"$set": {"state": "cancelled"}})
    return {"id": job["_id"], "state": "cancelled"}


@router.delete("/imports/{import_id}")
async def delete_import(request: Request, import_id: str) -> dict:
    user = await current_user(request)
    job = await _job(request, user, import_id)
    if not any("marksheet.upload" in grant["actions"] for grant in user["grants"]):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")
    if job["state"] == "committed":
        raise AppError(422, "import.committed", "A saved import cannot be deleted.")
    await db(request).import_jobs.delete_one({"_id": job["_id"], "institute_id": user["institute_id"]})
    if job.get("file_id"):
        stored = await db(request).files.find_one({"_id": job["file_id"], "institute_id": user["institute_id"]})
        if stored:
            _remove_stored_file(request, stored.get("path"))
            await db(request).files.delete_one({"_id": stored["_id"]})
    await write_audit(
        request,
        user,
        "import.delete",
        {"filename": job.get("filename") or ""},
        None,
        {"institute_id": user["institute_id"]},
    )
    return {"id": job["_id"]}


def _remove_stored_file(request: Request, raw: str | None) -> None:
    if not raw:
        return
    root = request.app.state.settings.data_dir.resolve()
    path = Path(raw)
    try:
        resolved = path.resolve()
    except OSError:
        return
    if resolved.is_relative_to(root) and resolved.is_file():
        resolved.unlink()


def _import_allowed(user: dict, batch: dict) -> bool:
    resource = {
        "institute_id": user["institute_id"],
        "branch_id": batch.get("branch_id"),
        "course_id": batch.get("course_id"),
        "batch_id": batch["_id"],
    }
    return any(allows(user["grants"], action, resource, True) for action in ("marksheet.view", "marksheet.upload"))


@router.get("/imports")
async def list_imports(
    request: Request, state: str = "", branch_id: str = "", course_id: str = "", batch_id: str = "", exam_id: str = ""
) -> dict:
    user = await current_user(request)
    if not any(
        "marksheet.upload" in grant["actions"] or "marksheet.view" in grant["actions"] for grant in user["grants"]
    ):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")
    query: dict = {"institute_id": user["institute_id"]}
    if state:
        query["state"] = state
    if exam_id:
        query["exam_id"] = exam_id
    narrowed = bool(branch_id or course_id or batch_id or state == "failed")
    allowed = await _import_batch_ids(request, user, branch_id, course_id, batch_id) if narrowed else set()
    items = []
    cursor = db(request).import_jobs.find(query).sort("created_at", -1)
    async for job in cursor:
        if narrowed and not set(job.get("batch_ids") or []).intersection(allowed):
            continue
        items.append(
            {
                "id": job["_id"],
                "filename": job.get("filename"),
                "state": job.get("state"),
                "created_at": job.get("created_at").isoformat() if job.get("created_at") else None,
                "batch_ids": list(job.get("batch_ids") or []),
                "outcomes": job.get("outcomes") or [],
                "summary": job.get("summary"),
                "marksheet_id": job.get("marksheet_id") or "",
                "source": job.get("source") or "workbook",
                "bundle": bool(job.get("bundle")),
            }
        )
    return {"items": items, "total": len(items)}


async def _import_batch_ids(request: Request, user: dict, branch_id: str, course_id: str, batch_id: str) -> set[str]:
    rows = await db(request).batches.find({"institute_id": user["institute_id"]}).to_list(length=500)
    allowed = set()
    for batch in rows:
        if branch_id and batch.get("branch_id") != branch_id:
            continue
        if course_id and batch.get("course_id") != course_id:
            continue
        if batch_id and batch["_id"] != batch_id:
            continue
        if _import_allowed(user, batch):
            allowed.add(batch["_id"])
    return allowed


@router.post("/marksheets")
async def create_marksheet(request: Request, body: MarksheetCreate) -> dict:
    user = await current_user(request)
    if body.exam_type not in {"unit", "part", "chapter"}:
        raise AppError(422, "marksheet.exam_type", "Choose unit, part, or chapter.")
    if body.attempt not in {"original", "retest"}:
        raise AppError(422, "marksheet.attempt", "Choose an original attempt or a retest.")
    database = db(request)
    institute_id = user["institute_id"]
    branch = await database.branches.find_one({"_id": body.branch_id, "institute_id": institute_id})
    course = await database.courses.find_one({"_id": body.course_id, "institute_id": institute_id})
    batch = await database.batches.find_one({"_id": body.batch_id, "institute_id": institute_id})
    subject = await database.subjects.find_one({"_id": body.subject_id, "institute_id": institute_id})
    paper = await database.papers.find_one({"_id": body.paper_id, "institute_id": institute_id})
    if not all((branch, course, batch, subject, paper)):
        raise AppError(404, "not_found", "That academic record was not found.")
    if (
        batch.get("course_id") != course["_id"]
        or subject.get("course_id") != course["_id"]
        or paper.get("subject_id") != subject["_id"]
    ):
        raise AppError(422, "marksheet.context", "That batch, subject, and paper do not belong to the same course.")
    if batch.get("branch_id") and batch.get("branch_id") != branch["_id"]:
        raise AppError(422, "marksheet.context", "That batch does not belong to the chosen branch.")
    offering = await database.offerings.find_one(
        {"institute_id": institute_id, "branch_id": branch["_id"], "course_id": course["_id"]}
    )
    if offering is None:
        raise AppError(422, "marksheet.context", "That branch does not offer the chosen course.")
    resource = {
        "institute_id": institute_id,
        "branch_id": branch["_id"],
        "course_id": course["_id"],
        "batch_id": batch["_id"],
        "subject_id": subject["_id"],
        "paper_id": paper["_id"],
    }
    if not allows(user["grants"], "marksheet.upload", resource):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")
    title = body.title.strip()
    identity = "|".join(("manual", name_key(title), body.exam_type, body.attempt, batch["_id"], paper["_id"]))
    existing = await database.assessments.find_one({"institute_id": institute_id, "identity_key": identity})
    if existing:
        found = await database.marksheets.find_one({"assessment_id": existing["_id"], "institute_id": institute_id})
        raise AppError(
            409,
            "marksheet.exists",
            "A marksheet with that title already exists for this batch.",
            id=found["_id"] if found else None,
        )
    sheet_id = str(ObjectId())
    assessment_id = str(ObjectId())
    attempt_id = str(ObjectId())
    client = request.app.state.client
    async with await client.start_session() as session:
        async with session.start_transaction():
            await database.assessments.insert_one(
                {
                    "_id": assessment_id,
                    "institute_id": institute_id,
                    "identity_key": identity,
                    "title": title,
                    "title_key": name_key(title),
                    "exam_type": body.exam_type,
                    "format": "manual",
                    "subject_id": subject["_id"],
                    "paper_id": paper["_id"],
                    "branch_id": branch["_id"],
                    "course_id": course["_id"],
                    "batch_ids": [batch["_id"]],
                    "exam_date": body.exam_date,
                    "maximum": body.maximum,
                },
                session=session,
            )
            await database.attempts.insert_one(
                {
                    "_id": attempt_id,
                    "institute_id": institute_id,
                    "assessment_id": assessment_id,
                    "kind": body.attempt,
                    "maximum": body.maximum,
                },
                session=session,
            )
            await database.marksheets.insert_one(
                {
                    "_id": sheet_id,
                    "institute_id": institute_id,
                    "assessment_id": assessment_id,
                    "attempt_id": attempt_id,
                    "cohort_key": batch["_id"],
                    "branch_id": branch["_id"],
                    "course_id": course["_id"],
                    "subject_id": subject["_id"],
                    "paper_id": paper["_id"],
                    "batch_ids": [batch["_id"]],
                    "title": title,
                    "maximum": body.maximum,
                    "exam_type": body.exam_type,
                    "exam_date": body.exam_date,
                    "status": "draft",
                    "revision": 1,
                    "active_revision": None,
                    "source_sheet": "Entered here",
                    "uploader_id": user["_id"],
                    "edit_version": 0,
                    "attempt_kind": body.attempt,
                },
                session=session,
            )
    await database.audit_events.insert_one(
        {
            "_id": str(ObjectId()),
            "institute_id": institute_id,
            "action": "marksheet.create",
            "actor_id": user["_id"],
            "at": _now(),
            "context": {"marksheet_id": sheet_id, "title": title},
        }
    )
    return {"id": sheet_id, "status": "draft"}


@router.post("/marksheets/grid")
async def save_marksheet_grid(request: Request, body: MarksheetGrid) -> dict:
    """Save a pasted roster. Each row is matched by enrollment number, not by name."""
    filled = [row for row in body.rows if row.score.strip()]
    if not filled:
        raise AppError(422, "marksheet.grid", "Enter a score, or type absent, for at least one student.")
    database = db(request)
    codes = [row.enrollment_code.strip() for row in filled]
    if any(not code for code in codes):
        raise AppError(422, "marksheet.grid", "Each pasted row needs an enrollment number.")
    if len(codes) != len(set(codes)):
        raise AppError(422, "marksheet.grid", "An enrollment number is repeated in the pasted rows.")
    enrollments = await database.enrollments.find(
        {
            "institute_id": (await current_user(request))["institute_id"],
            "batch_id": body.batch_id,
            "enrollment_code": {"$in": codes},
        }
    ).to_list(length=len(codes))
    by_code = {item.get("enrollment_code"): item for item in enrollments}
    missing = [code for code in codes if code not in by_code]
    if missing:
        shown = ", ".join(missing[:8])
        raise AppError(422, "marksheet.grid", f"These enrollment numbers are not in this batch: {shown}.")
    prepared = []
    for row in filled:
        raw = row.score.strip()
        enrollment = by_code[row.enrollment_code.strip()]
        if raw.casefold() in {"a", "absent"}:
            prepared.append((enrollment["student_id"], "absent", None))
            continue
        try:
            score = float(raw)
        except ValueError as exc:
            raise AppError(
                422, "import.invalid_mark", f"{row.enrollment_code.strip()} needs a score or absent."
            ) from exc
        if score < 0 or score > body.maximum:
            raise AppError(422, "import.invalid_mark", f"{row.enrollment_code.strip()} is outside 0 through the maximum.")
        prepared.append((enrollment["student_id"], "scored", score))
    created = await create_marksheet(
        request,
        MarksheetCreate(
            title=body.title,
            branch_id=body.branch_id,
            course_id=body.course_id,
            batch_id=body.batch_id,
            subject_id=body.subject_id,
            paper_id=body.paper_id,
            exam_type=body.exam_type,
            maximum=body.maximum,
            exam_date=body.exam_date,
        ),
    )
    for student_id, status, score in prepared:
        await add_result(request, created["id"], ResultCreate(student_id=student_id, status=status, score=score))
    user = await current_user(request)
    await database.import_jobs.insert_one(
        {
            "_id": str(ObjectId()),
            "institute_id": user["institute_id"],
            "actor_id": user["_id"],
            "filename": body.title.strip(),
            "source": "grid",
            "marksheet_id": created["id"],
            "state": "committed",
            "batch_ids": [body.batch_id],
            "outcomes": [{"name": body.title.strip(), "status": "committed"}],
            "summary": {"counts": {"results": {"insert": len(prepared), "update": 0}}},
            "created_at": _now(),
        }
    )
    return {"id": created["id"], "saved": len(prepared)}


@router.post("/marksheets/{marksheet_id}/results")
async def add_result(request: Request, marksheet_id: str, body: ResultCreate) -> dict:
    user, sheet = await _editable(request, marksheet_id, "marksheet.edit_draft")
    if sheet.get("status") not in {"draft", "submitted"}:
        raise AppError(422, "marksheet.state", "Add students on a draft. A published sheet uses a correction.")
    if body.status not in {"scored", "absent", "missing"}:
        raise AppError(422, "import.invalid_mark", "Choose scored or absent.")
    score = body.score
    status = "absent" if body.status == "missing" else body.status
    if status == "scored":
        if score is None or score < 0 or (sheet.get("maximum") is not None and score > sheet["maximum"]):
            raise AppError(422, "import.invalid_mark", "The score must be from 0 through the maximum.")
    else:
        score = None
    enrollment = await db(request).enrollments.find_one(
        {"student_id": body.student_id, "batch_id": {"$in": sheet.get("batch_ids") or []}}
    )
    if enrollment is None:
        raise AppError(422, "marksheet.not_enrolled", "That student is not enrolled in this batch.")
    revision = sheet.get("revision") or 1
    duplicate = await db(request).results.find_one(
        {"marksheet_id": sheet["_id"], "revision": revision, "student_id": body.student_id}
    )
    if duplicate:
        raise AppError(409, "marksheet.duplicate", "That student is already on this marksheet.")
    result_id = str(ObjectId())
    await db(request).results.insert_one(
        {
            "_id": result_id,
            "institute_id": user["institute_id"],
            "marksheet_id": sheet["_id"],
            "enrollment_id": enrollment["_id"],
            "student_id": body.student_id,
            "revision": revision,
            "active": False,
            "status": status,
            "score": score,
            "maximum": sheet.get("maximum"),
            "exam_date": sheet.get("exam_date"),
            "title": sheet.get("title"),
            "assessment_id": sheet.get("assessment_id"),
            "attempt_kind": sheet.get("attempt_kind") or "original",
            "exam_type": sheet.get("exam_type"),
            "counts_in_aggregate": True,
            "branch_id": sheet.get("branch_id"),
            "course_id": sheet.get("course_id"),
            "batch_id": enrollment["batch_id"],
            "subject_id": sheet.get("subject_id"),
            "paper_id": sheet.get("paper_id"),
            "source": {"sheet": "Entered here"},
        }
    )
    version = sheet.get("edit_version", 0) + 1
    await db(request).marksheets.update_one({"_id": sheet["_id"]}, {"$set": {"edit_version": version}})
    from app.assessments.marks import refresh_sheet_marks

    await refresh_sheet_marks(db(request), sheet)
    return {"id": result_id, "edit_version": version}


def marksheet_visible(user: dict, sheet: dict) -> bool:
    resource = {
        "institute_id": sheet["institute_id"],
        "branch_id": sheet.get("branch_id"),
        "course_id": sheet.get("course_id"),
        "subject_id": sheet.get("subject_id"),
    }
    return allows(user["grants"], "marksheet.view", resource)


@router.get("/marksheets")
async def list_marksheets(
    request: Request,
    status: str = "",
    exam_type: str = "",
    branch_id: str = "",
    course_id: str = "",
    batch_id: str = "",
    subject_id: str = "",
    paper_id: str = "",
    exam_date: str = "",
    exam_date_from: str = "",
    exam_date_to: str = "",
    period: str = "",
    q: str = "",
    sort: str = "title",
    page: int = 0,
    page_size: int = 8,
) -> dict:
    from app.http import _restrict_batches, resolve_academic

    user = await current_user(request)
    if not any("marksheet.view" in grant["actions"] for grant in user["grants"]):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")
    academic, _dropped, period_ids = await resolve_academic(
        request, user, "marksheet.view", branch_id, course_id, batch_id, period
    )
    query: dict = {"institute_id": user["institute_id"]}
    for field, value in (
        ("status", status),
        ("exam_type", exam_type),
        ("branch_id", academic.get("branch_id", "")),
        ("course_id", academic.get("course_id", "")),
        ("subject_id", subject_id),
        ("paper_id", paper_id),
        ("exam_date", exam_date),
    ):
        if value:
            query[field] = value
    _restrict_batches(query, "batch_ids", academic.get("batch_id", ""), period_ids)
    if not exam_date and (exam_date_from or exam_date_to):
        window = {}
        if exam_date_from:
            window["$gte"] = exam_date_from
        if exam_date_to:
            window["$lte"] = exam_date_to
        query["exam_date"] = window
    named_sheets: set[str] = set()
    if q:
        matched_students = (
            await db(request)
            .students.find(
                {"institute_id": user["institute_id"], "display_name": {"$regex": q, "$options": "i"}},
                {"_id": 1},
            )
            .to_list(length=200)
        )
        student_ids = [item["_id"] for item in matched_students]
        if student_ids:
            named_sheets = {
                item["marksheet_id"]
                for item in await db(request)
                .results.find(
                    {"student_id": {"$in": student_ids}},
                    {"marksheet_id": 1},
                )
                .to_list(length=2000)
            }
    items = []
    names = {}
    cursor = db(request).marksheets.find(query)
    async for sheet in cursor:
        if not marksheet_visible(user, sheet):
            continue
        summary = _marksheet_summary(sheet)
        haystack = " ".join(
            str(summary.get(field) or "") for field in ("title", "source_file", "source_sheet")
        ).casefold()
        if q and q.casefold() not in haystack and sheet["_id"] not in named_sheets:
            continue
        for field, key in (("uploader_id", "uploader_name"), ("reviewer_id", "reviewer_name")):
            person_id = sheet.get(field)
            if person_id and person_id not in names:
                person = await db(request).users.find_one({"_id": person_id})
                names[person_id] = person["name"] if person else ""
            summary[key] = names.get(person_id, "")
        items.append(summary)
    items.sort(
        key=lambda item: str(item.get(sort if sort in {"exam_date", "status", "title"} else "title") or "").casefold()
    )
    total = len(items)
    current = page or 1
    if page:
        start = (page - 1) * page_size
        items = items[start : start + page_size]
    pages = max(1, (total + page_size - 1) // page_size)
    return {"items": items, "total": total, "page": current, "pages": pages}


@router.post("/marksheets/{marksheet_id}/submit")
async def submit_route(request: Request, marksheet_id: str) -> dict:
    user, sheet = await _editable(request, marksheet_id, "marksheet.submit")
    return await submit_marksheet(db(request), user, sheet)


@router.post("/marksheets/{marksheet_id}/reject")
async def reject_route(request: Request, marksheet_id: str) -> dict:
    user, sheet = await _editable(request, marksheet_id, "marksheet.publish")
    return await reject_marksheet(db(request), user, sheet)


@router.post("/marksheets/{marksheet_id}/correct")
async def correct_route(request: Request, marksheet_id: str, body: ReasonBody) -> dict:
    user, sheet = await _editable(request, marksheet_id, "marksheet.correct")
    return await correct_marksheet(
        db(request), request.app.state.client, user, sheet, body.reason, body.changes, body.edit_version
    )


@router.post("/marksheets/{marksheet_id}/withdraw")
async def withdraw_route(request: Request, marksheet_id: str, body: ReasonBody) -> dict:
    user, sheet = await _editable(request, marksheet_id, "marksheet.withdraw")
    return await withdraw_marksheet(db(request), request.app.state.client, user, sheet, body.reason, body.edit_version)


@router.patch("/marksheets/{marksheet_id}")
async def patch_marksheet(request: Request, marksheet_id: str, body: MarksheetPatch) -> dict:
    wants_bands = "bands" in body.model_fields_set
    if body.exam_date and wants_bands:
        _user, sheet = await _editable(request, marksheet_id, "marksheet.edit_draft")
    elif wants_bands:
        sheet = await db(request).marksheets.find_one({"_id": marksheet_id})
        action = "marksheet.edit_draft" if sheet and sheet.get("status") == "draft" else "marksheet.correct"
        _user, sheet = await _editable(request, marksheet_id, action)
    else:
        _user, sheet = await _editable(request, marksheet_id, "marksheet.edit_draft")
    updates = {}
    unset = {}
    if body.exam_date:
        if sheet.get("status") != "draft":
            raise AppError(422, "marksheet.state", "Only a draft can be edited.")
        updates["exam_date"] = body.exam_date
    if wants_bands:
        if body.bands is None:
            unset["bands"] = ""
        else:
            from app.calculating import prepare_bands

            try:
                updates["bands"] = prepare_bands([item.model_dump() for item in body.bands])
            except ValueError as error:
                raise AppError(422, "policy.bands", str(error)) from None
    if updates or unset:
        if body.exam_date:
            updates["edit_version"] = sheet.get("edit_version", 0) + 1
        change = {}
        if updates:
            change["$set"] = updates
        if unset:
            change["$unset"] = unset
        await db(request).marksheets.update_one({"_id": sheet["_id"]}, change)
    from app.assessments.marks import refresh_sheet_marks
    from app.calculating import public_bands

    fresh = {**sheet, **updates}
    if "bands" in unset:
        fresh.pop("bands", None)
    if wants_bands:
        await refresh_sheet_marks(db(request), fresh)
    saved = updates.get("bands")
    if "bands" in unset:
        shown = None
    elif saved:
        shown = public_bands(saved)
    elif sheet.get("bands"):
        shown = public_bands(sheet.get("bands"))
    else:
        shown = None
    return {"id": sheet["_id"], **{key: value for key, value in updates.items() if key != "bands"}, "bands": shown}


@router.patch("/marksheets/{marksheet_id}/results/{result_id}")
async def patch_result(request: Request, marksheet_id: str, result_id: str, body: ResultPatch) -> dict:
    _user, sheet = await _editable(request, marksheet_id, "marksheet.edit_draft")
    require_current_version(sheet, body.edit_version)
    if sheet.get("status") != "draft":
        raise AppError(422, "marksheet.state", "Only a draft result can be corrected here.")
    result = await db(request).results.find_one({"_id": result_id, "marksheet_id": sheet["_id"]})
    if result is None:
        raise AppError(404, "not_found", "That record was not found.")
    updates = {"previous_score": result.get("score")}
    if body.status in {"absent", "missing"}:
        updates["status"] = "absent"
        updates["score"] = None
    elif body.status == "scored" or body.score is not None:
        if body.score is None:
            raise AppError(422, "import.invalid_mark", "Enter a score, or mark the row absent.")
        if body.score < 0 or (sheet.get("maximum") is not None and body.score > sheet["maximum"]):
            raise AppError(422, "import.invalid_mark", "The score must be from 0 through the maximum.")
        updates["score"] = body.score
        updates["status"] = "scored"
    elif body.status:
        raise AppError(422, "import.invalid_mark", "Choose scored or absent.")
    version = sheet.get("edit_version", 0) + 1
    await db(request).results.update_one({"_id": result_id}, {"$set": updates})
    await db(request).marksheets.update_one({"_id": sheet["_id"]}, {"$set": {"edit_version": version}})
    from app.assessments.marks import refresh_sheet_marks

    await refresh_sheet_marks(db(request), sheet)
    return {"id": result_id, "edit_version": version, **updates}


async def _plan_totals(database, institute_id: str, plan: dict, known: dict) -> dict:
    seen = {key: list(value) for key, value in known.items()}
    create = 0
    reuse = 0
    for sheet in plan["sheets"]:
        if not sheet["included"]:
            continue
        for group in sheet["groups"]:
            if group["interpretation"] == "skip_duplicate":
                continue
            for student in group["students"]:
                if not str(student.get("display_name") or "").strip():
                    continue
                sessions = seen.get(student["name_key"], [])
                if student.get("batch_session") in sessions:
                    reuse += 1
                else:
                    create += 1
                    if student.get("batch_session"):
                        seen.setdefault(student["name_key"], []).append(student["batch_session"])
    branch = await database.branches.find_one({"institute_id": institute_id})
    return {"students_to_create": create, "students_to_reuse": reuse, "catalog_started": branch is not None}


async def _job(request: Request, user: dict, import_id: str) -> dict:
    job = await db(request).import_jobs.find_one({"_id": import_id, "institute_id": user["institute_id"]})
    if job is None:
        raise AppError(404, "not_found", "That record was not found.")
    if not any(
        "marksheet.upload" in grant["actions"] or "marksheet.view" in grant["actions"] for grant in user["grants"]
    ):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")
    return job


async def _editable(request: Request, marksheet_id: str, action: str) -> tuple[dict, dict]:
    user = await current_user(request)
    sheet = await db(request).marksheets.find_one({"_id": marksheet_id, "institute_id": user["institute_id"]})
    if sheet is None:
        raise AppError(404, "not_found", "That record was not found.")
    resource = {
        "institute_id": sheet["institute_id"],
        "branch_id": sheet.get("branch_id"),
        "course_id": sheet.get("course_id"),
        "subject_id": sheet.get("subject_id"),
    }
    if not allows(user["grants"], "marksheet.view", resource):
        raise AppError(404, "not_found", "That record was not found.")
    if not allows(user["grants"], action, resource):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")
    return user, sheet


def _stored_rows(request: Request, job: dict) -> dict:
    file_id = job.get("file_id")
    if not file_id:
        return {}
    path = request.app.state.settings.data_dir / "files" / f"{file_id}.xlsx"
    if not path.exists():
        return {}
    try:
        return workbook_rows(path.read_bytes())
    except ValueError:
        return {}


def _public_sheet(sheet: dict, rows: list | None = None) -> dict:
    return {
        "id": sheet["id"],
        "name": sheet["name"],
        "included": sheet["included"],
        "title": sheet["title"],
        "suggested_title": sheet["suggested_title"],
        "branch_name": sheet["branch_name"],
        "course_name": sheet["course_name"],
        "subject_name": sheet["subject_name"],
        "paper_number": sheet["paper_number"],
        "batches": sheet["batches"],
        "heading_conflict": sheet["heading_conflict"],
        "heading_acknowledged": sheet["heading_acknowledged"],
        "exam_date": sheet["exam_date"],
        "rows": rows or [],
        "attempt": sheet["attempt"],
        "exam_type": sheet["exam_type"],
        "planned": sheet.get("planned"),
        "exam": sheet.get("exam"),
        "similar_exams": sheet.get("similar_exams") or [],
        "matched_exam": sheet.get("matched_exam"),
        "selected_batch_ids": sheet.get("selected_batch_ids") or [],
        "cover_roster": bool(sheet.get("cover_roster")),
        "roster_reason": sheet.get("roster_reason") or "",
        "groups": [
            {
                "id": group["id"],
                "maximum": group["maximum"],
                "interpretation": group["interpretation"],
                "confirmed": group["confirmed"],
                "duplicate_of": group["duplicate_of"],
                "students": group["students"],
            }
            for group in sheet["groups"]
        ],
        "blockers": sheet["blockers"],
    }


def _marksheet_summary(sheet: dict) -> dict:
    return {
        "id": sheet["_id"],
        "title": sheet.get("title") or sheet.get("status"),
        "status": sheet.get("status"),
        "branch_id": sheet.get("branch_id"),
        "course_id": sheet.get("course_id"),
        "subject_id": sheet.get("subject_id"),
        "paper_id": sheet.get("paper_id"),
        "batch_ids": sheet.get("batch_ids") or [],
        "exam_date": sheet.get("exam_date"),
        "exam_type": sheet.get("exam_type"),
        "maximum": sheet.get("maximum"),
        "source_file": sheet.get("source_file"),
        "source_sheet": sheet.get("source_sheet"),
        "file_id": sheet.get("file_id"),
        "revision": sheet.get("active_revision") or sheet.get("revision"),
        "edit_version": sheet.get("edit_version", 0),
    }

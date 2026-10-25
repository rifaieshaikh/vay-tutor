"""Enquiry and admission intake, separate from mark-sheet import."""

from __future__ import annotations

import hashlib
import json
from datetime import datetime, timezone
from io import BytesIO

from bson import ObjectId
from fastapi import APIRouter, File, Form, Header, Request, UploadFile
from fastapi.responses import Response
from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict
from pymongo import ReturnDocument
from pymongo.errors import DuplicateKeyError

from app.acl import allows
from app.db import allocate_enrollment_code, allocate_student_code
from app.deps import current_user, db, write_audit
from app.enrollment import (
    assign_enrollment_code,
    batch_is_closed,
    enrollment_document,
    name_conflicts,
    name_document,
    require_open_batch,
)
from app.errors import AppError
from app.security import name_key, phone_key

router = APIRouter(prefix="/api/v1")

INTERESTS = ("High", "Medium", "Low")
CLOSURE_REASONS = (
    "Not interested",
    "Joined elsewhere",
    "Unable to contact",
    "Duplicate entry",
    "Created by mistake",
    "Other",
)
PERSON_FIELDS = ("name", "whatsapp", "email", "present_institution")
GENDERS = ("Female", "Male", "Other")
BLOOD_GROUPS = {"A+", "A-", "B+", "B-", "O+", "O-", "AB+", "AB-"}
ANSWER_TYPES = {"boolean", "yes_no", "choice", "text", "textarea"}
_MAX_BYTES = 5 * 1024 * 1024
_RECORD_LIMITS = {
    "qualification": 120,
    "stream": 120,
    "institution": 200,
    "mark": 40,
    "board": 160,
}


def _now() -> datetime:
    return datetime.now(timezone.utc)


def _moment_text(value) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _expect_current(row: dict, expected: str) -> None:
    if not expected:
        return
    if _moment_text(row.get("updated_at")) != expected:
        raise AppError(409, "admission.conflict", "This application changed. Reload it before saving.")


def _reader(user: dict) -> None:
    if not any("student.lookup" in (grant.get("actions") or []) for grant in user["grants"]):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")


def _writer(user: dict) -> None:
    if not any("student.manage" in (grant.get("actions") or []) for grant in user["grants"]):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")


def _catalog_editor(user: dict) -> None:
    allowed = any(
        "catalog.correct" in (grant.get("actions") or []) or "catalog.manage" in (grant.get("actions") or [])
        for grant in user["grants"]
    )
    if not allowed:
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")


def _text(value: str | None, limit: int, field: str, code: str = "intake.field") -> str:
    text = " ".join(str(value or "").split())
    if len(text) > limit:
        raise AppError(422, code, f"Use {limit} characters or fewer.", field=field)
    return text


def _long_text(value: str | None, limit: int, field: str) -> str:
    text = str(value or "").strip()
    if len(text) > limit:
        raise AppError(422, "intake.field", f"Use {limit} characters or fewer.", field=field)
    return text


def _choice(value: str | None, allowed: tuple[str, ...] | set[str], field: str, label: str) -> str:
    text = " ".join(str(value or "").split())
    if not text:
        return ""
    match = next((item for item in allowed if item.casefold() == text.casefold()), "")
    if not match:
        raise AppError(422, "intake.field", label, field=field)
    return match


def _blood(value: str | None) -> str:
    text = "".join(str(value or "").split()).upper()
    if not text:
        return ""
    match = next((item for item in BLOOD_GROUPS if item.upper() == text), "")
    if not match:
        raise AppError(422, "intake.field", "Choose a blood group.", field="blood_group")
    return match


def _email(value: str | None) -> str:
    text = _text(value, 120, "email")
    if text and ("@" not in text[1:-1] or "." not in text.split("@", 1)[1]):
        raise AppError(422, "intake.field", "Use an email address.", field="email")
    return text


def _pin(value: str | None) -> str:
    text = "".join(str(value or "").split())
    if text and (len(text) != 6 or not text.isdigit()):
        raise AppError(422, "intake.field", "Use a 6-digit PIN.", field="pin")
    return text


def _date(value: str | None, field: str, *, future_ok: bool = True) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    try:
        parsed = datetime.strptime(text, "%Y-%m-%d").date()
    except ValueError as exc:
        raise AppError(422, "intake.field", "Use a date as YYYY-MM-DD.", field=field) from exc
    if not future_ok and parsed > _now().date():
        raise AppError(422, "intake.field", "The date of birth is in the future.", field=field)
    return text


def _year(value: str | None) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    year = int(text) if len(text) == 4 and text.isdigit() else 0
    if year < 1950 or year > _now().year + 1:
        raise AppError(422, "intake.field", "Use a year of passing as YYYY.", field="year_of_passing")
    return text


def clean_academic_records(items: list | None) -> list[dict]:
    records = []
    for item in items or []:
        if not isinstance(item, dict) and not hasattr(item, "qualification"):
            raise AppError(
                422, "intake.field", "Each academic record needs qualification, stream, institution, mark, and year."
            )
        source = item if isinstance(item, dict) else item.model_dump()
        record = {
            "qualification": _text(source.get("qualification"), _RECORD_LIMITS["qualification"], "qualification"),
            "stream": _text(source.get("stream"), _RECORD_LIMITS["stream"], "stream"),
            "institution": _text(source.get("institution"), _RECORD_LIMITS["institution"], "institution"),
            "mark": _text(source.get("mark"), _RECORD_LIMITS["mark"], "mark"),
            "year_of_passing": _year(source.get("year_of_passing")),
            "board": _text(source.get("board"), _RECORD_LIMITS["board"], "board"),
        }
        if any(record.values()):
            records.append(record)
    if len(records) > 12:
        raise AppError(422, "intake.field", "Use 12 academic records or fewer.", field="academic_records")
    return records


def displayed_records(student: dict) -> list[dict]:
    stored = student.get("academic_records")
    if isinstance(stored, list) and stored:
        return stored
    legacy = {
        "qualification": student.get("qualification") or "",
        "stream": "",
        "institution": student.get("institution") or "",
        "mark": "",
        "year_of_passing": student.get("passing_year") or "",
        "board": student.get("board") or "",
    }
    if any(legacy.values()):
        return [legacy]
    return []


def _mirror_legacy(records: list[dict]) -> dict[str, str]:
    first = records[0] if records else {}
    return {
        "qualification": first.get("qualification") or "",
        "institution": first.get("institution") or "",
        "board": first.get("board") or "",
        "passing_year": first.get("year_of_passing") or "",
    }


class QuestionBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    label: str = ""
    answer_type: str = "text"
    choices: list[str] = []
    required: bool = False
    place: str = ""
    section: str = "general"
    required_stage: str = ""
    on_profile: bool = False


class QuestionPatch(BaseModel):
    model_config = ConfigDict(extra="forbid")
    label: str | None = None
    answer_type: str | None = None
    choices: list[str] | None = None
    required: bool | None = None
    order: int | None = None
    place: str | None = None
    section: str | None = None
    required_stage: str | None = None
    on_profile: bool | None = None


class AnswerBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    question_id: str = ""
    value: str = ""
    label: str = ""


class FollowUpBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    at: str = ""
    note: str = ""
    interest: str = ""
    next_on: str = ""
    kind: str = ""
    task_id: str = ""


class EnquiryBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    name: str = ""
    whatsapp: str = ""
    email: str = ""
    present_institution: str = ""
    course_id: str = ""
    interest: str = ""
    references: str = ""
    details: str = ""
    answers: list[AnswerBody] = []
    prospect_id: str = ""
    episode_note: str = ""


class CloseBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    reason: str = ""
    note: str = ""
    expected_updated_at: str = ""


class ReopenBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    reason: str = ""
    expected_updated_at: str = ""


class ImportSourceBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    sheet: str = ""
    header_row: int = 1


class ImportMappingBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    column_id: str = ""
    target: str = ""


class ImportDefaultBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    course_id: str = ""
    interest: str = ""
    mode: str = "blanks"


class ImportDecisionBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    key: str = ""
    action: str = ""
    prospect_id: str = ""
    student_id: str = ""
    episode_note: str = ""
    values: dict[str, str] = {}


class ImportPlanBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    revision: int = 0
    delimiter: str = ""
    sources: list[ImportSourceBody] = []
    mappings: list[ImportMappingBody] = []
    defaults: ImportDefaultBody = ImportDefaultBody()
    courses: dict[str, str] = {}
    interests: dict[str, str] = {}
    date_format: str = ""
    formulas: dict[str, dict] = {}
    decisions: list[ImportDecisionBody] = []
    identity_mode: str = "manual"


class ImportCommitBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    revision: int
    idempotency_key: str = ""


class MappingSaveBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    import_id: str = ""
    name: str = ""


class MappingApplyBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    mapping_id: str = ""


class EnquiryCommit(BaseModel):
    model_config = ConfigDict(extra="forbid")
    rows: list[EnquiryBody]


class AdmissionBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    course_id: str = ""
    name: str = ""
    phone: str = ""
    date_of_birth: str = ""
    gender: str = ""
    blood_group: str = ""
    email: str = ""
    street: str = ""
    address_line_2: str = ""
    place: str = ""
    district: str = ""
    state: str = ""
    pin: str = ""
    guardian_name: str = ""
    guardian_phone: str = ""
    academic_records: list[dict] = []
    course_answers: list[AnswerBody] = []
    business_answers: list[AnswerBody] = []
    prospect_id: str = ""
    episode_note: str = ""
    confirm_course_change: bool = False
    expected_updated_at: str = ""


class AdmissionCommitRow(AdmissionBody):
    decision: str = ""
    student_id: str = ""


class AdmissionCommit(BaseModel):
    model_config = ConfigDict(extra="forbid")
    rows: list[AdmissionCommitRow]


class ConvertBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    branch_id: str
    batch_id: str
    decision: str = ""
    student_id: str = ""
    expected_updated_at: str = ""


class IdentityBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    decision: str
    student_id: str = ""
    expected_updated_at: str = ""


class SelectionBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    branch_id: str
    batch_id: str
    expected_updated_at: str = ""


class CorrectionBody(AdmissionBody):
    note: str = ""


def _question_view(row: dict) -> dict:
    return {
        "id": row["_id"],
        "course_id": row.get("course_id"),
        "label": row.get("label") or "",
        "answer_type": row.get("answer_type") or "text",
        "choices": row.get("choices") or [],
        "required": bool(row.get("required")),
        "order": row.get("order") or 0,
        "place": _place_of(row.get("forms") or []),
        "section": row.get("section") or "general",
        "required_stage": row.get("required_stage") or "",
        "on_profile": bool(row.get("on_profile")),
        "retired": bool(row.get("retired")),
    }


def _clean_question(body: QuestionBody, existing: dict | None = None) -> dict:
    label = _text(body.label, 120, "label")
    if not label:
        raise AppError(422, "intake.field", "Enter the question.", field="label")
    answer_type = body.answer_type or (existing or {}).get("answer_type") or "text"
    if answer_type not in ANSWER_TYPES:
        raise AppError(
            422,
            "intake.field",
            "Choose yes/no, a consent checkbox, a dropdown, a text box, or a text area.",
            field="answer_type",
        )
    choices = []
    for item in body.choices:
        choice = _text(item, 80, "choices")
        if choice and choice not in choices:
            choices.append(choice)
    if answer_type == "choice" and not choices:
        raise AppError(422, "intake.field", "Add at least one choice.", field="choices")
    if answer_type != "choice":
        choices = []
    section = (body.section or "general").strip() or "general"
    if section not in {"general", "contact", "guardian", "education", "course"}:
        raise AppError(422, "intake.field", "Choose a section.", field="section")
    stage = (body.required_stage or "").strip()
    if stage not in {"", "admit", "enroll"}:
        raise AppError(422, "intake.field", "Choose when the answer is required.", field="required_stage")
    return {
        "label": label,
        "label_key": name_key(label),
        "answer_type": answer_type,
        "choices": choices,
        "required": bool(body.required) or bool(stage),
        "section": section,
        "required_stage": stage,
        "on_profile": bool(body.on_profile),
    }


def _forms_for(place: str) -> list[str]:
    chosen = place or "both"
    if chosen == "both":
        return ["enquiry", "admission"]
    if chosen not in {"enquiry", "admission"}:
        raise AppError(422, "intake.field", "Choose enquiry, admission, or both.", field="place")
    return [chosen]


def _place_of(forms: list) -> str:
    names = set(forms or [])
    if "enquiry" in names and "admission" in names:
        return "both"
    if "enquiry" in names:
        return "enquiry"
    if "admission" in names:
        return "admission"
    return ""


async def _course(request: Request, user: dict, course_id: str) -> dict:
    if not course_id:
        raise AppError(422, "intake.field", "Choose a course.", field="course_id")
    course = await db(request).courses.find_one({"_id": course_id, "institute_id": user["institute_id"]})
    if course is None:
        raise AppError(422, "intake.field", "Choose a course.", field="course_id")
    return course


async def _questions_for(request: Request, course_id: str, include_retired: bool = False) -> list[dict]:
    if not course_id:
        return []
    query: dict = {"course_id": course_id}
    if not include_retired:
        query["retired"] = {"$ne": True}
    return await db(request).course_questions.find(query).sort("order", 1).to_list(length=50)


async def _form_questions(request: Request, institute_id: str, form: str) -> list[dict]:
    return (
        await db(request)
        .form_questions.find({"institute_id": institute_id, "forms": form, "retired": {"$ne": True}})
        .sort("order", 1)
        .to_list(length=50)
    )


async def clean_form_answers(
    request: Request, institute_id: str, form: str, answers: list, *, require: bool
) -> list[dict]:
    questions = await _form_questions(request, institute_id, form)
    return _clean_answers(questions, answers, require=require)


async def remember_answers(
    request: Request, institute_id: str, form: str, existing: list, incoming: list, *, require: bool
) -> list[dict]:
    questions = await _form_questions(request, institute_id, form)
    cleaned = _clean_answers(questions, incoming, require=require)
    current = {item["_id"] for item in questions}
    kept = [item for item in existing or [] if item.get("question_id") not in current]
    return kept + cleaned


def _answer_value(question: dict, value: str) -> str:
    kind = question.get("answer_type") or "text"
    label = question.get("label") or "the question"
    if kind == "textarea":
        return _long_text(value, 2000, "answers")
    text = " ".join(str(value or "").split())
    if not text:
        return ""
    if kind == "yes_no":
        if text.casefold() in {"yes", "true", "1", "y"}:
            return "Yes"
        if text.casefold() in {"no", "false", "0", "n"}:
            return "No"
        raise AppError(422, "intake.field", f"Choose Yes or No for {label}.", field="answers")
    if kind == "boolean":
        if text.casefold() in {"yes", "true", "1", "y"}:
            return "Yes"
        if text.casefold() in {"no", "false", "0", "n"}:
            return ""
        raise AppError(422, "intake.field", f"Tick {label} or leave it clear.", field="answers")
    if kind == "text":
        return _text(text, 300, "answers")
    if kind == "number":
        try:
            float(text)
        except ValueError as exc:
            raise AppError(422, "intake.field", f"Enter a number for {label}.", field="answers") from exc
        return text
    if kind == "date":
        return _date(text, "answers")
    choices = question.get("choices") or []
    match = next((item for item in choices if item.casefold() == text.casefold()), "")
    if not match:
        raise AppError(
            422,
            "intake.field",
            f"Choose an answer for {label}.",
            field="course_answers",
        )
    return match


def _clean_answers(questions: list[dict], answers: list[AnswerBody] | list[dict], *, require: bool) -> list[dict]:
    by_id = {item["_id"]: item for item in questions}
    cleaned = []
    seen = set()
    for item in answers:
        source = item if isinstance(item, dict) else item.model_dump()
        question_id = source.get("question_id") or ""
        question = by_id.get(question_id)
        if question is None:
            raise AppError(422, "intake.field", "One of the answers does not match a question.", field="answers")
        if question_id in seen:
            continue
        seen.add(question_id)
        value = _answer_value(question, source.get("value") or "")
        cleaned.append({"question_id": question_id, "label": question.get("label") or "", "value": value})
    if require:
        missing = [
            item
            for item in questions
            if item.get("required")
            and not next((answer["value"] for answer in cleaned if answer["question_id"] == item["_id"]), "")
        ]
        if missing:
            label = missing[0].get("label") or "the question"
            message = f"Tick {label}." if missing[0].get("answer_type") == "boolean" else f"Answer {label}."
            raise AppError(422, "intake.field", message, field="answers")
    return cleaned


def _enquiry_fields(body: EnquiryBody) -> dict:
    name = _text(body.name, 120, "name")
    if not name:
        raise AppError(422, "intake.field", "Enter the name.", field="name")
    whatsapp = _text(body.whatsapp, 30, "whatsapp")
    email = _email(body.email)
    return {
        "name": name,
        "name_key": name_key(name),
        "whatsapp": whatsapp,
        "whatsapp_key": phone_key(whatsapp),
        "email": email,
        "email_key": name_key(email),
        "present_institution": _text(body.present_institution, 200, "present_institution"),
        "course_id": body.course_id,
        "interest": _choice(body.interest, INTERESTS, "interest", "Choose High, Medium, or Low."),
        "references": _long_text(body.references, 500, "references"),
        "details": _long_text(body.details, 1000, "details"),
    }


def _admission_contact(body: AdmissionBody) -> dict:
    name = _text(body.name, 120, "name")
    phone = _text(body.phone, 30, "phone")
    if not name:
        raise AppError(422, "intake.field", "Enter the name.", field="name")
    return {
        "name": name,
        "name_key": name_key(name),
        "phone": phone,
        "phone_key": phone_key(phone),
        "date_of_birth": _date(body.date_of_birth, "date_of_birth", future_ok=False),
        "gender": _choice(body.gender, GENDERS, "gender", "Choose Female, Male, or Other."),
        "blood_group": _blood(body.blood_group),
        "email": _email(body.email),
        "street": _long_text(body.street, 200, "street"),
        "address_line_2": _long_text(body.address_line_2, 200, "address_line_2"),
        "place": _text(body.place, 80, "place"),
        "district": _text(body.district, 80, "district"),
        "state": _text(body.state, 80, "state"),
        "pin": _pin(body.pin),
        "guardian_name": _text(body.guardian_name, 120, "guardian_name"),
        "guardian_phone": _text(body.guardian_phone, 30, "guardian_phone"),
        "academic_records": clean_academic_records(body.academic_records),
    }


def _merged_question(body: QuestionPatch, row: dict) -> QuestionBody:
    return QuestionBody(
        label=body.label if body.label is not None else row.get("label") or "",
        answer_type=body.answer_type if body.answer_type is not None else row.get("answer_type") or "text",
        choices=body.choices if body.choices is not None else row.get("choices") or [],
        required=body.required if body.required is not None else bool(row.get("required")),
        place=body.place if body.place is not None else _place_of(row.get("forms") or []),
        section=body.section if body.section is not None else row.get("section") or "general",
        required_stage=body.required_stage if body.required_stage is not None else row.get("required_stage") or "",
        on_profile=body.on_profile if body.on_profile is not None else bool(row.get("on_profile")),
    )


def enquiry_status(row: dict) -> str:
    if row.get("admission_id"):
        return "converted"
    status = row.get("status") or ""
    if status in {"open", "closed", "converted"}:
        return status
    return "open"


def enquiry_next_on(row: dict) -> str:
    if "pending_on" in row:
        return str(row.get("pending_on") or "")
    follow_ups = row.get("follow_ups") or []
    if not follow_ups:
        return ""
    return str(follow_ups[-1].get("next_on") or "")


def follow_up_state(row: dict, today: str) -> str:
    """Overdue, today, and upcoming use the institute timezone's calendar date."""
    status = enquiry_status(row)
    if status == "converted":
        return "admitted"
    if status == "closed":
        return "closed"
    next_on = enquiry_next_on(row)
    if not next_on:
        return "unscheduled"
    if next_on < today:
        return "overdue"
    if next_on == today:
        return "today"
    return "upcoming"


def enquiry_progress(row: dict, today: str) -> tuple[str, str]:
    status = enquiry_status(row)
    if status == "converted":
        return "Converted", "Open the admission"
    if status == "closed":
        return "Closed", row.get("closure_reason") or "Closed"
    state = follow_up_state(row, today)
    if state == "unscheduled":
        return "Open", "Record a follow-up"
    if state == "upcoming":
        return "Open", f"Follow up on {enquiry_next_on(row)}"
    return "Open", "Follow up"


def admission_lifecycle(row: dict) -> str:
    status = row.get("status") or "open"
    if status == "converted":
        return "enrolled"
    if status == "closed":
        return "closed"
    return "open"


def _missing_answer(question: dict, answers: list) -> str:
    value = next(
        (item.get("value") for item in answers if item.get("question_id") == question.get("_id")),
        "",
    )
    if value:
        return ""
    label = question.get("label") or "the question"
    if question.get("answer_type") == "boolean":
        return f"Tick {label}"
    return f"Answer {label}"


def admission_requirements(
    row: dict, course_questions: list[dict], business_questions: list[dict], *, available: bool = True
) -> list[dict]:
    """Structured gaps. Enrolled and closed records are historical and are not incomplete."""
    if not available:
        return [
            {
                "key": "requirements.unavailable",
                "state": "unavailable",
                "section": "requirements",
                "field": "",
                "label": "Requirements",
                "message": "Requirements could not be loaded.",
            }
        ]
    if admission_lifecycle(row) != "open":
        return []
    items = []
    if not (row.get("phone") or "").strip():
        items.append(
            {
                "key": "contact.phone",
                "state": "missing",
                "section": "contact",
                "field": "phone",
                "label": "Contact number",
                "message": "Add the contact number",
            }
        )
    if not row.get("course_id"):
        items.append(
            {
                "key": "course",
                "state": "missing",
                "section": "course",
                "field": "course_id",
                "label": "Course",
                "message": "Choose a course",
            }
        )
    for question in business_questions:
        if not question.get("required"):
            continue
        message = _missing_answer(question, row.get("business_answers") or [])
        if message:
            items.append(
                {
                    "key": f"business.{question.get('_id')}",
                    "state": "missing",
                    "section": "additional",
                    "field": question.get("_id") or "",
                    "label": question.get("label") or "Question",
                    "message": message,
                }
            )
    for question in course_questions:
        if not question.get("required"):
            continue
        message = _missing_answer(question, row.get("course_answers") or [])
        if message:
            items.append(
                {
                    "key": f"course.{question.get('_id')}",
                    "state": "missing",
                    "section": "course",
                    "field": question.get("_id") or "",
                    "label": question.get("label") or "Question",
                    "message": message,
                }
            )
    return items


def admission_blockers(row: dict, course_questions: list[dict], business_questions: list[dict]) -> list[str]:
    return [item["message"] for item in admission_requirements(row, course_questions, business_questions)]


def identity_resolved(row: dict) -> bool:
    decision = row.get("match_decision") or ""
    if decision == "create":
        return True
    return decision == "attach" and bool(row.get("match_student_id"))


def admission_readiness(row: dict, requirements: list[dict], *, batch_open: bool = False) -> str:
    lifecycle = admission_lifecycle(row)
    if lifecycle != "open":
        return lifecycle
    if any(item.get("state") == "unavailable" for item in requirements):
        return "unavailable"
    if requirements:
        return "needs_information"
    if not identity_resolved(row):
        return "needs_identity"
    if not row.get("branch_id") or not row.get("batch_id") or not batch_open:
        return "needs_batch"
    return "ready"


def admission_progress(
    row: dict, blockers: list[str] | None = None, matches: list[dict] | None = None, readiness: str = ""
) -> tuple[str, str]:
    lifecycle = admission_lifecycle(row)
    if lifecycle == "enrolled":
        return "Enrolled", "Open the student"
    if lifecycle == "closed":
        return "Closed", row.get("closure_reason") or "Closed"
    current = readiness or (
        "needs_information"
        if blockers
        else "needs_identity"
        if not identity_resolved(row)
        else "needs_batch"
        if not row.get("batch_id")
        else "ready"
    )
    if current == "unavailable":
        return "Open", "Retry requirements"
    if current == "needs_information":
        return "Open", (blockers or ["Complete the application"])[0]
    if current == "needs_identity":
        return "Open", "Confirm identity"
    if current == "needs_batch":
        return "Open", "Choose a batch"
    return "Open", "Enroll"


def _enquiry_view(row: dict, course_name: str = "", today: str = "") -> dict:
    current = today or _now().date().isoformat()
    stage, next_action = enquiry_progress(row, current)
    state = follow_up_state(row, current)
    return {
        "id": row["_id"],
        "reference": row.get("reference") or "",
        "prospect_id": row.get("prospect_id") or "",
        "name": row.get("name") or "",
        "whatsapp": row.get("whatsapp") or "",
        "email": row.get("email") or "",
        "present_institution": row.get("present_institution") or "",
        "course_id": row.get("course_id") or "",
        "course_name": course_name,
        "interest": row.get("interest") or "",
        "references": row.get("references") or "",
        "details": row.get("details") or "",
        "status": enquiry_status(row),
        "stage": stage,
        "next_action": next_action,
        "next_on": enquiry_next_on(row),
        "pending_task_id": row.get("pending_task_id") or "",
        "last_contact_on": row.get("last_contact_on") or "",
        "follow_up_state": state,
        "closure_reason": row.get("closure_reason") or "",
        "closure_note": row.get("closure_note") or "",
        "episode_note": row.get("episode_note") or "",
        "follow_ups": [
            {
                "id": item.get("id"),
                "at": item.get("at") or "",
                "note": item.get("note") or "",
                "interest": item.get("interest") or "",
                "next_on": item.get("next_on") or "",
                "actor": item.get("actor") or "",
            }
            for item in row.get("follow_ups") or []
        ],
        "activity": _activity(row),
        "admission_id": row.get("admission_id") or None,
        "created_at": row.get("created_at"),
        "today": current,
    }


def _activity(row: dict) -> list[dict]:
    items = []
    for item in row.get("follow_ups") or []:
        items.append(
            {
                "id": item.get("id") or "",
                "kind": "contact",
                "at": item.get("at") or "",
                "note": item.get("note") or "",
                "actor": item.get("actor") or "",
            }
        )
    for item in row.get("events") or []:
        items.append(
            {
                "id": item.get("id") or "",
                "kind": item.get("kind") or "",
                "at": item.get("at") or "",
                "note": item.get("note") or "",
                "actor": item.get("actor") or "",
            }
        )
    items.sort(key=lambda item: (str(item.get("at") or ""), str(item.get("id") or "")))
    return items


def _admission_view(
    row: dict,
    course_name: str = "",
    matches: list[dict] | None = None,
    questions: list[dict] | None = None,
    business_questions: list[dict] | None = None,
    *,
    branch_name: str = "",
    batch_name: str = "",
    batch_open: bool = False,
    requirements_available: bool = True,
) -> dict:
    requirements = row.get("_requirements")
    if requirements is None:
        requirements = admission_requirements(
            row, questions or [], business_questions or [], available=requirements_available
        )
    readiness = row.get("_readiness") or admission_readiness(row, requirements, batch_open=batch_open)
    blockers = [item["message"] for item in requirements if item.get("state") == "missing"]
    stage, next_action = admission_progress(row, blockers, matches, readiness)
    if readiness == "ready" and batch_name:
        next_action = f"Enroll in {batch_name}"
    lifecycle = admission_lifecycle(row)
    return {
        "id": row["_id"],
        "reference": row.get("reference") or "",
        "prospect_id": row.get("prospect_id") or "",
        "course_id": row.get("course_id") or "",
        "course_name": course_name,
        "name": row.get("name") or "",
        "phone": row.get("phone") or "",
        "date_of_birth": row.get("date_of_birth") or "",
        "gender": row.get("gender") or "",
        "blood_group": row.get("blood_group") or "",
        "email": row.get("email") or "",
        "street": row.get("street") or "",
        "address_line_2": row.get("address_line_2") or "",
        "place": row.get("place") or "",
        "district": row.get("district") or "",
        "state": row.get("state") or "",
        "pin": row.get("pin") or "",
        "guardian_name": row.get("guardian_name") or "",
        "guardian_phone": row.get("guardian_phone") or "",
        "academic_records": row.get("academic_records") or [],
        "course_answers": row.get("course_answers") or [],
        "business_answers": row.get("business_answers") or [],
        "questions": [_question_view(item) for item in questions or []],
        "status": row.get("status") or "open",
        "lifecycle": lifecycle,
        "readiness": readiness,
        "requirements": requirements,
        "stage": stage,
        "next_action": next_action,
        "blockers": blockers,
        "student_id": row.get("student_id") or None,
        "enrollment_id": row.get("enrollment_id") or None,
        "branch_id": row.get("branch_id") or None,
        "branch_name": branch_name,
        "batch_id": row.get("batch_id") or None,
        "batch_name": batch_name,
        "enquiry_id": row.get("enquiry_id") or None,
        "closure_reason": row.get("closure_reason") or "",
        "closure_note": row.get("closure_note") or "",
        "activity": _activity(row),
        "match_decision": row.get("match_decision") or "",
        "match_student_id": row.get("match_student_id") or "",
        "matches": matches or [],
        "created_at": row.get("created_at"),
        "updated_at": _moment_text(row.get("updated_at")),
        "episode_note": row.get("episode_note") or "",
        "source": (
            {"sheet": (row.get("source") or {}).get("sheet") or "", "row": (row.get("source") or {}).get("row") or ""}
            if row.get("source")
            else None
        ),
    }


_PROFILE_FIELDS = (
    ("email", "Email"),
    ("date_of_birth", "Date of birth"),
    ("gender", "Gender"),
    ("street", "Address"),
    ("place", "City"),
    ("guardian_name", "Guardian"),
    ("guardian_phone", "Guardian phone"),
)


def _profile_differences(application: dict, student: dict) -> list[dict]:
    found = []
    for key, label in _PROFILE_FIELDS:
        application_value = str(application.get(key) or "")
        student_value = str(student.get(key) or "")
        if application_value and student_value and application_value != student_value:
            found.append(
                {"field": key, "label": label, "application": application_value, "student": student_value}
            )
    return found


def _match_view(student: dict, application: dict | None = None) -> dict:
    view = {
        "id": student["_id"],
        "display_name": student.get("display_name") or "",
        "student_code": student.get("student_code") or "",
        "phone": student.get("phone") or "",
        "email": student.get("email") or "",
        "differences": _profile_differences(application, student) if application else [],
    }
    return view


async def _course_name(request: Request, course_id: str) -> str:
    if not course_id:
        return ""
    course = await db(request).courses.find_one({"_id": course_id})
    return (course or {}).get("name") or ""


async def find_matches(
    request: Request, institute_id: str, name: str, phone: str, application: dict | None = None
) -> list[dict]:
    key = name_key(name)
    number = phone_key(phone)
    if not key or not number:
        return []
    people = (
        await db(request)
        .students.find({"institute_id": institute_id, "name_key": key, "archived": {"$ne": True}})
        .to_list(length=50)
    )
    matched = []
    for person in people:
        stored = person.get("phone_key") or phone_key(person.get("phone") or "")
        if stored == number:
            matched.append(_match_view(person, application))
    return matched


async def student_extras(database, student: dict) -> dict:
    profiles = []
    for item in student.get("course_profiles") or []:
        course_id = item.get("course_id") or ""
        course = await database.courses.find_one({"_id": course_id})
        questions = {
            row["_id"]: row for row in await database.course_questions.find({"course_id": course_id}).to_list(length=50)
        }
        answers = []
        for answer in item.get("answers") or []:
            question = questions.get(answer.get("question_id")) or {}
            answers.append(
                {
                    "question_id": answer.get("question_id") or "",
                    "label": answer.get("label") or question.get("label") or "",
                    "value": answer.get("value") or "",
                    "answer_type": question.get("answer_type") or "text",
                    "choices": question.get("choices") or [],
                }
            )
        profiles.append(
            {
                "course_id": course_id,
                "course_name": (course or {}).get("name") or "",
                "answers": answers,
            }
        )
    return {
        "academic_records": displayed_records(student),
        "course_profiles": profiles,
        "business_answers": student.get("business_answers") or [],
    }


def apply_record_update(student: dict, records: list[dict]) -> tuple[dict, dict, dict, dict]:
    cleaned = clean_academic_records(records)
    sets: dict = {"academic_records": cleaned}
    unsets: dict = {}
    before = {"academic_records": _record_summary(displayed_records(student))}
    after = {"academic_records": _record_summary(cleaned)}
    for field, value in _mirror_legacy(cleaned).items():
        if value:
            sets[field] = value
        else:
            unsets[field] = ""
    return sets, unsets, before, after


def _record_summary(records: list[dict]) -> str:
    if not records:
        return "Not recorded"
    parts = []
    for item in records:
        label = ", ".join(
            item.get(key) or ""
            for key in ("qualification", "stream", "institution", "mark", "year_of_passing", "board")
            if item.get(key)
        )
        if label:
            parts.append(label)
    return "; ".join(parts) or "Not recorded"


async def apply_course_profiles(request: Request, institute_id: str, profiles: list) -> list[dict]:
    cleaned = []
    for item in profiles or []:
        source = item if isinstance(item, dict) else {}
        course_id = source.get("course_id") or ""
        course = await db(request).courses.find_one({"_id": course_id, "institute_id": institute_id})
        if course is None:
            raise AppError(422, "intake.field", "Choose a course for the extra answers.", field="course_profiles")
        questions = await _questions_for(request, course_id)
        answers = _clean_answers(questions, source.get("answers") or [], require=False)
        cleaned.append({"course_id": course_id, "answers": answers})
    return cleaned


@router.get("/form-questions")
async def list_form_questions(request: Request, form: str = "", include_retired: bool = False) -> dict:
    user = await current_user(request)
    _reader(user)
    query: dict = {"institute_id": user["institute_id"]}
    if form:
        if form not in {"enquiry", "admission"}:
            raise AppError(422, "intake.field", "Choose enquiry or admission.", field="form")
        query["forms"] = form
    if not include_retired:
        query["retired"] = {"$ne": True}
    rows = await db(request).form_questions.find(query).sort("order", 1).to_list(length=50)
    return {"items": [_question_view(row) for row in rows]}


@router.post("/form-questions")
async def create_form_question(request: Request, body: QuestionBody) -> dict:
    user = await current_user(request)
    _catalog_editor(user)
    fields = _clean_question(body)
    forms = _forms_for(body.place or "both")
    taken = await db(request).form_questions.find_one(
        {"institute_id": user["institute_id"], "label_key": fields["label_key"]}
    )
    if taken:
        raise AppError(409, "question.exists", "That question already exists.")
    last = (
        await db(request)
        .form_questions.find({"institute_id": user["institute_id"]})
        .sort("order", -1)
        .to_list(length=1)
    )
    order = (last[0].get("order") or 0) + 1 if last else 1
    question_id = str(ObjectId())
    await db(request).form_questions.insert_one(
        {"_id": question_id, "institute_id": user["institute_id"], "order": order, "forms": forms, **fields}
    )
    await write_audit(
        request, user, "form_question.create", None, {"label": fields["label"], "place": _place_of(forms)}, {}
    )
    return {"id": question_id}


@router.patch("/form-questions/{question_id}")
async def update_form_question(request: Request, question_id: str, body: QuestionPatch) -> dict:
    user = await current_user(request)
    _catalog_editor(user)
    row = await db(request).form_questions.find_one({"_id": question_id, "institute_id": user["institute_id"]})
    if row is None:
        raise AppError(404, "not_found", "That record was not found.")
    merged = _merged_question(body, row)
    fields = _clean_question(merged)
    fields["forms"] = _forms_for(merged.place or "both")
    taken = await db(request).form_questions.find_one(
        {"institute_id": user["institute_id"], "label_key": fields["label_key"], "_id": {"$ne": question_id}}
    )
    if taken:
        raise AppError(409, "question.exists", "That question already exists.")
    if body.order is not None:
        if body.order < 1 or body.order > 50:
            raise AppError(422, "intake.field", "Use an order from 1 to 50.", field="order")
        fields["order"] = body.order
    await db(request).form_questions.update_one({"_id": question_id}, {"$set": fields})
    return {"id": question_id}


@router.delete("/form-questions/{question_id}")
async def delete_form_question(request: Request, question_id: str) -> dict:
    user = await current_user(request)
    _catalog_editor(user)
    row = await db(request).form_questions.find_one({"_id": question_id, "institute_id": user["institute_id"]})
    if row is None:
        raise AppError(404, "not_found", "That record was not found.")
    return await _retire_or_delete(request, user, "form_questions", question_id, row, "form_question.delete")


async def _question_is_used(request: Request, institute_id: str, collection: str, question_id: str) -> bool:
    database = db(request)
    if collection == "form_questions":
        checks = (
            ("enquiries", "answers.question_id"),
            ("admissions", "business_answers.question_id"),
            ("students", "business_answers.question_id"),
        )
    else:
        checks = (
            ("admissions", "course_answers.question_id"),
            ("students", "course_profiles.answers.question_id"),
        )
    for name, field in checks:
        found = await database[name].find_one({"institute_id": institute_id, field: question_id}, {"_id": 1})
        if found is not None:
            return True
    return False


async def _retire_or_delete(
    request: Request, user: dict, collection: str, question_id: str, row: dict, action: str, scope: dict | None = None
) -> dict:
    if await _question_is_used(request, user["institute_id"], collection, question_id):
        await db(request)[collection].update_one({"_id": question_id}, {"$set": {"retired": True}})
        await write_audit(request, user, action, {"label": row.get("label") or "", "retired": True}, None, scope or {})
        return {"id": question_id, "retired": True}
    await db(request)[collection].delete_one({"_id": question_id})
    await write_audit(request, user, action, {"label": row.get("label") or ""}, None, scope or {})
    return {"id": question_id}


def intake_course_visible(grants: list, course_id: str) -> bool:
    """An enquiry or admission has a course and no branch. A grant's branch does not hide it."""
    for grant in grants:
        if "student.lookup" not in (grant.get("actions") or []):
            continue
        scoped = (grant.get("scope") or {}).get("course_id")
        if scoped and scoped != course_id:
            continue
        if scoped and not course_id:
            continue
        return True
    return False


def applied_intake_course(grants: list, course_id: str) -> tuple[str, bool]:
    if not course_id:
        return "", False
    if intake_course_visible(grants, course_id):
        return course_id, False
    return "", True


def enquiry_awaiting_follow_up(row: dict, today: str) -> bool:
    if enquiry_status(row) != "open":
        return False
    return follow_up_state(row, today) in {"overdue", "today", "unscheduled"}


@router.get("/enquiries")
async def list_enquiries(
    request: Request,
    q: str = "",
    page: int = 1,
    page_size: int = 20,
    follow_up: str = "",
    course_id: str = "",
    sort: str = "",
    direction: str = "",
    stage: str = "",
    interest: str = "",
    closure_reason: str = "",
    created_from: str = "",
    created_to: str = "",
    next_from: str = "",
    next_to: str = "",
) -> dict:
    user = await current_user(request)
    _reader(user)
    return await _list_people(
        request,
        user,
        "enquiries",
        q,
        page,
        page_size,
        follow_up=follow_up,
        course_id=course_id,
        sort=sort,
        direction=direction,
        stage=stage,
        interest=interest,
        closure_reason=closure_reason,
        created_from=created_from,
        created_to=created_to,
        next_from=next_from,
        next_to=next_to,
    )


@router.get("/prospects")
async def suggest_prospects(request: Request, q: str = "") -> dict:
    user = await current_user(request)
    _reader(user)
    needle = " ".join(q.split())
    if len(needle) < 2:
        return {"items": []}
    folded = name_key(needle)
    digits = phone_key(needle)
    found = []
    cursor = db(request).prospects.find({"institute_id": user["institute_id"]})
    async for prospect in cursor:
        name = prospect.get("name_key") or ""
        phone = prospect.get("whatsapp_key") or ""
        email = prospect.get("email_key") or ""
        if folded not in name and folded not in email and (not digits or digits not in phone):
            continue
        enquiries = await db(request).enquiries.find(
            {"institute_id": user["institute_id"], "prospect_id": prospect["_id"]}
        ).to_list(length=20)
        found.append(
            {
                "id": prospect["_id"],
                "name": prospect.get("name") or "",
                "whatsapp": prospect.get("whatsapp") or "",
                "email": prospect.get("email") or "",
                "enquiries": [
                    {
                        "id": item["_id"],
                        "course_id": item.get("course_id") or "",
                        "status": enquiry_status(item),
                        "reference": item.get("reference") or "",
                    }
                    for item in enquiries
                ],
            }
        )
        if len(found) >= 8:
            break
    return {"items": found}


@router.post("/enquiries")
async def create_enquiry(
    request: Request, body: EnquiryBody, idempotency_key: str = Header(default="", alias="Idempotency-Key")
) -> dict:
    user = await current_user(request)
    _writer(user)
    replay = await _claim_operation(request, user, "enquiry.create", idempotency_key, body.model_dump())
    if replay is not None:
        return replay
    try:
        if body.course_id:
            await _course(request, user, body.course_id)
        fields = _enquiry_fields(body)
        fields["answers"] = await clean_form_answers(
            request, user["institute_id"], "enquiry", body.answers, require=False
        )
        enquiry_id = await _insert_enquiry(
            request, user, fields, prospect_id=body.prospect_id, episode_note=body.episode_note
        )
        await write_audit(request, user, "enquiry.create", None, {"name": fields["name"]}, {"enquiry_id": enquiry_id})
        from app.business import write_business

        await write_business(
            db(request),
            user,
            action="Added an enquiry",
            record=fields["name"],
            href=f"/students/enquiries/{enquiry_id}",
            permission="student.lookup",
            collection="enquiries",
            record_id=enquiry_id,
            course_id=fields.get("course_id") or "",
        )
        result = {"id": enquiry_id}
        await _finish_operation(request, user, "enquiry.create", idempotency_key, result)
        return result
    except Exception:
        await _release_operation(request, user, "enquiry.create", idempotency_key)
        raise


@router.get("/enquiries/{enquiry_id}")
async def get_enquiry(request: Request, enquiry_id: str) -> dict:
    user = await current_user(request)
    _reader(user)
    row = await _prepare_enquiry(request, user, enquiry_id)
    course = await _course_name(request, row.get("course_id") or "")
    today = await _institute_today(request, user)
    view = _enquiry_view(row, course, today)
    questions = await _form_questions(request, user["institute_id"], "enquiry")
    view["questions"] = [_question_view(item) for item in questions]
    view["answers"] = row.get("answers") or []
    view["related_count"] = await db(request).enquiries.count_documents(
        {"institute_id": user["institute_id"], "prospect_id": row.get("prospect_id") or ""}
    )
    return view


@router.patch("/enquiries/{enquiry_id}")
async def update_enquiry(request: Request, enquiry_id: str, body: EnquiryBody) -> dict:
    user = await current_user(request)
    _writer(user)
    row = await _prepare_enquiry(request, user, enquiry_id)
    if enquiry_status(row) == "converted":
        raise AppError(409, "enquiry.converted", "This enquiry is an admission. Open the admission to continue.")
    if body.course_id:
        await _course(request, user, body.course_id)
    fields = _enquiry_fields(body)
    fields["answers"] = await remember_answers(
        request, user["institute_id"], "enquiry", row.get("answers") or [], body.answers, require=False
    )
    fields["updated_at"] = _now()
    prospect_id = row.get("prospect_id") or ""
    person_keys = ("name", "name_key", "whatsapp", "whatsapp_key", "email", "email_key", "present_institution")
    person = {key: fields[key] for key in person_keys}
    if prospect_id:
        await db(request).prospects.update_one({"_id": prospect_id}, {"$set": person})
        await db(request).enquiries.update_many(
            {"institute_id": user["institute_id"], "prospect_id": prospect_id},
            {"$set": person},
        )
    specific = {key: fields[key] for key in fields if key not in person}
    specific["updated_at"] = fields["updated_at"]
    await db(request).enquiries.update_one({"_id": enquiry_id}, {"$set": specific})
    related = await db(request).enquiries.count_documents(
        {"institute_id": user["institute_id"], "prospect_id": prospect_id}
    ) if prospect_id else 1
    await write_audit(
        request,
        user,
        "enquiry.update",
        {"name": row.get("name") or ""},
        {"name": fields["name"]},
        {"enquiry_id": enquiry_id},
    )
    return {"id": enquiry_id, "related_count": related}


@router.post("/enquiries/{enquiry_id}/follow-ups")
async def add_follow_up(
    request: Request,
    enquiry_id: str,
    body: FollowUpBody,
    idempotency_key: str = Header(default="", alias="Idempotency-Key"),
) -> dict:
    user = await current_user(request)
    _writer(user)
    replay = await _claim_operation(request, user, f"follow:{enquiry_id}", idempotency_key, body.model_dump())
    if replay is not None:
        return replay
    try:
        row = await _prepare_enquiry(request, user, enquiry_id)
        _require_open_follow_up(row)
        today = await _institute_today(request, user)
        result = await _save_follow_up(request, user, row, body, today)
        await _finish_operation(request, user, f"follow:{enquiry_id}", idempotency_key, result)
        return result
    except Exception:
        await _release_operation(request, user, f"follow:{enquiry_id}", idempotency_key)
        raise


@router.delete("/enquiries/{enquiry_id}/follow-ups/{follow_id}")
async def delete_follow_up(request: Request, enquiry_id: str, follow_id: str) -> dict:
    user = await current_user(request)
    _writer(user)
    await _enquiry(request, user, enquiry_id)
    raise AppError(409, "enquiry.history", "Completed follow-ups stay on the record.")


@router.post("/enquiries/{enquiry_id}/close")
async def close_enquiry(request: Request, enquiry_id: str, body: CloseBody) -> dict:
    user = await current_user(request)
    _writer(user)
    row = await _prepare_enquiry(request, user, enquiry_id)
    if enquiry_status(row) == "converted":
        raise AppError(409, "enquiry.converted", "This enquiry is an admission. Open the admission to continue.")
    if enquiry_status(row) == "closed":
        return {"id": enquiry_id, "existing": True}
    reason = _choice(body.reason, CLOSURE_REASONS, "reason", "Choose a closure reason.")
    if not reason:
        raise AppError(422, "intake.field", "Choose a closure reason.", field="reason")
    note = _long_text(body.note, 500, "note")
    if reason == "Other" and not note:
        raise AppError(422, "intake.field", "Enter a note for Other.", field="note")
    actor = user.get("name") or ""
    event = _event("closed", f"Closed: {reason}" + (f". {note}" if note else ""), actor)
    if row.get("pending_on"):
        event = _event(
            "task_cancelled",
            f"Pending follow-up cancelled because the enquiry was closed ({reason}).",
            actor,
        )
        events = [event, _event("closed", f"Closed: {reason}" + (f". {note}" if note else ""), actor)]
    else:
        events = [event]
    await _replace_enquiry(
        request,
        row,
        {
            "status": "closed",
            "closure_reason": reason,
            "closure_note": note,
            "pending_on": "",
            "pending_task_id": "",
            "updated_at": _now(),
        },
        events,
    )
    await write_audit(request, user, "enquiry.close", None, {"reason": reason}, {"enquiry_id": enquiry_id})
    return {"id": enquiry_id}


@router.post("/enquiries/{enquiry_id}/reopen")
async def reopen_enquiry(request: Request, enquiry_id: str, body: ReopenBody) -> dict:
    user = await current_user(request)
    _writer(user)
    row = await _prepare_enquiry(request, user, enquiry_id)
    if enquiry_status(row) == "converted":
        raise AppError(409, "enquiry.converted", "Converted enquiries stay linked to the admission.")
    if enquiry_status(row) != "closed":
        raise AppError(409, "enquiry.open", "This enquiry is already open.")
    reason = _long_text(body.reason, 500, "reason")
    if not reason:
        raise AppError(422, "intake.field", "Enter a reason for reopening.", field="reason")
    actor = user.get("name") or ""
    await _replace_enquiry(
        request,
        row,
        {"status": "open", "closure_reason": "", "closure_note": "", "updated_at": _now()},
        [_event("reopened", reason, actor)],
    )
    await write_audit(request, user, "enquiry.reopen", None, {"reason": reason}, {"enquiry_id": enquiry_id})
    return {"id": enquiry_id}


@router.post("/enquiries/{enquiry_id}/convert")
async def convert_enquiry(request: Request, enquiry_id: str) -> dict:
    user = await current_user(request)
    _writer(user)
    row = await _prepare_enquiry(request, user, enquiry_id)
    if row.get("admission_id"):
        return {"id": row["admission_id"], "existing": True}
    if enquiry_status(row) == "closed":
        raise AppError(409, "enquiry.closed", "Reopen the enquiry before creating an admission.")
    if not row.get("course_id"):
        raise AppError(422, "intake.field", "Choose a course before creating an admission.", field="course_id")
    await _course(request, user, row["course_id"])
    admission_id = str(ObjectId())
    records = []
    if row.get("present_institution"):
        records.append(
            {
                "qualification": "",
                "stream": "",
                "institution": row.get("present_institution") or "",
                "mark": "",
                "year_of_passing": "",
                "board": "",
            }
        )
    phone = row.get("whatsapp") or ""
    if not phone:
        raise AppError(422, "intake.field", "Enter a WhatsApp number before admitting this enquiry.", field="whatsapp")
    reference = await _allocate_admission_reference(request, user["institute_id"])
    actor = user.get("name") or ""
    enquiry_answers = await clean_form_answers(
        request, user["institute_id"], "enquiry", row.get("answers") or [], require=True
    )
    admission_questions = await _form_questions(request, user["institute_id"], "admission")
    shared_ids = {item["_id"] for item in admission_questions}
    shared_answers = [item for item in enquiry_answers if item["question_id"] in shared_ids]
    await db(request).admissions.insert_one(
        {
            "_id": admission_id,
            "institute_id": user["institute_id"],
            "reference": reference,
            "prospect_id": row.get("prospect_id") or "",
            "course_id": row.get("course_id") or "",
            "name": row.get("name") or "",
            "name_key": row.get("name_key") or name_key(row.get("name") or ""),
            "phone": phone,
            "phone_key": phone_key(phone),
            "date_of_birth": "",
            "gender": "",
            "blood_group": "",
            "email": row.get("email") or "",
            "events": [_event("created", "Created from the enquiry.", actor)],
            "street": "",
            "address_line_2": "",
            "place": "",
            "district": "",
            "state": "",
            "pin": "",
            "guardian_name": "",
            "guardian_phone": "",
            "academic_records": records,
            "course_answers": [],
            "business_answers": shared_answers,
            "status": "open",
            "enquiry_id": enquiry_id,
            "created_at": _now(),
            "updated_at": _now(),
        }
    )
    actor = user.get("name") or ""
    events = [_event("converted", "Created the admission and cancelled any pending follow-up.", actor)]
    await _replace_enquiry(
        request,
        row,
        {
            "admission_id": admission_id,
            "status": "converted",
            "pending_on": "",
            "pending_task_id": "",
            "updated_at": _now(),
        },
        events,
    )
    await write_audit(
        request,
        user,
        "enquiry.convert",
        None,
        {"name": row.get("name") or "", "admission_id": admission_id},
        {"enquiry_id": enquiry_id},
    )
    from app.business import write_business

    await write_business(
        db(request),
        user,
        action="Created an admission",
        record=row.get("name") or "",
        href=f"/students/admissions/{admission_id}",
        permission="student.lookup",
        collection="admissions",
        record_id=admission_id,
        course_id=row.get("course_id") or "",
    )
    return {"id": admission_id, "existing": False}


@router.post("/enquiries/import")
async def preview_enquiries(request: Request, file: UploadFile = File(...)) -> dict:
    user = await current_user(request)
    _writer(user)
    rows = _sheet_rows(await _workbook_bytes(file))
    courses = await db(request).courses.find({"institute_id": user["institute_id"]}).to_list(length=200)
    by_name = {item.get("name_key"): item for item in courses}
    form_questions = await _form_questions(request, user["institute_id"], "enquiry")
    preview = []
    for offset, cells in rows:
        mapped, unknown = _map_headers(cells, _ENQUIRY_HEADERS)
        extras, leftover = _question_columns(cells, unknown, form_questions)
        item = {"row": offset, "errors": [], "values": {}, "answers": []}
        if leftover:
            item["errors"].append(f"Unknown column {leftover[0]}.")
        try:
            course_id = ""
            course_label = mapped.get("course") or ""
            if course_label:
                course = by_name.get(name_key(course_label))
                if course is None:
                    raise AppError(422, "intake.field", f"No course named {course_label}.", field="course_id")
                course_id = course["_id"]
            body = EnquiryBody(
                name=mapped.get("name") or "",
                whatsapp=mapped.get("whatsapp") or "",
                present_institution=mapped.get("present_institution") or "",
                course_id=course_id,
                interest=mapped.get("interest") or "",
                references=mapped.get("references") or "",
                details=mapped.get("details") or "",
                answers=extras,
            )
            fields = _enquiry_fields(body)
            answers = _clean_answers(form_questions, extras, require=False)
            item["answers"] = answers
            item["values"] = {
                "name": fields["name"],
                "whatsapp": fields["whatsapp"],
                "present_institution": fields["present_institution"],
                "course_id": course_id,
                "course_name": course_label,
                "interest": fields["interest"],
                "references": fields["references"],
                "details": fields["details"],
            }
        except AppError as exc:
            item["errors"].append(exc.message)
        preview.append(item)
    return {"rows": preview}


@router.post("/enquiries/import/commit")
async def commit_enquiries(request: Request, body: EnquiryCommit) -> dict:
    user = await current_user(request)
    _writer(user)
    if not body.rows:
        raise AppError(422, "intake.field", "There are no rows to save.")
    created = []
    for item in body.rows:
        if item.course_id:
            await _course(request, user, item.course_id)
        fields = _enquiry_fields(item)
        fields["answers"] = await clean_form_answers(
            request, user["institute_id"], "enquiry", item.answers, require=False
        )
        enquiry_id = await _insert_enquiry(request, user, fields)
        created.append(enquiry_id)
    await write_audit(
        request, user, "enquiry.import", None, {"count": len(created)}, {"institute_id": user["institute_id"]}
    )
    return {"created": len(created), "ids": created}


async def _import_scope(request: Request, user: dict, branch_id: str, course_id: str) -> None:
    branch = await db(request).branches.find_one({"_id": branch_id, "institute_id": user["institute_id"]})
    course = await db(request).courses.find_one({"_id": course_id, "institute_id": user["institute_id"]})
    offering = await db(request).offerings.find_one(
        {"institute_id": user["institute_id"], "branch_id": branch_id, "course_id": course_id}
    )
    if branch is None or course is None or offering is None:
        raise AppError(422, "import.scope", "Choose a branch and a course that branch offers.")


def _import_summary(row: dict) -> dict:
    created = (row.get("result") or {}).get("created") or []
    return {
        "id": row["_id"],
        "filename": row.get("filename") or "",
        "status": row.get("status") or "open",
        "created_at": row["created_at"].isoformat() if row.get("created_at") is not None else None,
        "branch_id": row.get("branch_id") or "",
        "course_id": row.get("course_id") or "",
        "batch_id": row.get("batch_id") or "",
        "created_count": len(created) if isinstance(created, list) else 0,
    }


@router.get("/enquiry-imports")
async def list_enquiry_imports(request: Request) -> dict:
    user = await current_user(request)
    _reader(user)
    rows = await db(request).enquiry_imports.find(
        {"institute_id": user["institute_id"]}, {"content": 0}
    ).sort("created_at", -1).to_list(200)
    return {"items": [_import_summary(row) for row in rows]}


@router.get("/enquiry-imports/template")
async def enquiry_import_template(request: Request) -> Response:
    from app.admissions.enquiry_import import template_workbook

    user = await current_user(request)
    _reader(user)
    return Response(
        content=template_workbook(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=enquiry-template.xlsx"},
    )


@router.post("/enquiry-imports")
async def create_enquiry_import(
    request: Request,
    file: UploadFile = File(...),
    branch_id: str = Form(...),
    course_id: str = Form(...),
) -> dict:
    from app.admissions.enquiry_import import store_upload

    user = await current_user(request)
    _writer(user)
    await _import_scope(request, user, branch_id, course_id)
    data = await file.read()
    return await store_upload(request, user, file.filename or "import", data, branch_id, course_id)


@router.delete("/enquiry-imports/{operation_id}")
async def delete_enquiry_import(request: Request, operation_id: str) -> dict:
    from app.admissions.enquiry_import import _operation

    user = await current_user(request)
    _writer(user)
    row = await _operation(request, user, operation_id)
    if row.get("status") == "committed":
        raise AppError(422, "import.committed", "A saved import cannot be deleted.")
    await db(request).enquiry_imports.delete_one({"_id": row["_id"], "institute_id": user["institute_id"]})
    await write_audit(
        request,
        user,
        "enquiry.import.delete",
        {"filename": row.get("filename") or ""},
        None,
        {"institute_id": user["institute_id"]},
    )
    return {"id": row["_id"]}


@router.get("/enquiry-imports/{operation_id}")
async def read_enquiry_import(request: Request, operation_id: str) -> dict:
    from app.admissions.enquiry_import import _operation, present_operation

    user = await current_user(request)
    _reader(user)
    return await present_operation(request, user, await _operation(request, user, operation_id))


@router.put("/enquiry-imports/{operation_id}")
async def plan_enquiry_import(request: Request, operation_id: str, body: ImportPlanBody) -> dict:
    from app.admissions.enquiry_import import save_plan

    user = await current_user(request)
    _writer(user)
    return await save_plan(request, user, operation_id, body.model_dump())


@router.post("/enquiry-imports/{operation_id}/commit")
async def commit_enquiry_import(request: Request, operation_id: str, body: ImportCommitBody) -> dict:
    from app.admissions.enquiry_import import commit_saved_plan

    user = await current_user(request)
    _writer(user)
    return await commit_saved_plan(request, user, operation_id, body.revision, body.idempotency_key)


@router.get("/enquiry-imports/{operation_id}/errors")
async def enquiry_import_errors(request: Request, operation_id: str) -> Response:
    from app.admissions.enquiry_import import error_report

    user = await current_user(request)
    _reader(user)
    report = await error_report(request, user, operation_id)
    return Response(
        content=report,
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=enquiry-import-errors.csv"},
    )


@router.post("/enquiry-mappings")
async def save_enquiry_mapping(request: Request, body: MappingSaveBody) -> dict:
    from app.admissions.enquiry_import import remember_mapping

    user = await current_user(request)
    _writer(user)
    return await remember_mapping(request, user, body.import_id, body.name)


@router.post("/enquiry-mappings/{mapping_id}/publish")
async def share_enquiry_mapping(request: Request, mapping_id: str) -> dict:
    from app.admissions.enquiry_import import publish_mapping

    user = await current_user(request)
    _writer(user)
    return await publish_mapping(request, user, mapping_id)


@router.post("/enquiry-imports/{operation_id}/mapping")
async def use_enquiry_mapping(request: Request, operation_id: str, body: MappingApplyBody) -> dict:
    from app.admissions.enquiry_import import apply_mapping

    user = await current_user(request)
    _writer(user)
    return await apply_mapping(request, user, operation_id, body.mapping_id)


@router.get("/admission-imports")
async def list_admission_imports(request: Request) -> dict:
    user = await current_user(request)
    _reader(user)
    rows = await db(request).admission_imports.find(
        {"institute_id": user["institute_id"]}, {"content": 0}
    ).sort("created_at", -1).to_list(200)
    return {"items": [_import_summary(row) for row in rows]}


@router.get("/admission-imports/template")
async def admission_import_template(request: Request) -> Response:
    from app.admissions.admission_import import template_workbook

    user = await current_user(request)
    _reader(user)
    return Response(
        content=template_workbook(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=admission-template.xlsx"},
    )


@router.post("/admission-imports")
async def create_admission_import(
    request: Request,
    file: UploadFile = File(...),
    branch_id: str = Form(...),
    course_id: str = Form(...),
    batch_id: str = Form(...),
) -> dict:
    from app.admissions.admission_import import store_upload

    user = await current_user(request)
    _writer(user)
    await _import_scope(request, user, branch_id, course_id)
    await _eligible_batch(request, user, {"course_id": course_id}, branch_id, batch_id)
    return await store_upload(
        request, user, file.filename or "import", await file.read(), branch_id, course_id, batch_id
    )


@router.delete("/admission-imports/{operation_id}")
async def delete_admission_import(request: Request, operation_id: str) -> dict:
    from app.admissions.admission_import import _operation

    user = await current_user(request)
    _writer(user)
    row = await _operation(request, user, operation_id)
    if row.get("status") == "committed":
        raise AppError(422, "import.committed", "A saved import cannot be deleted.")
    await db(request).admission_imports.delete_one({"_id": row["_id"], "institute_id": user["institute_id"]})
    await write_audit(
        request,
        user,
        "admission.import.delete",
        {"filename": row.get("filename") or ""},
        None,
        {"institute_id": user["institute_id"]},
    )
    return {"id": row["_id"]}


@router.get("/admission-imports/{operation_id}")
async def read_admission_import(request: Request, operation_id: str) -> dict:
    from app.admissions.admission_import import _operation, present_operation

    user = await current_user(request)
    _reader(user)
    return await present_operation(request, user, await _operation(request, user, operation_id))


@router.put("/admission-imports/{operation_id}")
async def plan_admission_import(request: Request, operation_id: str, body: ImportPlanBody) -> dict:
    from app.admissions.admission_import import save_plan

    user = await current_user(request)
    _writer(user)
    return await save_plan(request, user, operation_id, body.model_dump())


@router.post("/admission-imports/{operation_id}/commit")
async def commit_admission_import(request: Request, operation_id: str, body: ImportCommitBody) -> dict:
    from app.admissions.admission_import import commit_saved_plan

    user = await current_user(request)
    _writer(user)
    return await commit_saved_plan(request, user, operation_id, body.revision, body.idempotency_key)


@router.get("/admission-imports/{operation_id}/errors")
async def admission_import_errors(request: Request, operation_id: str) -> Response:
    from app.admissions.admission_import import error_report

    user = await current_user(request)
    _reader(user)
    report = await error_report(request, user, operation_id)
    return Response(
        content=report,
        media_type="text/csv",
        headers={"Content-Disposition": "attachment; filename=admission-import-errors.csv"},
    )


@router.post("/admission-mappings")
async def save_admission_mapping(request: Request, body: MappingSaveBody) -> dict:
    from app.admissions.admission_import import remember_mapping

    user = await current_user(request)
    _writer(user)
    return await remember_mapping(request, user, body.import_id, body.name)


@router.post("/admission-mappings/{mapping_id}/publish")
async def share_admission_mapping(request: Request, mapping_id: str) -> dict:
    from app.admissions.admission_import import publish_mapping

    user = await current_user(request)
    _writer(user)
    return await publish_mapping(request, user, mapping_id)


@router.post("/admission-imports/{operation_id}/mapping")
async def use_admission_mapping(request: Request, operation_id: str, body: MappingApplyBody) -> dict:
    from app.admissions.admission_import import apply_mapping

    user = await current_user(request)
    _writer(user)
    return await apply_mapping(request, user, operation_id, body.mapping_id)


@router.get("/admissions")
async def list_admissions(
    request: Request,
    q: str = "",
    page: int = 1,
    page_size: int = 20,
    status: str = "",
    course_id: str = "",
    readiness: str = "",
    sort: str = "",
    direction: str = "",
    created_from: str = "",
    created_to: str = "",
    branch_id: str = "",
    batch_id: str = "",
) -> dict:
    user = await current_user(request)
    _reader(user)
    return await _list_people(
        request,
        user,
        "admissions",
        q,
        page,
        page_size,
        status=status,
        course_id=course_id,
        readiness=readiness,
        sort=sort,
        direction=direction,
        created_from=created_from,
        created_to=created_to,
        branch_id=branch_id,
        batch_id=batch_id,
    )


@router.post("/admissions")
async def create_admission(
    request: Request, body: AdmissionBody, idempotency_key: str = Header(default="", alias="Idempotency-Key")
) -> dict:
    user = await current_user(request)
    _writer(user)
    replay = await _claim_operation(request, user, "admission.create", idempotency_key, body.model_dump())
    if replay:
        return replay
    try:
        result = await _save_admission(request, user, None, body)
    except Exception:
        await _release_operation(request, user, "admission.create", idempotency_key)
        raise
    await _finish_operation(request, user, "admission.create", idempotency_key, result)
    return result


@router.get("/admissions/{admission_id}")
async def get_admission(request: Request, admission_id: str) -> dict:
    user = await current_user(request)
    _reader(user)
    row = await _ensure_admission_reference(request, user, await _admission(request, user, admission_id))
    available = True
    try:
        questions = await _questions_for(request, row.get("course_id") or "")
        business = await _form_questions(request, user["institute_id"], "admission")
    except Exception:
        questions = []
        business = []
        available = False
    matches = await find_matches(
        request, user["institute_id"], row.get("name") or "", row.get("phone") or "", row if available else None
    )
    context = await _selection_context(request, row)
    view = _admission_view(
        row,
        await _course_name(request, row.get("course_id") or ""),
        matches,
        questions,
        business,
        branch_name=context["branch_name"],
        batch_name=context["batch_name"],
        batch_open=context["batch_open"],
        requirements_available=available,
    )
    view["prospect"] = await _prospect_view(request, user, row.get("prospect_id") or "")
    view["repeats"] = await _open_repeats(request, user, row)
    view["batches"] = await _eligible_batch_choices(request, user, row)
    return view


@router.patch("/admissions/{admission_id}")
async def update_admission(request: Request, admission_id: str, body: AdmissionBody) -> dict:
    user = await current_user(request)
    _writer(user)
    row = await _admission(request, user, admission_id)
    _expect_current(row, body.expected_updated_at)
    if row.get("status") == "converted":
        raise AppError(409, "admission.converted", "This admission is already a student. Edit the student instead.")
    if row.get("status") == "closed":
        raise AppError(409, "admission.closed", "Reopen the application before editing it.")
    return await _save_admission(request, user, row, body)


@router.get("/admissions/{admission_id}/course-change")
async def preview_course_change(request: Request, admission_id: str, course_id: str = "") -> dict:
    user = await current_user(request)
    _reader(user)
    row = await _admission(request, user, admission_id)
    return await _course_change_summary(request, user, row, course_id)


@router.post("/admissions/{admission_id}/correction")
async def correct_admission(request: Request, admission_id: str, body: CorrectionBody) -> dict:
    user = await current_user(request)
    _writer(user)
    row = await _admission(request, user, admission_id)
    _expect_current(row, body.expected_updated_at)
    if row.get("status") != "converted":
        raise AppError(409, "admission.open", "Source corrections are for enrolled applications. Edit this application directly.")
    note = _long_text(body.note, 500, "note")
    if not note:
        raise AppError(422, "intake.field", "Enter what this correction changes.", field="note")
    saved = await _save_admission(request, user, row, body, correction_note=note)
    return saved


@router.post("/admissions/{admission_id}/identity")
async def confirm_admission_identity(request: Request, admission_id: str, body: IdentityBody) -> dict:
    user = await current_user(request)
    _writer(user)
    row = await _admission(request, user, admission_id)
    _expect_current(row, body.expected_updated_at)
    _require_open_admission(row)
    questions = await _questions_for(request, row.get("course_id") or "")
    business = await _form_questions(request, user["institute_id"], "admission")
    if admission_requirements(row, questions, business):
        raise AppError(422, "admission.requirements", "Complete the application before confirming identity.")
    decision = body.decision
    if decision not in {"create", "attach"}:
        raise AppError(422, "admission.identity", "Choose an existing student or a new student.", field="decision")
    matches = await find_matches(request, user["institute_id"], row.get("name") or "", row.get("phone") or "")
    student_id = ""
    if decision == "attach":
        student_id = body.student_id
        if student_id not in {item["id"] for item in matches}:
            raise AppError(422, "admission.match", "Choose one of the matching students.")
    note = "Confirmed a new student." if decision == "create" else "Confirmed an existing student."
    await db(request).admissions.update_one(
        {"_id": admission_id},
        {
            "$set": {
                "match_decision": decision,
                "match_student_id": student_id,
                "updated_at": _now(),
            },
            "$push": {"events": _event("identity", note, user.get("name") or "")},
        },
    )
    await write_audit(
        request, user, "admission.identity", None, {"decision": decision}, {"admission_id": admission_id}
    )
    return {"id": admission_id, "decision": decision, "student_id": student_id}


@router.post("/admissions/{admission_id}/selection")
async def save_admission_selection(request: Request, admission_id: str, body: SelectionBody) -> dict:
    user = await current_user(request)
    _writer(user)
    row = await _admission(request, user, admission_id)
    _expect_current(row, body.expected_updated_at)
    _require_open_admission(row)
    if not identity_resolved(row):
        raise AppError(409, "admission.identity", "Confirm identity before choosing a batch.")
    batch = await _eligible_batch(request, user, row, body.branch_id, body.batch_id)
    await db(request).admissions.update_one(
        {"_id": admission_id},
        {
            "$set": {"branch_id": body.branch_id, "batch_id": body.batch_id, "updated_at": _now()},
            "$push": {"events": _event("selection", f"Selected {batch.get('name') or 'a batch'}.", user.get("name") or "")},
        },
    )
    return {"id": admission_id, "branch_id": body.branch_id, "batch_id": body.batch_id, "batch_name": batch.get("name") or ""}


@router.post("/admissions/{admission_id}/close")
async def close_admission(request: Request, admission_id: str, body: CloseBody) -> dict:
    user = await current_user(request)
    _writer(user)
    row = await _admission(request, user, admission_id)
    _expect_current(row, body.expected_updated_at)
    if row.get("status") == "converted":
        raise AppError(409, "admission.enrolled", "Enrolled applications stay with the student record.")
    if row.get("status") == "closed":
        return {"id": admission_id, "existing": True}
    reason = _choice(body.reason, CLOSURE_REASONS, "reason", "Choose a closure reason.")
    if not reason:
        raise AppError(422, "intake.field", "Choose a closure reason.", field="reason")
    note = _long_text(body.note, 500, "note")
    if reason == "Other" and not note:
        raise AppError(422, "intake.field", "Enter a note for Other.", field="note")
    actor = user.get("name") or ""
    await db(request).admissions.update_one(
        {"_id": admission_id},
        {
            "$set": {
                "status": "closed",
                "closure_reason": reason,
                "closure_note": note,
                "updated_at": _now(),
            },
            "$push": {"events": _event("closed", f"Closed: {reason}" + (f". {note}" if note else ""), actor)},
        },
    )
    await write_audit(request, user, "admission.close", None, {"reason": reason}, {"admission_id": admission_id})
    return {"id": admission_id}


@router.post("/admissions/{admission_id}/reopen")
async def reopen_admission(request: Request, admission_id: str, body: ReopenBody) -> dict:
    user = await current_user(request)
    _writer(user)
    row = await _admission(request, user, admission_id)
    _expect_current(row, body.expected_updated_at)
    if row.get("status") == "converted":
        raise AppError(409, "admission.enrolled", "Enrolled applications stay with the student record.")
    if row.get("status") != "closed":
        raise AppError(409, "admission.open", "This application is already open.")
    reason = _long_text(body.reason, 500, "reason")
    if not reason:
        raise AppError(422, "intake.field", "Enter a reason for reopening.", field="reason")
    actor = user.get("name") or ""
    await db(request).admissions.update_one(
        {"_id": admission_id},
        {
            "$set": {"status": "open", "closure_reason": "", "closure_note": "", "updated_at": _now()},
            "$push": {"events": _event("reopened", reason, actor)},
        },
    )
    await write_audit(request, user, "admission.reopen", None, {"reason": reason}, {"admission_id": admission_id})
    return {"id": admission_id}


@router.post("/admissions/{admission_id}/convert")
async def convert_admission(
    request: Request,
    admission_id: str,
    body: ConvertBody,
    idempotency_key: str = Header(default="", alias="Idempotency-Key"),
) -> dict:
    user = await current_user(request)
    _writer(user)
    replay = await _claim_operation(
        request, user, "admission.convert", idempotency_key, {"id": admission_id, **body.model_dump()}
    )
    if replay:
        return replay
    try:
        result = await _convert_admission(request, user, admission_id, body)
    except Exception:
        await _release_operation(request, user, "admission.convert", idempotency_key)
        raise
    await _finish_operation(request, user, "admission.convert", idempotency_key, result)
    return result


async def _convert_admission(request: Request, user: dict, admission_id: str, body: ConvertBody) -> dict:
    row = await _admission(request, user, admission_id)
    _expect_current(row, body.expected_updated_at)
    if row.get("status") == "converted":
        return {"id": row.get("student_id"), "existing": True, "enrollment_id": row.get("enrollment_id")}
    if row.get("status") == "closed":
        raise AppError(409, "admission.closed", "Reopen the application before enrolling.")
    if not allows(
        user["grants"],
        "student.manage",
        {
            "institute_id": user["institute_id"],
            "branch_id": body.branch_id,
            "course_id": row.get("course_id"),
            "batch_id": body.batch_id,
        },
    ):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")
    batch = await db(request).batches.find_one({"_id": body.batch_id, "institute_id": user["institute_id"]})
    if batch is None or batch.get("branch_id") != body.branch_id or batch.get("course_id") != row.get("course_id"):
        raise AppError(422, "enrollment.context", "Choose a branch and batch for this course.")
    require_open_batch(batch)
    questions = await _questions_for(request, row.get("course_id") or "")
    answers = _clean_answers(questions, row.get("course_answers") or [], require=True)
    business = await clean_form_answers(
        request, user["institute_id"], "admission", row.get("business_answers") or [], require=True
    )
    row["business_answers"] = business
    if not (row.get("phone") or "").strip():
        raise AppError(422, "intake.field", "Enter the contact number.", field="phone")
    matches = await find_matches(request, user["institute_id"], row.get("name") or "", row.get("phone") or "")
    decision = body.decision or row.get("match_decision") or ""
    student_id = body.student_id or row.get("match_student_id") or ""
    if decision == "attach":
        if student_id not in {item["id"] for item in matches}:
            raise AppError(422, "admission.match", "Choose one of the matching students.")
    elif decision == "create":
        student_id = ""
    elif matches:
        raise AppError(409, "admission.match", "Confirm which student this admission belongs to.", matches=matches)
    else:
        raise AppError(409, "admission.identity", "Confirm that this applicant is a new student.")
    result = await _commit_student(
        request, user, row, batch, answers, student_id or None, admission_id, decision or "create"
    )
    await write_audit(
        request,
        user,
        "admission.convert",
        None,
        {"name": row.get("name") or "", "student_id": result["id"], "created": result["created"]},
        {"student_id": result["id"], "branch_id": body.branch_id, "batch_id": body.batch_id},
    )
    return result


@router.post("/admissions/import")
async def preview_admissions(request: Request, file: UploadFile = File(...)) -> dict:
    user = await current_user(request)
    _writer(user)
    rows = _sheet_rows(await _workbook_bytes(file))
    courses = await db(request).courses.find({"institute_id": user["institute_id"]}).to_list(length=200)
    by_name = {item.get("name_key"): item for item in courses}
    question_rows = await db(request).course_questions.find({"institute_id": user["institute_id"]}).to_list(length=500)
    business_questions = await _form_questions(request, user["institute_id"], "admission")
    preview = []
    for offset, cells in rows:
        mapped, unknown = _map_headers(cells, _ADMISSION_HEADERS)
        item = {"row": offset, "errors": [], "matches": [], "values": {}, "answers": [], "business_answers": []}
        course = by_name.get(name_key(mapped.get("course") or ""))
        questions = [question for question in question_rows if course and question.get("course_id") == course["_id"]]
        extras, leftover = _question_columns(cells, unknown, questions)
        business_extras, leftover = _question_columns(cells, leftover, business_questions)
        if leftover:
            item["errors"].append(f"Unknown column {leftover[0]}.")
        try:
            if course is None:
                label = mapped.get("course") or ""
                message = f"No course named {label}." if label else "Choose a course."
                raise AppError(422, "intake.field", message, field="course_id")
            body = _admission_from_map(mapped, course["_id"], extras, business_extras)
            fields = _admission_contact(body)
            answers = _clean_answers(questions, body.course_answers, require=False)
            business = _clean_answers(business_questions, body.business_answers, require=False)
            matches = await find_matches(request, user["institute_id"], fields["name"], fields["phone"])
            item["matches"] = matches
            item["proposed_action"] = "create_application"
            item["answers"] = answers
            item["business_answers"] = business
            item["values"] = {**fields, "course_id": course["_id"], "course_name": course.get("name") or ""}
            item["values"].pop("name_key", None)
            item["values"].pop("phone_key", None)
        except AppError as exc:
            item["errors"].append(exc.message)
        preview.append(item)
    return {"rows": preview}


@router.post("/admissions/import/commit")
async def commit_admissions(request: Request, body: AdmissionCommit) -> dict:
    user = await current_user(request)
    _writer(user)
    if not body.rows:
        raise AppError(422, "intake.field", "There are no rows to save.")
    created = []
    for item in body.rows:
        matches = await find_matches(request, user["institute_id"], item.name, item.phone)
        decision = item.decision
        if matches and decision not in {"attach", "create"}:
            raise AppError(
                409, "admission.match", f"Confirm whether {item.name} is an existing student.", matches=matches
            )
        if decision == "attach" and item.student_id not in {match["id"] for match in matches}:
            raise AppError(422, "admission.match", f"Choose a matching student for {item.name}.")
        if decision not in {"attach", "create"}:
            decision = ""
        saved = await _save_admission(request, user, None, item, decision=decision, student_id=item.student_id)
        created.append(saved["id"])
    await write_audit(
        request, user, "admission.import", None, {"count": len(created)}, {"institute_id": user["institute_id"]}
    )
    return {"created": len(created), "ids": created}


async def _save_admission(
    request: Request,
    user: dict,
    existing: dict | None,
    body: AdmissionBody,
    decision: str = "",
    student_id: str = "",
    session=None,
    quiet: bool = False,
    source: dict | None = None,
    correction_note: str = "",
) -> dict:
    await _course(request, user, body.course_id)
    if existing is not None and (existing.get("course_id") or "") != body.course_id:
        summary = await _course_change_summary(request, user, existing, body.course_id)
        if summary["needs_confirmation"] and not body.confirm_course_change:
            raise AppError(409, "admission.course", summary["message"], field="course_id")
    fields = _admission_contact(body)
    questions = await _questions_for(request, body.course_id)
    cleaned_course = _clean_answers(questions, body.course_answers, require=False)
    current_course = {item["_id"] for item in questions}
    kept_course = [
        item
        for item in (existing or {}).get("course_answers") or []
        if item.get("question_id") not in current_course
    ]
    fields["course_answers"] = kept_course + cleaned_course
    fields["business_answers"] = await remember_answers(
        request,
        user["institute_id"],
        "admission",
        (existing or {}).get("business_answers") or [],
        body.business_answers,
        require=False,
    )
    fields["course_id"] = body.course_id
    fields["updated_at"] = _now()
    if body.prospect_id:
        fields["prospect_id"] = body.prospect_id
    elif existing is not None:
        fields["prospect_id"] = existing.get("prospect_id") or ""
    if body.episode_note.strip():
        fields["episode_note"] = _long_text(body.episode_note, 500, "episode_note")
    elif existing is not None:
        fields["episode_note"] = existing.get("episode_note") or ""
    if source:
        fields["source"] = {"operation_id": source.get("operation_id") or "", "sheet": source.get("sheet") or "", "row": source.get("row") or 0}
    actor = user.get("name") or ""
    if decision:
        fields["match_decision"] = decision
        fields["match_student_id"] = student_id if decision == "attach" else ""
    if existing is not None and not correction_note:
        if existing.get("name_key") != fields["name_key"] or existing.get("phone_key") != fields["phone_key"]:
            fields["match_decision"] = ""
            fields["match_student_id"] = ""
        if (existing.get("course_id") or "") != body.course_id:
            fields["branch_id"] = ""
            fields["batch_id"] = ""
    identity_changed = existing is None or (
        existing.get("name_key") != fields["name_key"]
        or existing.get("phone_key") != fields["phone_key"]
        or (existing.get("course_id") or "") != body.course_id
    )
    if identity_changed:
        await _guard_admission_repeat(request, user, fields, existing["_id"] if existing else "", session=session)
    if existing is None:
        admission_id = str(ObjectId())
        note = "Created the application."
        if fields.get("episode_note"):
            note = f"Created the application. {fields['episode_note']}"
        await db(request).admissions.insert_one(
            {
                "_id": admission_id,
                "institute_id": user["institute_id"],
                "status": "open",
                "reference": await _allocate_admission_reference(request, user["institute_id"], session=session),
                "created_at": _now(),
                "events": [_event("created", note, actor)],
                **fields,
            },
            session=session,
        )
        if not quiet:
            await write_audit(
                request, user, "admission.create", None, {"name": fields["name"]}, {"admission_id": admission_id}
            )
            from app.business import write_business

            await write_business(
                db(request),
                user,
                action="Created an admission",
                record=fields["name"],
                href=f"/students/admissions/{admission_id}",
                permission="student.lookup",
                collection="admissions",
                record_id=admission_id,
                course_id=fields.get("course_id") or "",
            )
        return {"id": admission_id}
    event = (
        _event("correction", correction_note, actor)
        if correction_note
        else _event("updated", "Updated the application.", actor)
    )
    await db(request).admissions.update_one(
        {"_id": existing["_id"]},
        {"$set": fields, "$push": {"events": event}},
        session=session,
    )
    if not quiet:
        await write_audit(
            request,
            user,
            "admission.correct" if correction_note else "admission.update",
            {"name": existing.get("name") or ""},
            {"name": fields["name"], "note": correction_note},
            {"admission_id": existing["_id"]},
        )
    return {"id": existing["_id"]}


async def _commit_student(
    request: Request,
    user: dict,
    admission: dict,
    batch: dict,
    answers: list[dict],
    student_id: str | None,
    admission_id: str,
    decision: str,
    session=None,
    announce: bool = True,
) -> dict:
    institute_id = user["institute_id"]
    profile = _profile_from_admission(admission)
    records = admission.get("academic_records") or []
    business = admission.get("business_answers") or []

    async def enroll(active) -> dict:
        created_now = student_id is None
        current_id = student_id
        if created_now:
            code = await allocate_student_code(db(request), active, "IAM")
            current_id = str(ObjectId())
            document = {
                "_id": current_id,
                "institute_id": institute_id,
                "student_code": code,
                "display_name": admission.get("name") or "",
                "name_key": admission.get("name_key") or name_key(admission.get("name") or ""),
                "edit_version": 0,
                "academic_records": records,
                "course_profiles": [{"course_id": admission.get("course_id"), "answers": answers}] if answers else [],
                "business_answers": business,
                **{key: value for key, value in profile.items() if value},
                **{key: value for key, value in _mirror_legacy(records).items() if value},
            }
            if profile.get("phone"):
                document["phone_key"] = phone_key(profile["phone"])
            await db(request).students.insert_one(document, session=active)
        else:
            student = await db(request).students.find_one({"_id": current_id}, session=active)
            if student is None:
                raise AppError(404, "not_found", "That record was not found.")
        enrollment_id = await _ensure_enrollment(
            request, active, institute_id, current_id, batch, admission.get("name_key") or ""
        )
        await db(request).admissions.update_one(
            {"_id": admission_id},
            {
                "$set": {
                    "status": "converted",
                    "student_id": current_id,
                    "enrollment_id": enrollment_id,
                    "branch_id": batch.get("branch_id"),
                    "batch_id": batch["_id"],
                    "updated_at": _now(),
                    "match_decision": decision or "create",
                    "match_student_id": "" if created_now else current_id,
                },
                "$push": {
                    "events": _event(
                        "enrolled",
                        f"Enrolled in {batch.get('name') or 'the batch'}.",
                        user.get("name") or "",
                    )
                },
            },
            session=active,
        )
        return {"id": current_id, "created": created_now, "enrollment_id": enrollment_id}

    try:
        if session is not None:
            result = await enroll(session)
        else:
            async with await request.app.state.client.start_session() as owned:
                async with owned.start_transaction():
                    result = await enroll(owned)
    except DuplicateKeyError as exc:
        raise AppError(
            409,
            "student.identity",
            "Another student in that batch already uses this name. Add the course to that student instead.",
        ) from exc
    if announce:
        from app.business import write_business

        await write_business(
            db(request),
            user,
            action="Enrolled a student",
            record=admission.get("name") or "",
            href=f"/students/{result['id']}",
            permission="student.lookup",
            collection="students",
            record_id=result["id"],
            course_id=admission.get("course_id") or "",
            branch_id=batch.get("branch_id") or "",
            batch_id=batch["_id"],
        )
    return result


def _profile_from_admission(admission: dict) -> dict:
    return {
        "date_of_birth": admission.get("date_of_birth") or "",
        "gender": admission.get("gender") or "",
        "blood_group": admission.get("blood_group") or "",
        "phone": admission.get("phone") or "",
        "email": admission.get("email") or "",
        "street": admission.get("street") or "",
        "address_line_2": admission.get("address_line_2") or "",
        "place": admission.get("place") or "",
        "district": admission.get("district") or "",
        "state": admission.get("state") or "",
        "pin": admission.get("pin") or "",
        "guardian_name": admission.get("guardian_name") or "",
        "guardian_phone": admission.get("guardian_phone") or "",
    }


async def _ensure_enrollment(
    request: Request, session, institute_id: str, student_id: str, batch: dict, key: str
) -> str:
    batch_id = batch["_id"]
    already = await db(request).enrollments.find_one({"student_id": student_id, "batch_id": batch_id}, session=session)
    if already:
        if not already.get("enrollment_code"):
            await assign_enrollment_code(db(request), already, session)
        return already["_id"]
    require_open_batch(batch)
    if key and await name_conflicts(db(request), batch_id, [key], student_id, session):
        raise AppError(
            409,
            "student.identity",
            "Another student in that batch already uses this name. Add the course to that student instead.",
        )
    enrollment_id = str(ObjectId())
    await db(request).enrollments.insert_one(
        enrollment_document(
            enrollment_id=enrollment_id,
            institute_id=institute_id,
            student_id=student_id,
            branch_id=batch.get("branch_id"),
            course_id=batch.get("course_id"),
            batch_id=batch_id,
            enrollment_code=await allocate_enrollment_code(db(request), session),
        ),
        session=session,
    )
    if key:
        own = await db(request).enrollment_names.find_one(
            {"batch_id": batch_id, "name_key": key, "student_id": student_id}, session=session
        )
        if own is None:
            await db(request).enrollment_names.insert_one(
                name_document(
                    institute_id=institute_id, batch_id=batch_id, student_id=student_id, name_key=key
                ),
                session=session,
            )
    return enrollment_id


async def _enquiry(request: Request, user: dict, enquiry_id: str) -> dict:
    row = await db(request).enquiries.find_one({"_id": enquiry_id, "institute_id": user["institute_id"]})
    if row is None:
        raise AppError(404, "not_found", "That record was not found.")
    return row


async def _admission(request: Request, user: dict, admission_id: str) -> dict:
    row = await db(request).admissions.find_one({"_id": admission_id, "institute_id": user["institute_id"]})
    if row is None:
        raise AppError(404, "not_found", "That record was not found.")
    return row


def _event(kind: str, note: str, actor: str) -> dict:
    return {"id": str(ObjectId()), "kind": kind, "at": _now().isoformat(), "note": note, "actor": actor}


def _require_open_follow_up(row: dict) -> None:
    status = enquiry_status(row)
    if status == "converted":
        raise AppError(409, "enquiry.converted", "This enquiry is an admission. Follow-ups stay on the admission.")
    if status == "closed":
        raise AppError(409, "enquiry.closed", "Reopen the enquiry before recording a follow-up.")


async def _institute_today(request: Request, user: dict) -> str:
    from app.clock import local_today

    institute = await db(request).institutes.find_one({"_id": user["institute_id"]}, {"timezone": 1})
    return local_today((institute or {}).get("timezone"))


def _require_open_admission(row: dict) -> None:
    if row.get("status") == "converted":
        raise AppError(409, "admission.enrolled", "This application is already enrolled.")
    if row.get("status") == "closed":
        raise AppError(409, "admission.closed", "Reopen the application before continuing.")


async def _course_change_summary(request: Request, user: dict, row: dict, course_id: str) -> dict:
    await _course(request, user, course_id)
    same = (row.get("course_id") or "") == course_id
    questions = await _questions_for(request, course_id)
    current_ids = {item["_id"] for item in questions}
    preserved = [
        item.get("label") or "Previous answer"
        for item in row.get("course_answers") or []
        if item.get("question_id") not in current_ids and (item.get("value") or "")
    ]
    context = await _selection_context(request, row)
    clears_batch = bool(row.get("batch_id") or row.get("branch_id")) and row.get("status") != "converted" and not same
    parts = []
    if preserved:
        parts.append("Previous answers stay on the application: " + ", ".join(preserved) + ".")
    if clears_batch:
        batch = context["batch_name"] or "the selected batch"
        parts.append(f"The saved batch {batch} will be cleared.")
    incoming = [item.get("label") or "Question" for item in questions if item.get("required")]
    if incoming and not same:
        parts.append("This course asks: " + ", ".join(incoming) + ".")
    return {
        "changed": not same,
        "preserved": preserved,
        "clears_batch": clears_batch,
        "batch_name": context["batch_name"],
        "needs_confirmation": bool(preserved or clears_batch),
        "message": " ".join(parts) or "The course change does not affect saved answers or a batch.",
    }


async def _guard_admission_repeat(
    request: Request, user: dict, fields: dict, current_id: str = "", session=None
) -> None:
    phone = fields.get("phone_key") or ""
    if not phone or not fields.get("name_key") or not fields.get("course_id"):
        return
    rows = await db(request).admissions.find(
        {
            "institute_id": user["institute_id"],
            "status": "open",
            "name_key": fields["name_key"],
            "phone_key": phone,
            "course_id": fields["course_id"],
        },
        session=session,
    ).to_list(length=20)
    batch_id = fields.get("batch_id") or ""
    match = None
    for existing in rows:
        if existing["_id"] == current_id:
            continue
        other_batch = existing.get("batch_id") or ""
        if batch_id and other_batch and batch_id != other_batch:
            continue
        match = existing
        break
    if match is None or (fields.get("episode_note") or "").strip():
        return
    course = await _course_name(request, fields["course_id"])
    label = course or "this course"
    raise AppError(
        422,
        "admission.repeat",
        f"An open application already exists for {label}. Explain why this is a new application.",
        field="episode_note",
        admission_id=match["_id"],
    )


async def _open_repeats(request: Request, user: dict, row: dict) -> list[dict]:
    phone = row.get("phone_key") or phone_key(row.get("phone") or "")
    if not phone or not row.get("name_key"):
        return []
    found = []
    cursor = db(request).admissions.find(
        {
            "institute_id": user["institute_id"],
            "status": "open",
            "name_key": row.get("name_key"),
            "phone_key": phone,
            "course_id": row.get("course_id") or "",
        }
    )
    async for existing in cursor:
        if existing["_id"] == row["_id"]:
            continue
        found.append(
            {
                "id": existing["_id"],
                "reference": existing.get("reference") or "",
                "name": existing.get("name") or "",
            }
        )
    return found


async def _prospect_view(request: Request, user: dict, prospect_id: str) -> dict | None:
    if not prospect_id:
        return None
    person = await db(request).prospects.find_one({"_id": prospect_id, "institute_id": user["institute_id"]})
    if person is None:
        return None
    return {"id": person["_id"], "name": person.get("name") or "", "whatsapp": person.get("whatsapp") or ""}


async def _eligible_batch_choices(request: Request, user: dict, row: dict) -> list[dict]:
    course_id = row.get("course_id") or ""
    if not course_id:
        return []
    batches = await db(request).batches.find(
        {"institute_id": user["institute_id"], "course_id": course_id}
    ).to_list(length=200)
    branches = {
        item["_id"]: item.get("name") or ""
        for item in await db(request).branches.find({"institute_id": user["institute_id"]}).to_list(length=200)
    }
    choices = []
    for batch in batches:
        resource = {
            "institute_id": user["institute_id"],
            "branch_id": batch.get("branch_id"),
            "course_id": course_id,
            "batch_id": batch["_id"],
        }
        if not allows(user["grants"], "student.manage", resource) and not allows(user["grants"], "student.lookup", resource):
            continue
        choices.append(
            {
                "id": batch["_id"],
                "branch_id": batch.get("branch_id") or "",
                "branch_name": branches.get(batch.get("branch_id") or "", ""),
                "name": batch.get("name") or "",
                "started_on": batch.get("started_on") or "",
                "ended_on": batch.get("ended_on") or "",
                "open": not batch_is_closed(batch),
            }
        )
    choices.sort(key=lambda item: (item["branch_name"].casefold(), item["name"].casefold()))
    return choices


async def _eligible_batch(request: Request, user: dict, row: dict, branch_id: str, batch_id: str) -> dict:
    if not allows(
        user["grants"],
        "student.manage",
        {
            "institute_id": user["institute_id"],
            "branch_id": branch_id,
            "course_id": row.get("course_id"),
            "batch_id": batch_id,
        },
    ):
        raise AppError(403, "auth.forbidden", "You do not have permission for that action.")
    batch = await db(request).batches.find_one({"_id": batch_id, "institute_id": user["institute_id"]})
    if batch is None or batch.get("branch_id") != branch_id or batch.get("course_id") != row.get("course_id"):
        raise AppError(422, "enrollment.context", "Choose a branch and batch for this course.")
    require_open_batch(batch)
    return batch


async def _selection_context(request: Request, row: dict) -> dict:
    branch_id = row.get("branch_id") or ""
    batch_id = row.get("batch_id") or ""
    branch_name = ""
    batch_name = ""
    batch_open = False
    if branch_id:
        branch = await db(request).branches.find_one({"_id": branch_id})
        branch_name = (branch or {}).get("name") or ""
    if batch_id:
        batch = await db(request).batches.find_one({"_id": batch_id})
        if batch:
            batch_name = batch.get("name") or ""
            batch_open = (
                not batch_is_closed(batch)
                and batch.get("branch_id") == branch_id
                and batch.get("course_id") == (row.get("course_id") or "")
            )
    return {"branch_name": branch_name, "batch_name": batch_name, "batch_open": batch_open}


async def _allocate_admission_reference(request: Request, institute_id: str, session=None) -> str:
    counter = await db(request).counters.find_one_and_update(
        {"_id": f"admission:{institute_id}"},
        {"$inc": {"value": 1}},
        upsert=True,
        return_document=ReturnDocument.AFTER,
        session=session,
    )
    return f"ADM-{int(counter['value']):06d}"


async def _ensure_admission_reference(request: Request, user: dict, row: dict) -> dict:
    if row.get("reference"):
        return row
    reference = await _allocate_admission_reference(request, user["institute_id"])
    await db(request).admissions.update_one({"_id": row["_id"]}, {"$set": {"reference": reference}})
    row["reference"] = reference
    return row


def _local_day(value, zone: str) -> str:
    from zoneinfo import ZoneInfo

    from app.clock import coerce_timezone

    if isinstance(value, datetime):
        moment = value if value.tzinfo else value.replace(tzinfo=timezone.utc)
        return moment.astimezone(ZoneInfo(coerce_timezone(zone))).date().isoformat()
    return str(value or "")[:10]


def _sort_admissions(rows: list[dict], sort: str, direction: str) -> None:
    desc = (direction or "desc") != "asc"
    field = sort or "created"
    rows.sort(key=lambda row: str(row.get("_id") or ""))
    if field == "name":
        rows.sort(key=lambda row: str(row.get("name") or "").casefold(), reverse=desc)
        return
    if field == "updated":
        rows.sort(key=lambda row: row.get("updated_at") or row.get("created_at") or datetime.min.replace(tzinfo=timezone.utc), reverse=desc)
        return
    rows.sort(key=lambda row: row.get("created_at") or datetime.min.replace(tzinfo=timezone.utc), reverse=desc)


async def _allocate_reference(request: Request, institute_id: str, session=None) -> str:
    counter = await db(request).counters.find_one_and_update(
        {"_id": f"enquiry:{institute_id}"},
        {"$inc": {"value": 1}},
        upsert=True,
        return_document=ReturnDocument.AFTER,
        session=session,
    )
    return f"ENQ-{int(counter['value']):06d}"


async def _guard_repeat(
    request: Request, user: dict, prospect_id: str, course_id: str, episode_note: str, session=None
) -> None:
    if not prospect_id:
        return
    rows = await db(request).enquiries.find(
        {"institute_id": user["institute_id"], "prospect_id": prospect_id}, session=session
    ).to_list(length=50)
    for existing in rows:
        if enquiry_status(existing) != "open":
            continue
        if (existing.get("course_id") or "") != (course_id or ""):
            continue
        if episode_note.strip():
            return
        label = "this course" if course_id else "an unspecified course"
        raise AppError(
            422,
            "intake.field",
            f"An open enquiry already exists for {label}. Explain why this is a new enquiry.",
            field="episode_note",
            enquiry_id=existing["_id"],
        )


async def _insert_enquiry(
    request: Request,
    user: dict,
    fields: dict,
    *,
    prospect_id: str = "",
    episode_note: str = "",
    session=None,
) -> str:
    database = db(request)
    person = {
        "name": fields["name"],
        "name_key": fields["name_key"],
        "whatsapp": fields.get("whatsapp") or "",
        "whatsapp_key": fields.get("whatsapp_key") or "",
        "email": fields.get("email") or "",
        "email_key": fields.get("email_key") or "",
        "present_institution": fields.get("present_institution") or "",
    }
    if prospect_id:
        if prospect_id.startswith("row:"):
            raise AppError(422, "import.field", "Link the row to a person created in this import.", field="prospect_id")
        prospect = await database.prospects.find_one(
            {"_id": prospect_id, "institute_id": user["institute_id"]}, session=session
        )
        if prospect is None:
            raise AppError(404, "not_found", "That person was not found.")
        await _guard_repeat(request, user, prospect_id, fields.get("course_id") or "", episode_note, session=session)
        person = {
            "name": prospect.get("name") or "",
            "name_key": prospect.get("name_key") or "",
            "whatsapp": prospect.get("whatsapp") or "",
            "whatsapp_key": prospect.get("whatsapp_key") or "",
            "email": prospect.get("email") or "",
            "email_key": prospect.get("email_key") or "",
            "present_institution": prospect.get("present_institution") or "",
        }
    else:
        prospect_id = str(ObjectId())
        await database.prospects.insert_one(
            {"_id": prospect_id, "institute_id": user["institute_id"], "created_at": _now(), **person},
            session=session,
        )
    enquiry_id = str(ObjectId())
    reference = await _allocate_reference(request, user["institute_id"], session=session)
    await database.enquiries.insert_one(
        {
            "_id": enquiry_id,
            "institute_id": user["institute_id"],
            "prospect_id": prospect_id,
            "reference": reference,
            "status": "open",
            "pending_on": "",
            "pending_task_id": "",
            "last_contact_on": "",
            "closure_reason": "",
            "closure_note": "",
            "episode_note": episode_note.strip(),
            "follow_ups": [],
            "events": [],
            "created_at": _now(),
            "updated_at": _now(),
            **person,
            "course_id": fields.get("course_id") or "",
            "interest": fields.get("interest") or "",
            "references": fields.get("references") or "",
            "details": fields.get("details") or "",
            "answers": fields.get("answers") or [],
        },
        session=session,
    )
    return enquiry_id


async def _ensure_enquiry_row(request: Request, user: dict, row: dict) -> dict:
    sets: dict = {}
    if not row.get("prospect_id"):
        prospect_id = str(ObjectId())
        await db(request).prospects.insert_one(
            {
                "_id": prospect_id,
                "institute_id": user["institute_id"],
                "name": row.get("name") or "",
                "name_key": row.get("name_key") or name_key(row.get("name") or ""),
                "whatsapp": row.get("whatsapp") or "",
                "whatsapp_key": row.get("whatsapp_key") or phone_key(row.get("whatsapp") or ""),
                "email": row.get("email") or "",
                "email_key": name_key(row.get("email") or ""),
                "present_institution": row.get("present_institution") or "",
                "created_at": _now(),
            }
        )
        sets["prospect_id"] = prospect_id
    if not row.get("reference"):
        sets["reference"] = await _allocate_reference(request, user["institute_id"])
    if not row.get("status"):
        sets["status"] = "converted" if row.get("admission_id") else "open"
    if "pending_on" not in row:
        follow_ups = row.get("follow_ups") or []
        sets["pending_on"] = str(follow_ups[-1].get("next_on") or "") if follow_ups else ""
    if "events" not in row:
        sets["events"] = []
    if not sets:
        return row
    if not row.get("updated_at"):
        sets["updated_at"] = _now()
    await db(request).enquiries.update_one({"_id": row["_id"]}, {"$set": sets})
    row.update(sets)
    return row


async def _prepare_enquiry(request: Request, user: dict, enquiry_id: str) -> dict:
    return await _ensure_enquiry_row(request, user, await _enquiry(request, user, enquiry_id))


async def _replace_enquiry(
    request: Request, row: dict, sets: dict, events: list[dict], follow: dict | None = None
) -> None:
    update: dict = {"$set": sets}
    push: dict = {}
    if events:
        push["events"] = {"$each": events}
    if follow:
        push["follow_ups"] = follow
    if push:
        update["$push"] = push
    query: dict = {"_id": row["_id"]}
    if row.get("updated_at") is not None:
        query["updated_at"] = row["updated_at"]
    result = await db(request).enquiries.update_one(query, update)
    if result.modified_count == 0:
        raise AppError(409, "enquiry.conflict", "This enquiry changed. Reload and try again.")


async def _save_follow_up(request: Request, user: dict, row: dict, body: FollowUpBody, today: str) -> dict:
    kind = body.kind or ""
    if kind not in {"", "record", "schedule"}:
        raise AppError(422, "intake.field", "Choose record or schedule.", field="kind")
    actor = user.get("name") or ""
    interest = _choice(body.interest, INTERESTS, "interest", "Choose High, Medium, or Low.")
    if body.task_id and row.get("pending_task_id") and body.task_id != row.get("pending_task_id"):
        raise AppError(409, "enquiry.conflict", "That follow-up is no longer the pending one.")
    if kind == "schedule":
        next_on = _date(body.next_on, "next_on")
        if not next_on or next_on < today:
            raise AppError(422, "intake.field", "Schedule today or a later date.", field="next_on")
        note = _long_text(body.note, 500, "note")
        previous = row.get("pending_on") or ""
        task_id = str(ObjectId())
        message = f"Scheduled a follow-up on {next_on}."
        if previous and previous != next_on:
            message = f"Rescheduled the follow-up from {previous} to {next_on}."
        if note:
            message = f"{message} {note}"
        await _replace_enquiry(
            request,
            row,
            {"pending_on": next_on, "pending_task_id": task_id, "updated_at": _now()},
            [_event("scheduled", message, actor)],
        )
        await write_audit(request, user, "enquiry.schedule", None, {"next_on": next_on}, {"enquiry_id": row["_id"]})
        return {"id": task_id, "kind": "schedule", "next_on": next_on, "note": note}
    at = _date(body.at, "at")
    note = _long_text(body.note, 500, "note")
    next_on = _date(body.next_on, "next_on")
    if kind == "record":
        if not at or at > today:
            raise AppError(422, "intake.field", "Enter a contact date no later than today.", field="at")
        if not note:
            raise AppError(422, "intake.field", "Enter what was discussed.", field="note")
        if next_on and next_on < today:
            raise AppError(422, "intake.field", "The next follow-up must be today or later.", field="next_on")
    elif not at or not note:
        raise AppError(422, "intake.field", "Enter the follow-up date and what was discussed.", field="note")
    follow = {
        "id": str(ObjectId()),
        "at": at,
        "note": note,
        "interest": interest,
        "next_on": next_on,
        "actor": actor,
    }
    sets = {
        "pending_on": next_on,
        "pending_task_id": str(ObjectId()) if next_on else "",
        "last_contact_on": at,
        "updated_at": _now(),
    }
    if kind == "record" and interest:
        sets["interest"] = interest
    await _replace_enquiry(request, row, sets, [], follow)
    await write_audit(request, user, "enquiry.follow_up", None, {"at": at, "note": note}, {"enquiry_id": row["_id"]})
    return follow


def _operation_id(institute_id: str, purpose: str, key: str) -> str:
    return f"{institute_id}:{purpose}:{key}"


def _digest(payload: dict) -> str:
    raw = json.dumps(payload, sort_keys=True, default=str, separators=(",", ":"))
    return hashlib.sha256(raw.encode()).hexdigest()


async def _claim_operation(request: Request, user: dict, purpose: str, key: str, payload: dict):
    if not key:
        return None
    digest = _digest(payload)
    operation_id = _operation_id(user["institute_id"], purpose, key)
    found = await db(request).intake_operations.find_one({"_id": operation_id})
    if found:
        if found.get("digest") != digest:
            raise AppError(409, "intake.idempotency", "This retry does not match the original request.")
        if found.get("response") is not None:
            return found["response"]
        raise AppError(409, "intake.idempotency", "That request is already in progress.")
    try:
        await db(request).intake_operations.insert_one(
            {"_id": operation_id, "digest": digest, "response": None, "created_at": _now()}
        )
    except DuplicateKeyError:
        return await _claim_operation(request, user, purpose, key, payload)
    return None


async def _finish_operation(request: Request, user: dict, purpose: str, key: str, response: dict) -> None:
    if not key:
        return
    await db(request).intake_operations.update_one(
        {"_id": _operation_id(user["institute_id"], purpose, key)}, {"$set": {"response": response}}
    )


async def _release_operation(request: Request, user: dict, purpose: str, key: str) -> None:
    if not key:
        return
    await db(request).intake_operations.delete_one(
        {"_id": _operation_id(user["institute_id"], purpose, key), "response": None}
    )


def _local_day(stamp, zone: str) -> str:
    from zoneinfo import ZoneInfo

    if stamp is None or not hasattr(stamp, "astimezone"):
        return ""
    if stamp.tzinfo is None:
        stamp = stamp.replace(tzinfo=timezone.utc)
    return stamp.astimezone(ZoneInfo(zone)).date().isoformat()


def _bounded_date(value: str, field: str) -> str:
    if not value:
        return ""
    _date(value, field)
    return value


def _filter_enquiries(rows: list[dict], today: str, zone: str, **filters) -> list[dict]:
    follow_up = filters.get("follow_up") or ""
    if follow_up:
        if follow_up == "due":
            rows = [row for row in rows if enquiry_awaiting_follow_up(row, today)]
        elif follow_up in {"overdue", "today", "upcoming", "unscheduled", "admitted"}:
            rows = [row for row in rows if follow_up_state(row, today) == follow_up]
        else:
            raise AppError(422, "intake.field", "Choose a follow-up filter.", field="follow_up")
    stage = filters.get("stage") or ""
    if stage:
        if stage not in {"open", "converted", "closed"}:
            raise AppError(422, "intake.field", "Choose Open, Converted, or Closed.", field="stage")
        rows = [row for row in rows if enquiry_status(row) == stage]
    interest = filters.get("interest") or ""
    if interest:
        chosen = _choice(interest, INTERESTS, "interest", "Choose High, Medium, or Low.")
        rows = [row for row in rows if (row.get("interest") or "") == chosen]
    closure_reason = filters.get("closure_reason") or ""
    if closure_reason:
        chosen_reason = _choice(closure_reason, CLOSURE_REASONS, "closure_reason", "Choose a closure reason.")
        rows = [row for row in rows if (row.get("closure_reason") or "") == chosen_reason]
    created_from = _bounded_date(filters.get("created_from") or "", "created_from")
    created_to = _bounded_date(filters.get("created_to") or "", "created_to")
    next_from = _bounded_date(filters.get("next_from") or "", "next_from")
    next_to = _bounded_date(filters.get("next_to") or "", "next_to")
    if created_from or created_to:
        kept = []
        for row in rows:
            day = _local_day(row.get("created_at"), zone)
            if created_from and (not day or day < created_from):
                continue
            if created_to and (not day or day > created_to):
                continue
            kept.append(row)
        rows = kept
    if next_from or next_to:
        kept = []
        for row in rows:
            day = enquiry_next_on(row)
            if not day:
                continue
            if next_from and day < next_from:
                continue
            if next_to and day > next_to:
                continue
            kept.append(row)
        rows = kept
    return rows


def _stable_sort(rows: list[dict], key, direction: str, *, missing_last: bool) -> None:
    def filled(row: dict) -> bool:
        value = key(row)
        return value not in ("", None)

    present = [row for row in rows if filled(row)] if missing_last else list(rows)
    absent = [row for row in rows if not filled(row)] if missing_last else []
    present.sort(key=lambda row: row["_id"], reverse=direction == "desc")
    present.sort(key=key, reverse=direction == "desc")
    absent.sort(key=lambda row: row["_id"])
    rows[:] = present + absent


def _needs_order(rows: list[dict], today: str) -> None:
    groups = {"overdue": [], "today": [], "unscheduled": [], "rest": []}
    for row in rows:
        state = follow_up_state(row, today)
        groups[state if state in groups else "rest"].append(row)
    groups["overdue"].sort(key=lambda row: row["_id"])
    groups["overdue"].sort(key=lambda row: enquiry_next_on(row))
    groups["today"].sort(key=lambda row: row["_id"])
    groups["unscheduled"].sort(key=lambda row: row["_id"])
    groups["unscheduled"].sort(key=lambda row: _local_stamp(row.get("created_at")))
    rows[:] = groups["overdue"] + groups["today"] + groups["unscheduled"] + groups["rest"]


def _local_stamp(value) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _sort_enquiries(rows: list[dict], sort: str, direction: str, follow_up: str, today: str) -> None:
    if direction and direction not in {"asc", "desc"}:
        raise AppError(422, "intake.field", "Choose ascending or descending.", field="direction")
    if sort not in {"", "created", "name", "next_on", "last_contact", "needs"}:
        raise AppError(422, "intake.field", "Choose a sort field.", field="sort")
    if follow_up == "due" and sort in {"", "needs"}:
        _needs_order(rows, today)
        return
    if sort == "needs":
        _needs_order(rows, today)
        return
    if sort in {"", "created"}:
        _stable_sort(
            rows,
            lambda row: _local_stamp(row.get("created_at")),
            "desc" if direction != "asc" else "asc",
            missing_last=False,
        )
        return
    if sort == "name":
        _stable_sort(rows, lambda row: (row.get("name") or "").casefold(), direction or "asc", missing_last=False)
        return
    if sort == "next_on":
        _stable_sort(rows, enquiry_next_on, direction or "asc", missing_last=True)
        return
    _stable_sort(rows, lambda row: row.get("last_contact_on") or "", direction or "desc", missing_last=True)


async def _list_people(
    request: Request,
    user: dict,
    collection: str,
    q: str,
    page: int,
    page_size: int,
    follow_up: str = "",
    status: str = "",
    course_id: str = "",
    sort: str = "",
    direction: str = "",
    stage: str = "",
    interest: str = "",
    closure_reason: str = "",
    created_from: str = "",
    created_to: str = "",
    next_from: str = "",
    next_to: str = "",
    readiness: str = "",
    branch_id: str = "",
    batch_id: str = "",
) -> dict:
    size = min(max(page_size, 1), 100)
    number = max(page, 1)
    query: dict = {"institute_id": user["institute_id"]}
    limit = 10000 if collection == "enquiries" else 500
    rows = await db(request)[collection].find(query).sort("created_at", -1).to_list(length=limit)
    rows = [row for row in rows if intake_course_visible(user["grants"], row.get("course_id") or "")]
    selected_course, _dropped = applied_intake_course(user["grants"], course_id)
    if selected_course:
        rows = [row for row in rows if (row.get("course_id") or "") == selected_course]
    today = ""
    zone = "Asia/Kolkata"
    if collection == "enquiries":
        from app.clock import coerce_timezone, local_today

        institute = await db(request).institutes.find_one({"_id": user["institute_id"]}, {"timezone": 1})
        zone = coerce_timezone((institute or {}).get("timezone"))
        today = local_today(zone)
        prepared = []
        for row in rows:
            prepared.append(await _ensure_enquiry_row(request, user, row))
        rows = prepared
        rows = _filter_enquiries(
            rows,
            today,
            zone,
            follow_up=follow_up,
            stage=stage,
            interest=interest,
            closure_reason=closure_reason,
            created_from=created_from,
            created_to=created_to,
            next_from=next_from,
            next_to=next_to,
        )
        _sort_enquiries(rows, sort, direction, follow_up, today)
    admission_zone = "Asia/Kolkata"
    if collection == "admissions":
        from app.clock import coerce_timezone

        institute = await db(request).institutes.find_one({"_id": user["institute_id"]}, {"timezone": 1})
        admission_zone = coerce_timezone((institute or {}).get("timezone"))
        business_questions = await _form_questions(request, user["institute_id"], "admission")
        question_cache: dict[str, list] = {}
        batches = {
            item["_id"]: item
            for item in await db(request).batches.find({"institute_id": user["institute_id"]}).to_list(length=500)
        }
        branches = {
            item["_id"]: item.get("name") or ""
            for item in await db(request).branches.find({"institute_id": user["institute_id"]}).to_list(length=200)
        }
        prepared = []
        for row in rows:
            row = await _ensure_admission_reference(request, user, row)
            course_key = row.get("course_id") or ""
            if course_key not in question_cache:
                question_cache[course_key] = await _questions_for(request, course_key)
            requirements = admission_requirements(row, question_cache[course_key], business_questions)
            batch = batches.get(row.get("batch_id") or "")
            batch_open = bool(
                batch
                and not batch_is_closed(batch)
                and batch.get("branch_id") == (row.get("branch_id") or "")
                and batch.get("course_id") == course_key
            )
            row["_requirements"] = requirements
            row["_readiness"] = admission_readiness(row, requirements, batch_open=batch_open)
            row["_batch_name"] = (batch or {}).get("name") or ""
            row["_branch_name"] = branches.get(row.get("branch_id") or "", "")
            row["_batch_open"] = batch_open
            prepared.append(row)
        rows = prepared
        if status:
            stored = "converted" if status == "enrolled" else status
            rows = [row for row in rows if (row.get("status") or "open") == stored]
        if readiness:
            rows = [row for row in rows if row.get("_readiness") == readiness]
        if created_from:
            rows = [row for row in rows if _local_day(row.get("created_at"), admission_zone) >= created_from]
        if created_to:
            rows = [row for row in rows if _local_day(row.get("created_at"), admission_zone) <= created_to]
        if branch_id:
            rows = [row for row in rows if (row.get("branch_id") or "") == branch_id]
        if batch_id:
            rows = [row for row in rows if (row.get("batch_id") or "") == batch_id]
        _sort_admissions(rows, sort, direction)
    needle = q.casefold().strip()
    if needle and collection == "enquiries":
        digits = phone_key(q)
        rows = [
            row
            for row in rows
            if needle in _enquiry_search(row) or (digits and digits in phone_key(row.get("whatsapp") or ""))
        ]
    elif needle and collection == "admissions":
        digits = phone_key(q)
        rows = [
            row
            for row in rows
            if needle in _admission_search(row) or (digits and digits in phone_key(row.get("phone") or ""))
        ]
    elif needle:
        rows = [row for row in rows if needle in _search_text(row)]
    total = len(rows)
    start = (number - 1) * size
    chosen = rows[start : start + size]
    course_ids = list({row.get("course_id") for row in chosen if row.get("course_id")})
    names = {}
    if course_ids:
        for course in await db(request).courses.find({"_id": {"$in": course_ids}}).to_list(length=len(course_ids)):
            names[course["_id"]] = course.get("name") or ""
    business_questions = []
    course_questions: dict[str, list] = {}
    if collection == "admissions":
        business_questions = await _form_questions(request, user["institute_id"], "admission")
        for course_id in course_ids:
            course_questions[course_id] = await _questions_for(request, course_id)
    items = []
    for row in chosen:
        if collection == "enquiries":
            view = _enquiry_view(row, names.get(row.get("course_id"), ""), today)
            view["follow_up_count"] = len(row.get("follow_ups") or [])
            items.append(view)
        else:
            course_id = row.get("course_id") or ""
            blockers = admission_blockers(row, course_questions.get(course_id, []), business_questions)
            matches = []
            if (row.get("status") or "open") != "converted" and not blockers and not row.get("match_decision"):
                matches = await find_matches(
                    request, user["institute_id"], row.get("name") or "", row.get("phone") or ""
                )
            items.append(
                _admission_view(
                    row,
                    names.get(course_id, ""),
                    matches,
                    course_questions.get(course_id, []),
                    business_questions,
                    branch_name=row.get("_branch_name") or "",
                    batch_name=row.get("_batch_name") or "",
                    batch_open=bool(row.get("_batch_open")),
                )
            )
    payload = {"items": items, "total": total, "page": number, "pages": max(1, (total + size - 1) // size)}
    if collection == "enquiries":
        payload["today"] = today
    return payload


def _enquiry_search(row: dict) -> str:
    """Identity search uses reference, name, contact, and email. Notes stay out. Name order is casefold."""
    return " ".join(
        [
            str(row.get("reference") or ""),
            str(row.get("name") or ""),
            str(row.get("whatsapp") or ""),
            str(row.get("email") or ""),
        ]
    ).casefold()


def _admission_search(row: dict) -> str:
    return " ".join(
        [
            str(row.get("reference") or ""),
            str(row.get("name") or ""),
            str(row.get("phone") or ""),
            str(row.get("email") or ""),
        ]
    ).casefold()


def _search_text(row: dict) -> str:
    return " ".join(
        str(row.get(key) or "") for key in ("name", "whatsapp", "phone", "present_institution", "interest")
    ).casefold()


_ENQUIRY_HEADERS = {
    "name": "name",
    "whatsapp": "whatsapp",
    "whatsapp number": "whatsapp",
    "present institution": "present_institution",
    "interested course": "course",
    "course": "course",
    "level of interest": "interest",
    "interest": "interest",
    "references": "references",
    "details": "details",
    "other details": "details",
}
_ADMISSION_HEADERS = {
    "course": "course",
    "name": "name",
    "contact number": "phone",
    "phone": "phone",
    "contact": "phone",
    "dob": "date_of_birth",
    "date of birth": "date_of_birth",
    "gender": "gender",
    "blood group": "blood_group",
    "email": "email",
    "address line 1": "street",
    "address line 2": "address_line_2",
    "city": "place",
    "district": "district",
    "state": "state",
    "pin": "pin",
    "parent or guardian": "guardian_name",
    "name of parent/guardian": "guardian_name",
    "guardian": "guardian_name",
    "parent contact": "guardian_phone",
    "contact number of parent/guardian": "guardian_phone",
    "guardian phone": "guardian_phone",
    "qualification": "qualification",
    "stream": "stream",
    "stream/department": "stream",
    "department": "stream",
    "institution": "institution",
    "name of the institution": "institution",
    "mark": "mark",
    "year of passing": "year_of_passing",
    "year": "year_of_passing",
}


def _admission_from_map(
    mapped: dict, course_id: str, extras: list[AnswerBody], business: list[AnswerBody] | None = None
) -> AdmissionBody:
    record = {
        "qualification": mapped.get("qualification") or "",
        "stream": mapped.get("stream") or "",
        "institution": mapped.get("institution") or "",
        "mark": mapped.get("mark") or "",
        "year_of_passing": mapped.get("year_of_passing") or "",
        "board": "",
    }
    records = [record] if any(record.values()) else []
    return AdmissionBody(
        course_id=course_id,
        name=mapped.get("name") or "",
        phone=mapped.get("phone") or "",
        date_of_birth=mapped.get("date_of_birth") or "",
        gender=mapped.get("gender") or "",
        blood_group=mapped.get("blood_group") or "",
        email=mapped.get("email") or "",
        street=mapped.get("street") or "",
        address_line_2=mapped.get("address_line_2") or "",
        place=mapped.get("place") or "",
        district=mapped.get("district") or "",
        state=mapped.get("state") or "",
        pin=mapped.get("pin") or "",
        guardian_name=mapped.get("guardian_name") or "",
        guardian_phone=mapped.get("guardian_phone") or "",
        academic_records=records,
        course_answers=extras,
        business_answers=business or [],
    )


async def _workbook_bytes(file: UploadFile) -> bytes:
    filename = file.filename or "intake.xlsx"
    if not filename.lower().endswith(".xlsx"):
        raise AppError(422, "import.file_type", "Upload an .xlsx workbook.")
    data = await file.read()
    if len(data) > _MAX_BYTES:
        raise AppError(422, "import.too_large", "The workbook is larger than 5 MB.")
    return data


def _sheet_rows(data: bytes) -> list[tuple[int, dict[str, str]]]:
    try:
        workbook = load_workbook(BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        raise AppError(422, "import.unreadable", "This workbook could not be read.") from exc
    sheet = workbook.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise AppError(422, "import.unreadable", "The workbook has no rows.")
    headers = [_cell_text(value) for value in rows[0]]
    if not any(headers):
        raise AppError(422, "import.unreadable", "The first row needs column names.")
    parsed = []
    for index, row in enumerate(rows[1:], start=2):
        cells = {}
        for header, value in zip(headers, row):
            if not header:
                continue
            text = _cell_text(value)
            if text:
                cells[header] = text
        if cells:
            parsed.append((index, cells))
    return parsed


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def _map_headers(cells: dict[str, str], known: dict[str, str]) -> tuple[dict[str, str], list[str]]:
    mapped: dict[str, str] = {}
    unknown = []
    for header, value in cells.items():
        key = known.get(name_key(header))
        if key is None:
            unknown.append(header)
        else:
            mapped[key] = value
    return mapped, unknown


def _question_columns(
    cells: dict[str, str], unknown: list[str], questions: list[dict]
) -> tuple[list[AnswerBody], list[str]]:
    keys = {name_key(item.get("label") or ""): item for item in questions}
    extras = []
    leftover = []
    for header in unknown:
        question = keys.get(name_key(header))
        if question is None:
            leftover.append(header)
        else:
            extras.append(AnswerBody(question_id=question["_id"], value=cells.get(header) or ""))
    return extras, leftover

from __future__ import annotations

import hashlib
import re
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.acl import allows
from app.calculating import (
    aggregate,
    band_detail,
    dense_ranks,
    display_percentage,
    percentage,
    percentage_change,
    public_bands,
    selected_attempts,
)


async def policy_for(db, institute_id: str) -> dict:
    policy = await db.policies.find_one({"institute_id": institute_id}, sort=[("version", -1)])
    return policy or {"version": 1, "bands": {"danger_below": 40, "safe_above": 60}, "aggregate": "maximum-marks-weighted", "passing_threshold": None}


async def hydrate(db, results: list[dict]) -> list[dict]:
    sheet_ids = list({item.get("marksheet_id") for item in results if item.get("marksheet_id")})
    sheets = {}
    if sheet_ids:
        for sheet in await db.marksheets.find({"_id": {"$in": sheet_ids}}).to_list(length=len(sheet_ids)):
            sheets[sheet["_id"]] = sheet
    rows = []
    for result in results:
        sheet = sheets.get(result.get("marksheet_id"), {})
        row = dict(result)
        row["id"] = result["_id"]
        row["exam_date"] = result.get("exam_date") or sheet.get("exam_date")
        row["title"] = result.get("title") or sheet.get("title") or sheet.get("status")
        row["assessment_id"] = result.get("assessment_id") or sheet.get("assessment_id")
        row["exam_type"] = result.get("exam_type") or sheet.get("exam_type")
        row["exam_kind"] = result.get("exam_kind") or sheet.get("exam_kind") or ""
        row["exam_id"] = result.get("exam_id") or sheet.get("exam_id") or ""
        row["maximum"] = result.get("maximum") if result.get("maximum") is not None else sheet.get("maximum")
        row["attempt_kind"] = result.get("attempt_kind") or "original"
        row["published"] = sheet.get("status") == "published" or result.get("active") is True
        row["policy_version"] = sheet.get("policy_version")
        row["bands"] = sheet.get("bands")
        rows.append(row)
    return rows


def describe_part(
    subject_id: str,
    exam_type: str,
    paper_id: str,
    scope: dict | None,
    paper_names: dict[str, str],
    unit_names: dict[str, str],
    kind_labels: dict[str, str],
) -> tuple[str, str]:
    """Return a stable part key and the label shown beside the subject."""
    scope = scope or {}
    kind = str(scope.get("kind") or exam_type or "subject")
    item_ids = [str(item) for item in (scope.get("item_ids") or []) if item]
    if kind == "subject":
        return f"{subject_id}|subject", "Whole subject"
    if kind in {"paper", "unit"} and (paper_id or item_ids):
        chosen = item_ids or ([paper_id] if paper_id else [])
        labels = [paper_names.get(item) or unit_names.get(item) or "" for item in chosen]
        labels = [item for item in labels if item]
        text = ", ".join(labels) if labels else (paper_names.get(paper_id) or "Paper")
        identity = "|".join(chosen) or paper_id or kind
        return f"{subject_id}|paper|{identity}", text
    if item_ids:
        labels = [unit_names.get(item) or paper_names.get(item) or "" for item in item_ids]
        labels = [item for item in labels if item]
        kind_label = kind_labels.get(kind) or kind.replace("-", " ").title()
        text = f"{kind_label}: {', '.join(labels)}" if labels else kind_label
        return f"{subject_id}|{kind}|{'|'.join(item_ids)}", text
    if paper_id:
        return f"{subject_id}|paper|{paper_id}", paper_names.get(paper_id) or "Paper"
    if kind in {"", "subject"}:
        return f"{subject_id}|subject", "Whole subject"
    kind_label = kind_labels.get(kind) or kind.replace("-", " ").title()
    return f"{subject_id}|{kind}", kind_label


def exam_columns(rows: list[dict], types: list[dict], limit: int = 5) -> list[tuple[str, str]]:
    """Exam types that appear on these results, lightest weight first, at most five."""
    weights = {item["key"]: item.get("weight") or 99 for item in types}
    labels = {item["key"]: item.get("label") or item["key"] for item in types}
    keys = list({row.get("exam_kind") or "" for row in rows if row.get("exam_kind")})
    keys.sort(key=lambda key: (weights.get(key, 99), (labels.get(key) or key).casefold()))
    columns = [(key, labels.get(key) or key) for key in keys[:limit]]
    if any(not (row.get("exam_kind") or "") for row in rows) and len(columns) < limit:
        columns.append(("", "Exam"))
    return columns


async def attach_parts(db, rows: list[dict]) -> None:
    """Name the subject partition for each result from its exam scope."""
    if not rows:
        return
    from app.academics.curriculum import public_kinds

    exam_ids = list({row.get("exam_id") for row in rows if row.get("exam_id")})
    exams = {}
    if exam_ids:
        for exam in await db.exams.find({"_id": {"$in": exam_ids}}).to_list(length=len(exam_ids)):
            exams[exam["_id"]] = exam
    for row in rows:
        exam = exams.get(row.get("exam_id") or "")
        if exam and not row.get("exam_kind"):
            row["exam_kind"] = exam.get("exam_kind") or ""
    subject_ids = {row.get("subject_id") for row in rows if row.get("subject_id")}
    subject_ids.update(exam.get("subject_id") for exam in exams.values() if exam.get("subject_id"))
    subjects = {}
    if subject_ids:
        found = await db.subjects.find({"_id": {"$in": list(subject_ids)}}).to_list(length=len(subject_ids))
        subjects = {item["_id"]: item for item in found}
    paper_ids = {row.get("paper_id") for row in rows if row.get("paper_id")}
    for exam in exams.values():
        if exam.get("paper_id"):
            paper_ids.add(exam.get("paper_id"))
        for item in (exam.get("scope") or {}).get("item_ids") or []:
            if item:
                paper_ids.add(item)
    paper_names = {}
    if paper_ids:
        found = await db.papers.find({"_id": {"$in": list(paper_ids)}}).to_list(length=len(paper_ids))
        paper_names = {item["_id"]: item.get("name") or "" for item in found}
    institute_id = next((row.get("institute_id") for row in rows if row.get("institute_id")), "")
    institute = await db.institutes.find_one({"_id": institute_id}) if institute_id else None
    kind_labels = {item["key"]: item["label"] for item in public_kinds((institute or {}).get("curriculum_kinds"))}
    for row in rows:
        exam = exams.get(row.get("exam_id") or "") or {}
        subject = subjects.get(row.get("subject_id") or exam.get("subject_id") or "") or {}
        units = {unit.get("id"): unit.get("name") or "" for unit in subject.get("units") or [] if unit.get("id")}
        paper_id = row.get("paper_id") or exam.get("paper_id") or ""
        key, label = describe_part(
            row.get("subject_id") or "",
            row.get("exam_type") or "",
            paper_id,
            exam.get("scope") or {},
            paper_names,
            units,
            kind_labels,
        )
        row["part_key"] = key
        row["part_label"] = label


def row_policy(row: dict, policy: dict | None) -> dict | None:
    if row.get("bands"):
        return {"bands": row["bands"]}
    return policy


def present(
    row: dict,
    policy: dict,
    rank: int | None,
    batch_rank: int | None,
    top_score=None,
    batch_top_score=None,
) -> dict:
    exact = percentage(row.get("score"), row.get("maximum")) if row.get("status") == "scored" else None
    detail = band_detail(exact, row_policy(row, policy)) if row.get("status") == "scored" else {"band": None, "band_name": None, "band_place": None}
    return {
        "id": row["id"],
        "subject_id": row.get("subject_id"),
        "paper_id": row.get("paper_id"),
        "batch_id": row.get("batch_id"),
        "status": row.get("status"),
        "score": row.get("score"),
        "maximum": row.get("maximum"),
        "percentage": display_percentage(exact),
        "band": detail["band"],
        "band_name": detail["band_name"],
        "band_place": detail["band_place"],
        "rank": rank,
        "batch_rank": batch_rank,
        "rank_label": "Cohort rank" if rank is not None else None,
        "top_score": top_score,
        "batch_top_score": batch_top_score,
        "exam_date": row.get("exam_date"),
        "title": row.get("title"),
        "exam_type": row.get("exam_type"),
        "exam_kind": row.get("exam_kind") or "",
        "part_key": row.get("part_key") or "",
        "part_label": row.get("part_label") or "",
        "attempt": row.get("attempt_kind") or "original",
        "assessment_id": row.get("assessment_id"),
        "marksheet_id": row.get("marksheet_id"),
        "included_in_total": row.get("counts_in_aggregate") is not False,
    }


def comparisons(rows: list[dict], outside: list[dict] | None = None) -> list[dict]:
    outside = outside or []
    included = {row["id"] for row in rows}
    grouped: dict[str, list[dict]] = {}
    for row in list(rows) + list(outside):
        key = row.get("assessment_id") or row["id"]
        grouped.setdefault(key, []).append(row)
    items = []
    for group in grouped.values():
        ordered = sorted(group, key=lambda item: (str(item.get("exam_date") or ""), item.get("revision") or 0))
        if len(ordered) < 2 or ordered[-1]["id"] not in included:
            continue
        items.append(
            {
                "title": ordered[-1].get("title"),
                "original_percentage": display_percentage(percentage(ordered[0].get("score"), ordered[0].get("maximum"))),
                "latest_percentage": display_percentage(percentage(ordered[-1].get("score"), ordered[-1].get("maximum"))),
                "change": percentage_change(ordered[0], ordered[-1]),
                "baseline_outside_period": ordered[0]["id"] not in included,
                "batch_id": ordered[-1].get("batch_id"),
            }
        )
    return items


def ranks_for(row: dict, siblings: list[dict]) -> tuple[int | None, int | None]:
    if row.get("status") != "scored":
        return None, None
    cohort = dense_ranks(siblings)
    same_batch = [item for item in siblings if item.get("batch_id") == row.get("batch_id")]
    batch = dense_ranks(same_batch)
    combined = cohort.get(row["id"])
    per_batch = batch.get(row["id"]) if len({item.get("batch_id") for item in siblings}) > 1 else None
    return combined, per_batch


def _highest_score(rows: list[dict]):
    chosen = None
    best = None
    for item in rows:
        if item.get("status") != "scored" or item.get("score") is None:
            continue
        value = Decimal(str(item["score"]))
        if best is None or value > best:
            best = value
            chosen = item
    return None if chosen is None else chosen.get("score")


def tops_for(siblings: list[dict], batch_id: str | None):
    """Highest scored mark on the exam, and within one batch when the exam is shared."""
    exam = _highest_score(siblings)
    if exam is None:
        return None, None
    batch = None
    if len({item.get("batch_id") for item in siblings}) > 1:
        batch = _highest_score([item for item in siblings if item.get("batch_id") == batch_id])
    return exam, batch


def matches(row: dict, filters: dict) -> bool:
    for field in ("branch_id", "course_id", "batch_id", "subject_id", "paper_id", "exam_type"):
        if filters.get(field) and row.get(field) != filters[field]:
            return False
    exam_date = str(row.get("exam_date") or "")
    if filters.get("exam_date") and exam_date != filters["exam_date"]:
        return False
    if filters.get("exam_date_from") and exam_date < filters["exam_date_from"]:
        return False
    if filters.get("exam_date_to") and (not exam_date or exam_date > filters["exam_date_to"]):
        return False
    attempt = filters.get("attempt")
    if attempt in {"original", "retest"} and (row.get("attempt_kind") or "original") != attempt:
        return False
    return True


def view_sections(level: str, rows: list[dict], enrollments: list[dict], names: dict[str, str], policy: dict | None = None) -> list[dict]:
    key = {
        "institute": "branch_id",
        "branch": "course_id",
        "course": "batch_id",
        "batch": "subject_id",
        "subject": "paper_id",
        "paper": "assessment_id",
    }[level]
    grouped: dict[str, list[dict]] = {}
    for row in selected_attempts(rows):
        grouped.setdefault(row.get(key) or "", []).append(row)
    sections = []
    for section_id, section_rows in grouped.items():
        people = {row.get("student_id") for row in section_rows if row.get("student_id")}
        summary = aggregate(section_rows, policy)
        sections.append(
            {
                "id": section_id,
                "label": names.get(section_id) or section_id or "Unassigned",
                "students": len(people),
                "results": len(section_rows),
                "percentage": summary["percentage"],
                "band": summary["band"],
                "band_name": summary.get("band_name"),
                "band_place": summary.get("band_place"),
                "coverage": f"{summary['scored']}/{summary['expected']}" if summary["expected"] else "0/0",
                "missing": summary["missing"],
                "absent": summary["absent"],
                "marksheet_ids": sorted({row["marksheet_id"] for row in section_rows if row.get("marksheet_id")}),
            }
        )
    if level == "batch":
        listed = {item["student_id"] for item in enrollments}
        for section in sections:
            section["listed_students"] = len(listed)
    return sections


async def names_for(db, rows: list[dict]) -> dict[str, str]:
    names = {}
    for collection, field in (
        ("branches", "branch_id"),
        ("courses", "course_id"),
        ("batches", "batch_id"),
        ("subjects", "subject_id"),
        ("papers", "paper_id"),
        ("assessments", "assessment_id"),
    ):
        ids = list({row.get(field) for row in rows if row.get(field)})
        if not ids:
            continue
        for item in await db[collection].find({"_id": {"$in": ids}}).to_list(length=len(ids)):
            names[item["_id"]] = item.get("name") or item.get("title") or (f"Paper {item['number']}" if item.get("number") is not None else item["_id"])
    return names


def authorized(rows: list[dict], grants: list[dict], action: str) -> list[dict]:
    visible = []
    for row in rows:
        resource = {
            "institute_id": row.get("institute_id"),
            "branch_id": row.get("branch_id"),
            "course_id": row.get("course_id"),
            "batch_id": row.get("batch_id"),
            "subject_id": row.get("subject_id"),
            "paper_id": row.get("paper_id"),
        }
        if allows(grants, action, resource):
            visible.append(row)
    return visible


def participation_summary(rows: list[dict], eligible_by_batch: dict[str, set[str]]) -> dict:
    """One opportunity per confirmed eligible student and assessment. Exemptions stay out of the denominator."""
    cohorts: dict[tuple, dict] = {}
    for row in selected_attempts(rows):
        batch_id = row.get("batch_id") or ""
        if batch_id not in eligible_by_batch:
            continue
        assessment = row.get("assessment_id") or row.get("marksheet_id") or row.get("id")
        cohorts.setdefault((assessment, batch_id), {})[row.get("student_id")] = row.get("status")
    scored = missing = absent = exempt = 0
    for (_assessment, batch_id), statuses in cohorts.items():
        for student_id in eligible_by_batch[batch_id]:
            status = statuses.get(student_id)
            if status == "exempt":
                exempt += 1
            elif status == "scored":
                scored += 1
            elif status == "absent":
                absent += 1
            else:
                missing += 1
    expected = scored + missing + absent
    ratio = round(scored / expected, 4) if expected else None
    return {
        "participation": ratio,
        "participation_scored": scored,
        "participation_expected": expected,
        "participation_missing": missing,
        "participation_absent": absent,
        "participation_exempt": exempt,
        "participation_basis": "confirmed eligible students for each assessment in this report",
    }


SECTION_FIELD = {
    "institute": "course_id",
    "branch": "course_id",
    "course": "batch_id",
    "batch": "subject_id",
    "subject": "paper_id",
    "paper": "assessment_id",
}


def _ranked_students(by_student: dict[str, list[dict]], policy: dict, limit: int) -> tuple[list[dict], int]:
    items = []
    for student_id, student_rows in by_student.items():
        if not student_id:
            continue
        summary = aggregate(student_rows, policy)
        if summary.get("band_place") != "high":
            continue
        items.append(
            {
                "student_id": student_id,
                "percentage": summary["percentage"],
                "band": summary["band"],
                "band_name": summary.get("band_name"),
                "band_place": summary.get("band_place"),
                "scored": summary["scored"],
                "expected": summary["expected"],
            }
        )
    items.sort(key=lambda item: (-item["percentage"], item["student_id"]))
    return items[:limit], len(items)


def student_boards(level: str, rows: list[dict], policy: dict, names: dict[str, str], limit: int = 20, attention_page: int = 1) -> dict:
    """Highest students, and danger results grouped by the next academic step."""
    field = SECTION_FIELD[level]
    chosen = selected_attempts(rows)
    grouped: dict[str, dict[str, list[dict]]] = {}
    section_of = {row.get("id"): (row.get(field) or "") for row in rows}
    for row in chosen:
        section = section_of.get(row.get("id"), "")
        grouped.setdefault(section, {}).setdefault(row.get("student_id") or "", []).append(row)
    danger_by_section: dict[str, list[dict]] = {}
    for item in attention(rows, policy):
        danger_by_section.setdefault(section_of.get(item["id"], ""), []).append(item)
    boards = []
    for section_id in set(grouped) | set(danger_by_section):
        shown, total = _ranked_students(grouped.get(section_id, {}), policy, limit)
        danger_items = danger_by_section.get(section_id, [])
        pages = max(1, (len(danger_items) + limit - 1) // limit) if danger_items else 1
        current = min(max(attention_page, 1), pages)
        boards.append(
            {
                "id": section_id,
                "label": names.get(section_id) or section_id or "Unassigned",
                "top_students": shown,
                "top_student_count": total,
                "attention": danger_items,
                "attention_result_count": len(danger_items),
                "attention_student_count": len({item.get("student_id") for item in danger_items if item.get("student_id")}),
                "attention_page": current,
                "attention_pages": pages,
            }
        )
    boards = [item for item in boards if item["top_student_count"] or item["attention_result_count"]]
    boards.sort(key=lambda item: item["label"])
    return {"top_students": [], "top_student_count": 0, "boards": boards}


def attention(rows: list[dict], policy: dict) -> list[dict]:
    items = []
    for row in rows:
        if row.get("status") != "scored":
            continue
        exact = percentage(row.get("score"), row.get("maximum"))
        detail = band_detail(exact, row_policy(row, policy))
        if detail["band_place"] != "low":
            continue
        items.append(
            {
                "id": row["id"],
                "student_id": row.get("student_id"),
                "title": row.get("title"),
                "score": row.get("score"),
                "maximum": row.get("maximum"),
                "percentage": display_percentage(exact),
                "band": detail["band"],
                "band_name": detail["band_name"],
                "band_place": detail["band_place"],
                "marksheet_id": row.get("marksheet_id"),
            }
        )
    items.sort(key=lambda item: (item["percentage"] is None, item["percentage"] if item["percentage"] is not None else 0, item.get("title") or ""))
    return items


def rank_page(rows: list[dict], page: int, size: int = 20) -> dict:
    scored = [row for row in rows if row.get("status") == "scored" and row.get("score") is not None]
    places = dense_ranks(scored)
    ordered = sorted(scored, key=lambda row: (places.get(row["id"], 10**6), str(row.get("title") or "")))
    start = max(page - 1, 0) * size
    pages = max(1, (len(ordered) + size - 1) // size) if ordered else 1
    return {
        "page": page,
        "pages": pages,
        "items": [
            {
                "id": row["id"],
                "title": row.get("title"),
                "score": row.get("score"),
                "maximum": row.get("maximum"),
                "rank": places.get(row["id"]),
            }
            for row in ordered[start : start + size]
        ],
    }


_MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
_ILLEGAL_FILENAME = re.compile(r'[\\/:*?"<>|\x00-\x1f]+')


def export_filename(
    kind: str,
    *,
    student_name: str = "",
    student_code: str = "",
    batch_name: str = "",
    on: date | None = None,
) -> str:
    extension = "xlsx" if kind == "report_xlsx" else "pdf"
    name = " ".join(student_name.split())
    code = " ".join(student_code.split())
    batch = " ".join(batch_name.split())
    if name or code:
        parts = [name, code]
        if batch:
            parts.append(batch)
        parts.append("progress card")
        stem = " ".join(part for part in parts if part)
    else:
        stamp = (on or datetime.now(timezone.utc).date()).isoformat()
        stem = f"Progress cards {stamp}"
    stem = _ILLEGAL_FILENAME.sub(" ", stem)
    stem = re.sub(r"\s+", " ", stem).strip(" .")
    stem = stem[:120].rstrip(" .") or "Progress card"
    return f"{stem}.{extension}"


def _format_date(value) -> str:
    text = str(value or "")
    parts = text.split("-")
    if len(parts) != 3:
        return text or "-"
    try:
        label = _MONTHS[int(parts[1]) - 1]
        return f"{int(parts[2])} {label} {parts[0]}"
    except (ValueError, IndexError):
        return text


def _percent(value) -> str:
    if value is None:
        return "-"
    return f"{float(value):.2f}%"


def _signed(value) -> str:
    if value is None:
        return "-"
    number = float(value)
    sign = "+" if number > 0 else ""
    return f"{sign}{number:.2f}"


def _plain_number(value) -> str:
    if value is None:
        return "-"
    number = float(value)
    if number.is_integer():
        return str(int(number))
    return f"{number:.2f}"


def _status_label(status: str | None) -> str:
    labels = {"missing": "Missing", "absent": "Absent", "exempt": "Exempt", "scored": "Scored"}
    return labels.get(status or "", status or "-")


def _mark_text(row: dict) -> str:
    if row.get("status") == "scored":
        return f"{_plain_number(row.get('score'))} of {_plain_number(row.get('maximum'))}"
    return _status_label(row.get("status"))


def _rank_text(row: dict) -> str:
    rank = row.get("rank")
    batch_rank = row.get("batch_rank")
    if rank is None and batch_rank is None:
        return "-"
    if rank is not None and batch_rank is not None:
        return f"{rank} | batch {batch_rank}"
    return str(rank if rank is not None else batch_rank)


def _top_text(row: dict) -> str:
    top = row.get("top_score")
    batch = row.get("batch_top_score")
    if top is None and batch is None:
        return "-"
    maximum = row.get("maximum")
    text = _mark_pair(top, maximum) if top is not None else "-"
    if batch is not None:
        text = f"{text} | batch {_mark_pair(batch, maximum)}"
    return text


def _mark_pair(score, maximum) -> str:
    return f"{_plain_number(score)} of {_plain_number(maximum)}"


def _context_name(mapping: dict, rows: list[dict], scope: dict, field: str) -> str:
    if scope.get(field):
        return mapping.get(scope[field]) or ""
    found = []
    for row in rows:
        label = mapping.get(row.get(field) or "")
        if label and label not in found:
            found.append(label)
    if len(found) == 1:
        return found[0]
    return ""


async def _name_map(db, collection: str, ids: set) -> dict[str, str]:
    values = [item for item in ids if item]
    found = {}
    if not values:
        return found
    for document in await db[collection].find({"_id": {"$in": values}}).to_list(length=len(values)):
        found[document["_id"]] = document.get("name") or ""
    return found


async def progress_document(
    db,
    institute_id: str,
    rows: list[dict],
    policy: dict,
    card: dict,
    scope: dict,
    kind: str,
    cohort: list[dict] | None = None,
) -> dict:
    published = [row for row in rows if row.get("published")]
    siblings = [row for row in (cohort if cohort is not None else rows) if row.get("published")]
    by_sheet: dict[str, list[dict]] = {}
    for row in siblings:
        by_sheet.setdefault(row.get("marksheet_id") or row.get("id"), []).append(row)
    prepared = []
    for row in published:
        item = dict(row)
        group = by_sheet.get(row.get("marksheet_id") or row.get("id"), [row])
        rank, batch_rank = ranks_for(item, group)
        top_score, batch_top_score = tops_for(group, item.get("batch_id"))
        item["rank"] = rank
        item["batch_rank"] = batch_rank
        item["top_score"] = top_score
        item["batch_top_score"] = batch_top_score
        prepared.append(item)

    student_ids = list({row.get("student_id") for row in prepared if row.get("student_id")})
    if card.get("student_id") and card["student_id"] not in student_ids:
        student_ids.append(card["student_id"])
    people = {}
    if student_ids:
        for student in await db.students.find({"_id": {"$in": student_ids}}).to_list(length=len(student_ids)):
            people[student["_id"]] = student
    enrollments = []
    if card.get("student_id"):
        enrollments = await db.enrollments.find(
            {"student_id": card["student_id"], "institute_id": institute_id}
        ).to_list(length=50)
    sources = list(prepared) + list(enrollments) + [scope]
    names = {
        "branch": await _name_map(db, "branches", {item.get("branch_id") for item in sources}),
        "course": await _name_map(db, "courses", {item.get("course_id") for item in sources}),
        "batch": await _name_map(db, "batches", {item.get("batch_id") for item in sources}),
        "subject": await _name_map(db, "subjects", {item.get("subject_id") for item in prepared}),
        "paper": await _name_map(db, "papers", {item.get("paper_id") for item in prepared}),
    }
    institute = await db.institutes.find_one({"_id": institute_id})
    await attach_parts(db, prepared)
    from app.academics.curriculum import public_exam_types

    columns = exam_columns(prepared, public_exam_types((institute or {}).get("exam_types")))
    exam_labels = {key: label for key, label in columns}
    grouped: dict[str, list[dict]] = {}
    for row in prepared:
        grouped.setdefault(row.get("student_id") or "", []).append(row)
    ordered_ids = sorted(
        student_ids,
        key=lambda student_id: (
            (people.get(student_id, {}).get("display_name") or "").casefold(),
            people.get(student_id, {}).get("student_code") or "",
        ),
    )
    students = []
    for student_id in ordered_ids:
        person = people.get(student_id, {})
        student_rows = grouped.get(student_id, [])
        identity_rows = student_rows or [
            item for item in enrollments if item.get("student_id") in {student_id, None} or card.get("student_id") == student_id
        ]
        students.append(_student_block(person, student_rows, identity_rows, names, policy, scope, columns, exam_labels))
    single = bool(card.get("student_id")) or len(students) == 1
    focus = students[0] if single and students else None
    now = datetime.now(timezone.utc).date()
    filename = export_filename(
        kind,
        student_name=(focus or {}).get("name") or "",
        student_code=(focus or {}).get("code") or "",
        batch_name=(focus or {}).get("batch") or "",
        on=now,
    )
    start = card.get("exam_date_from") or ""
    end = card.get("exam_date_to") or ""
    period = ""
    if start or end:
        period = f"{_format_date(start) if start else 'Start'} to {_format_date(end) if end else 'today'}"
    context = []
    if card.get("exam_type"):
        context.append(f"Exam type {card['exam_type']}")
    if card.get("attempt"):
        context.append(f"Attempt {card['attempt'].capitalize()}")
    if scope.get("subject_id"):
        context.append(f"Subject {names['subject'].get(scope['subject_id']) or 'Subject'}")
    if scope.get("paper_id"):
        context.append(f"Paper {names['paper'].get(scope['paper_id']) or 'Paper'}")
    notes = [
        f"Policy version {policy.get('version', 1)}.",
        "Authorized subjects only. Drafts are excluded.",
        "One active revision is included.",
        "The total uses the latest published attempt, weighted by each maximum.",
    ]
    for item in public_bands(policy.get("bands")):
        notes.append(f"{item['name']}: {item['phrase']}.")
    return {
        "filename": filename,
        "title": "Progress card" if single and focus else "Progress cards",
        "institute": (institute or {}).get("name") or "",
        "generated": _format_date(now.isoformat()),
        "single": bool(single and focus),
        "period": period,
        "context": context,
        "notes": notes,
        "students": students,
        "subject_headers": ["Subject", "Part", *[label for _key, label in columns]],
    }


def _student_block(
    person: dict,
    rows: list[dict],
    identity_rows: list[dict],
    names: dict,
    policy: dict,
    scope: dict,
    columns: list[tuple[str, str]] | None = None,
    exam_labels: dict | None = None,
) -> dict:
    summary = aggregate(rows, policy) if rows else {
        "percentage": None,
        "band_name": None,
        "scored": 0,
        "expected": 0,
        "missing": 0,
        "absent": 0,
    }
    columns = columns or []
    exam_labels = exam_labels or {}
    visible = {key for key, _label in columns}
    grouped: dict[tuple, list[dict]] = {}
    for row in rows:
        if (row.get("exam_kind") or "") not in visible:
            continue
        grouped.setdefault((row.get("subject_id") or "", row.get("part_key") or ""), []).append(row)
    ordered = sorted(
        grouped,
        key=lambda token: (
            (names["subject"].get(token[0]) or "Subject").casefold(),
            0 if (grouped[token][0].get("part_label") == "Whole subject") else 1,
            (grouped[token][0].get("part_label") or "").casefold(),
        ),
    )
    subjects = []
    seen_subject = None
    for token in ordered:
        group = grouped[token]
        subject_name = names["subject"].get(token[0]) or "Subject"
        show_subject = "" if token[0] == seen_subject else subject_name
        seen_subject = token[0]
        cells = []
        for key, _label in columns:
            subset = [row for row in group if (row.get("exam_kind") or "") == key]
            cells.append(_percent(aggregate(subset, policy)["percentage"]) if subset else "-")
        subjects.append([show_subject, group[0].get("part_label") or "Whole subject", *cells])
    exams = []
    corrections = []
    for row in rows:
        shown = present(row, policy, row.get("rank"), row.get("batch_rank"), row.get("top_score"), row.get("batch_top_score"))
        mark = _mark_text(shown)
        if row.get("previous_score") is not None:
            mark = f"{mark} (previous {_plain_number(row.get('previous_score'))})"
        title = shown.get("title") or "Result"
        subject = names["subject"].get(shown.get("subject_id") or "")
        if subject and subject not in title:
            title = f"{title} - {subject}"
        part = shown.get("part_label") or ""
        if part and part not in title:
            title = f"{title} - {part}"
        kind = exam_labels.get(shown.get("exam_kind") or "")
        if kind and kind not in title:
            title = f"{title} - {kind}"
        if shown.get("attempt") == "retest":
            title = f"{title} - Retest"
        if shown.get("included_in_total") is False:
            title = f"{title} - kept out"
        if row.get("previous_score") is not None:
            corrections.append(f"{title}: previous score {_plain_number(row.get('previous_score'))}")
        exams.append(
            {
                "date": str(shown.get("exam_date") or ""),
                "title": title,
                "cells": [
                    _format_date(shown.get("exam_date")),
                    title,
                    _percent(shown["percentage"]) if shown.get("status") == "scored" else _status_label(shown.get("status")),
                    mark,
                    shown.get("band_name") or _status_label(shown.get("status")),
                    _rank_text(shown),
                    _top_text(shown),
                ],
                "percentage": shown.get("percentage"),
                "place": shown.get("band_place"),
                "status": shown.get("status"),
            }
        )
    exams.sort(key=lambda item: item["title"].casefold())
    exams.sort(key=lambda item: item["date"] or "0000", reverse=True)
    attention = [
        [item["title"], item["cells"][2], item["cells"][3], item["cells"][4]]
        for item in sorted(exams, key=lambda entry: (entry["percentage"] is None, entry["percentage"] if entry["percentage"] is not None else 0))
        if item["place"] == "low"
    ]
    gaps = [
        [item["title"], "Absent" if item["status"] == "absent" else "Missing"]
        for item in exams
        if item["status"] in {"missing", "absent"}
    ]
    retests = []
    for item in comparisons(rows):
        basis = (
            "Original is outside this period and is not in the total"
            if item.get("baseline_outside_period")
            else "Both attempts are in this report"
        )
        retests.append(
            [
                item.get("title") or "Result",
                _percent(item.get("original_percentage")),
                _percent(item.get("latest_percentage")),
                _signed(item.get("change")),
                basis,
            ]
        )
    return {
        "name": person.get("display_name") or "Student",
        "code": person.get("student_code") or "",
        "branch": _context_name(names["branch"], identity_rows, scope, "branch_id"),
        "course": _context_name(names["course"], identity_rows, scope, "course_id"),
        "batch": _context_name(names["batch"], identity_rows, scope, "batch_id"),
        "performance": _percent(summary.get("percentage")),
        "band": summary.get("band_name") or "-",
        "scored": f"{summary.get('scored', 0)} of {summary.get('expected', 0)}",
        "missing": str(summary.get("missing", 0)),
        "absent": str(summary.get("absent", 0)),
        "subject_headers": ["Subject", "Part", *[label for _key, label in columns]] if columns else _SUBJECT_HEADERS,
        "subjects": subjects,
        "exams": [item["cells"] for item in exams],
        "retests": retests,
        "attention": attention,
        "gaps": gaps,
        "corrections": corrections,
    }


def export_bytes(kind: str, report: dict) -> tuple[bytes, str]:
    filename = report.get("filename") or export_filename(kind)
    if kind == "report_xlsx":
        return _workbook(report), filename
    return _pdf_document(report), filename


_PAGE_W = 595
_PAGE_H = 842
_MARGIN_X = 40
_MARGIN_TOP = 46
_MARGIN_BOTTOM = 40
_CONTENT_W = _PAGE_W - _MARGIN_X * 2
_TEAL = (0.059, 0.431, 0.420)
_INK = (0.110, 0.141, 0.188)
_MUTED = (0.322, 0.400, 0.369)
_WHITE = (1, 1, 1)
_ZEBRA = (0.957, 0.969, 0.965)
_SUBJECT_HEADERS = ["Subject", "Paper", "Performance", "Results", "Band"]
_SUBJECT_WIDTHS = [150, 110, 85, 80, 90]
_TEXT_HEADERS = {"Subject", "Part", "Paper", "Assessment", "Basis", "Status", "Student", "Code", "Date"}


def _grid_widths(total: int) -> list[int]:
    if total <= 0:
        return []
    subject, part = 140, 120
    if total == 1:
        return [_CONTENT_W]
    if total == 2:
        return [subject, _CONTENT_W - subject]
    slots = total - 2
    rest = _CONTENT_W - subject - part
    each = rest // slots
    widths = [subject, part] + [each] * slots
    widths[-1] += _CONTENT_W - sum(widths)
    return widths
_EXAM_HEADERS = ["Date", "Assessment", "Performance", "Results", "Band", "Rank", "Top mark"]
_EXAM_WIDTHS = [68, 125, 64, 70, 60, 48, 80]
_RETEST_HEADERS = ["Assessment", "Original", "Latest", "Change", "Basis"]
_RETEST_WIDTHS = [155, 80, 80, 70, 130]
_ATTENTION_HEADERS = ["Assessment", "Performance", "Results", "Band"]
_ATTENTION_WIDTHS = [200, 90, 110, 115]
_GAP_HEADERS = ["Assessment", "Status"]
_GAP_WIDTHS = [360, 155]
_INDEX_HEADERS = ["Student", "Code", "Performance", "Results", "Band"]
_INDEX_WIDTHS = [180, 100, 80, 80, 75]


def _pdf_columns(headers: list[str], widths: list[int], rows: list[list[str]]) -> tuple[list[str], list[int], list[list[str]]]:
    kept = [index for index, header in enumerate(headers) if header != "Band"]
    selected = [widths[index] for index in kept]
    total = sum(selected) or 1
    scaled = [int(round(width / total * _CONTENT_W)) for width in selected]
    if scaled:
        scaled[-1] += _CONTENT_W - sum(scaled)
    shown = [[row[index] if index < len(row) else "" for index in kept] for row in rows]
    return [headers[index] for index in kept], scaled, shown


def _pdf_escape(value) -> str:
    text = "" if value is None else str(value)
    text = text.replace("—", "-").replace("–", "-").replace("·", "|")
    text = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    return text.encode("latin-1", "replace").decode("latin-1")


def _fit(value, width: float, size: int = 9) -> str:
    lines = _wrap_text(value, width, size, max_lines=1)
    return lines[0] if lines else "-"


def _char_limit(width: float, size: int) -> int:
    return max(int(width / (size * 0.5)) - 1, 4)


def _wrap_text(value, width: float, size: int = 9, max_lines: int = 2) -> list[str]:
    text = " ".join(str(value or "-").split()) or "-"
    limit = _char_limit(width, size)
    if len(text) <= limit or max_lines < 2:
        if len(text) <= limit:
            return [text]
        return [text[: limit - 3] + "..."]
    words = text.split(" ")
    first = ""
    used = 0
    for index, word in enumerate(words):
        trial = word if not first else f"{first} {word}"
        if len(trial) <= limit:
            first = trial
            used = index + 1
            continue
        if not first:
            first = word[:limit]
            used = index + 1
        break
    rest = " ".join(words[used:]).strip()
    if not rest:
        return [first]
    if len(rest) > limit:
        rest = rest[: limit - 3] + "..."
    return [first, rest]


def _pdf_color(color: tuple[float, float, float]) -> str:
    return f"{color[0]:.3f} {color[1]:.3f} {color[2]:.3f}"


def _pdf_text(x: float, y: float, text: str, font: str, size: int, color: tuple[float, float, float]) -> str:
    return f"{_pdf_color(color)} rg\nBT /{font} {size} Tf 1 0 0 1 {x:.2f} {y:.2f} Tm ({_pdf_escape(text)}) Tj ET"


def _pdf_rect(x: float, y: float, width: float, height: float, color: tuple[float, float, float]) -> str:
    return f"{_pdf_color(color)} rg {x:.2f} {y:.2f} {width:.2f} {height:.2f} re f"


def _pdf_box(x: float, y: float, width: float, height: float, color: tuple[float, float, float]) -> str:
    return f"{_pdf_color(color)} RG 0.7 w {x:.2f} {y:.2f} {width:.2f} {height:.2f} re S"


class _PdfCard:
    def __init__(self, report: dict):
        self.report = report
        self.pages: list[list[str]] = []
        self.commands: list[str] = []
        self.y = _PAGE_H - _MARGIN_TOP
        self.running = report.get("title") or "Progress card"

    def build(self) -> list[list[str]]:
        report = self.report
        self._open_page(running=False)
        self._banner()
        students = report.get("students") or []
        if not students:
            self._paragraph("No published results match this report.")
        elif report.get("single"):
            self.running = students[0].get("name") or self.running
            self._student(students[0], heading=False)
        else:
            self._heading("Students")
            self._table(
                _INDEX_HEADERS,
                _INDEX_WIDTHS,
                [
                    [item.get("name") or "Student", item.get("code") or "-", item.get("performance") or "-", item.get("scored") or "-", item.get("band") or "-"]
                    for item in students
                ],
                "No published results match this report.",
            )
            for student in students:
                self.running = student.get("name") or "Progress card"
                self._student(student, heading=True)
        self.pages.append(self.commands)
        return self.pages

    def _open_page(self, running: bool):
        if self.commands:
            self.pages.append(self.commands)
        self.commands = []
        title = self.running if running else (self.report.get("title") or "Progress card")
        self.commands.append(_pdf_rect(0, _PAGE_H - 36, _PAGE_W, 36, _TEAL))
        self.commands.append(_pdf_text(_MARGIN_X, _PAGE_H - 23, _fit(title, 300, 13), "F2", 13, _WHITE))
        institute = self.report.get("institute") or ""
        if institute:
            self.commands.append(_pdf_text(340, _PAGE_H - 22, _fit(institute, 200, 9), "F1", 9, _WHITE))
        self.y = _PAGE_H - 52

    def _need(self, height: float):
        if self.y - height < _MARGIN_BOTTOM:
            self._open_page(running=True)

    def _text(self, text: str, font: str, size: int, color: tuple[float, float, float], x: float | None = None):
        self.commands.append(_pdf_text(x if x is not None else _MARGIN_X, self.y - size, text, font, size, color))

    def _banner(self):
        report = self.report
        lines = []
        if report.get("generated"):
            lines.append(f"Generated {report['generated']}")
        if report.get("period"):
            lines.append(f"Reporting period {report['period']}")
        lines.extend(report.get("context") or [])
        for line in lines:
            self._need(13)
            self._text(line, "F1", 9, _MUTED)
            self.y -= 12
        if lines:
            self.y -= 6

    def _paragraph(self, text: str):
        self._need(16)
        self._text(text, "F1", 10, _INK)
        self.y -= 16

    def _heading(self, text: str):
        self._need(40)
        self.y -= 8
        self._text(text, "F2", 12, _TEAL)
        self.y -= 22

    def _identity(self, student: dict):
        facts = [
            ("Student code", student.get("code") or ""),
            ("Branch", student.get("branch") or ""),
            ("Course", student.get("course") or ""),
            ("Batch", student.get("batch") or ""),
        ]
        facts = [(label, value) for label, value in facts if value]
        column = _CONTENT_W / 2
        for index in range(0, len(facts), 2):
            self._need(24)
            self._fact(_MARGIN_X, facts[index], column - 12)
            if index + 1 < len(facts):
                self._fact(_MARGIN_X + column, facts[index + 1], column - 12)
            self.y -= 26
        self._stats(student)

    def _fact(self, x: float, fact: tuple[str, str], width: float):
        label, value = fact
        self.commands.append(_pdf_text(x, self.y - 9, label, "F1", 8, _MUTED))
        self.commands.append(_pdf_text(x, self.y - 22, _fit(value, width, 11), "F2", 11, _INK))

    def _stats(self, student: dict):
        cards = [
            ("Weighted performance", student.get("performance") or "-", ""),
            ("Scored results", student.get("scored") or "-", "Latest published attempt"),
            ("Missing", student.get("missing") or "0", "No mark entered"),
            ("Absent", student.get("absent") or "0", "Marked absent"),
        ]
        gap = 8
        height = 48
        width = (_CONTENT_W - gap * 3) / 4
        self._need(height + 8)
        top = self.y
        fill = (0.973, 0.980, 0.976)
        edge = (0.863, 0.898, 0.882)
        for index, (label, value, note) in enumerate(cards):
            x = _MARGIN_X + index * (width + gap)
            self.commands.append(_pdf_rect(x, top - height, width, height, fill))
            self.commands.append(_pdf_box(x, top - height, width, height, edge))
            self.commands.append(_pdf_text(x + 6, top - 12, _fit(label, width - 12, 7), "F1", 7, _MUTED))
            self.commands.append(_pdf_text(x + 6, top - 28, _fit(value, width - 12, 12), "F2", 12, _TEAL))
            if note:
                self.commands.append(_pdf_text(x + 6, top - 41, _fit(note, width - 12, 7), "F1", 7, _MUTED))
        self.y = top - height - 8

    def _student(self, student: dict, heading: bool):
        if heading:
            self._heading(student.get("name") or "Student")
        elif student.get("name"):
            self._need(22)
            self._text(student["name"], "F2", 16, _INK)
            self.y -= 22
        self._identity(student)
        self._heading("Subjects and papers")
        subject_headers = student.get("subject_headers") or _SUBJECT_HEADERS
        subject_widths = _SUBJECT_WIDTHS if subject_headers == _SUBJECT_HEADERS else _grid_widths(len(subject_headers))
        self._table(subject_headers, subject_widths, student.get("subjects") or [], "No published subjects match this report.")
        self._heading("Exams")
        self._table(_EXAM_HEADERS, _EXAM_WIDTHS, student.get("exams") or [], "No published exams match this report.")
        if student.get("retests"):
            self._heading("Original and retest")
            self._table(_RETEST_HEADERS, _RETEST_WIDTHS, student["retests"], "")
        if student.get("corrections"):
            self._heading("Corrections")
            for line in student["corrections"]:
                self._paragraph(line)
        if student.get("attention") or student.get("gaps"):
            self._heading("Needs attention")
            if student.get("attention"):
                self._table(_ATTENTION_HEADERS, _ATTENTION_WIDTHS, student["attention"], "")
            if student.get("gaps"):
                self._table(_GAP_HEADERS, _GAP_WIDTHS, student["gaps"], "")

    def _table(self, headers: list[str], widths: list[int], rows: list[list[str]], empty: str):
        headers, widths, rows = _pdf_columns(headers, widths, rows)
        if not rows:
            if empty:
                self._paragraph(empty)
            return
        self._header_row(headers, widths)
        for index, row in enumerate(rows):
            height = self._row_height(row, widths)
            if self.y - height < _MARGIN_BOTTOM:
                self._open_page(running=True)
                self._header_row(headers, widths)
            self._body_row(row, widths, shaded=index % 2 == 1, headers=headers)
        self.y -= 6

    def _row_height(self, row: list[str], widths: list[int]) -> float:
        lines = 1
        for value, width in zip(row, widths):
            lines = max(lines, len(_wrap_text(value, width - 8)))
        return 10 + lines * 11

    def _header_row(self, headers: list[str], widths: list[int]):
        if self.y - 18 < _MARGIN_BOTTOM:
            self._open_page(running=True)
        top = self.y
        self.commands.append(_pdf_rect(_MARGIN_X, top - 18, _CONTENT_W, 18, _TEAL))
        x = _MARGIN_X
        for header, width in zip(headers, widths):
            label = _fit(header, width - 8)
            self.commands.append(_pdf_text(self._align_x(x, width, label, header), top - 13, label, "F2", 9, _WHITE))
            x += width
        self.y = top - 18

    def _body_row(self, row: list[str], widths: list[int], shaded: bool, headers: list[str]):
        wrapped = [_wrap_text(value, width - 8) for value, width in zip(row, widths)]
        height = 10 + max(len(parts) for parts in wrapped) * 11
        top = self.y
        if shaded:
            self.commands.append(_pdf_rect(_MARGIN_X, top - height, _CONTENT_W, height, _ZEBRA))
        rule = top - height
        self.commands.append(f"0.863 0.898 0.882 RG 0.4 w {_MARGIN_X:.2f} {rule:.2f} m {_MARGIN_X + _CONTENT_W:.2f} {rule:.2f} l S")
        x = _MARGIN_X
        for parts, width, header in zip(wrapped, widths, headers):
            for index, part in enumerate(parts):
                self.commands.append(_pdf_text(self._align_x(x, width, part, header), top - 13 - index * 11, part, "F1", 9, _INK))
            x += width
        self.y = top - height

    def _align_x(self, x: float, width: float, text: str, header: str) -> float:
        if header in _TEXT_HEADERS:
            return x + 4
        used = len(text) * 4.5
        return max(x + 4, x + width - used - 6)


def _pdf_document(report: dict) -> bytes:
    pages = _PdfCard(report).build() or [[]]
    font_regular = 3 + len(pages) * 2
    font_bold = font_regular + 1
    info_id = font_bold + 1
    kids = " ".join(f"{3 + index * 2} 0 R" for index in range(len(pages)))
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode(),
    ]
    generated = _pdf_escape(f"Generated {report.get('generated') or ''}".strip())
    for index, commands in enumerate(pages):
        content_id = 4 + index * 2
        page = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {_PAGE_W} {_PAGE_H}] "
            f"/Contents {content_id} 0 R /Resources << /Font << /F1 {font_regular} 0 R /F2 {font_bold} 0 R >> >> >>"
        )
        footer = [
            _pdf_text(_MARGIN_X, 24, f"Page {index + 1} of {len(pages)}", "F1", 8, _MUTED),
            _pdf_text(360, 24, generated, "F1", 8, _MUTED),
        ]
        objects.append(page.encode())
        objects.append(_pdf_stream_bytes(commands + footer))
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >>")
    title = _pdf_escape(report.get("filename") or report.get("title") or "Progress card")
    objects.append(f"<< /Title ({title}) /Creator (Vay Tutor) >>".encode("latin-1", "replace"))
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(output)
    output += f"xref\n0 {len(objects) + 1}\n".encode()
    output += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        output += f"{offset:010d} 00000 n \n".encode()
    trailer = f"trailer << /Size {len(objects) + 1} /Root 1 0 R /Info {info_id} 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    output += trailer.encode()
    return bytes(output)


def _pdf_stream_bytes(commands: list[str]) -> bytes:
    stream = "\n".join(commands).encode("latin-1", "replace")
    return b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"


_HEADER_FILL = PatternFill("solid", fgColor="0F6E6B")
_HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
_TITLE_FONT = Font(bold=True, color="0F6E6B", name="Calibri", size=18)
_SECTION_FONT = Font(bold=True, color="143F3D", name="Calibri", size=13)
_LABEL_FONT = Font(bold=True, color="52665E", name="Calibri", size=11)
_BODY_FONT = Font(name="Calibri", size=11, color="1C2430")
_ZEBRA_FILL = PatternFill("solid", fgColor="F4F7F6")
_THIN = Border(
    left=Side(style="thin", color="DCE5E1"),
    right=Side(style="thin", color="DCE5E1"),
    top=Side(style="thin", color="DCE5E1"),
    bottom=Side(style="thin", color="DCE5E1"),
)


def _write_row(sheet, row: int, values: list, font=None, fill=None, border=None) -> int:
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row, column, value)
        cell.font = font or _BODY_FONT
        if fill is not None:
            cell.fill = fill
        if border is not None:
            cell.border = border
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    return row + 1


def _write_table(sheet, row: int, headers: list[str], rows: list[list]) -> int:
    row = _write_row(sheet, row, headers, font=_HEADER_FONT, fill=_HEADER_FILL, border=_THIN)
    for index, values in enumerate(rows):
        row = _write_row(
            sheet,
            row,
            values,
            font=_BODY_FONT,
            fill=_ZEBRA_FILL if index % 2 else None,
            border=_THIN,
        )
    return row + 1


_STAT_FILL = PatternFill("solid", fgColor="F8FAF9")
_STAT_FONT = Font(bold=True, color="0F6E6B", name="Calibri", size=16)
_NAME_FONT = Font(bold=True, color="1C2430", name="Calibri", size=16)


def _sheet_page(sheet, report: dict, landscape: bool = False):
    for index, width in enumerate([24, 22, 18, 16, 18, 16, 18], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.sheet_properties.tabColor = "0F6E6B"
    sheet.page_margins.left = 0.55
    sheet.page_margins.right = 0.55
    sheet.page_margins.top = 0.7
    sheet.page_margins.bottom = 0.55
    sheet.page_margins.header = 0.28
    sheet.page_margins.footer = 0.28
    sheet.oddHeader.left.text = report.get("institute") or ""
    sheet.oddHeader.right.text = report.get("title") or "Progress card"
    sheet.oddFooter.left.text = f"Generated {report.get('generated') or ''}".strip()
    sheet.oddFooter.right.text = "Page &P of &N"
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.horizontalCentered = True
    sheet.print_options.horizontalCentered = True


def _write_student_sheet(sheet, row: int, student: dict, named: bool) -> int:
    if named and student.get("name"):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        sheet.row_dimensions[row].height = 22
        row = _write_row(sheet, row, [student.get("name") or "Student"], font=_NAME_FONT)
    facts = [
        ("Student code", student.get("code") or "-"),
        ("Branch", student.get("branch") or "-"),
        ("Course", student.get("course") or "-"),
        ("Batch", student.get("batch") or "-"),
    ]
    for offset in range(0, len(facts), 2):
        left, right = facts[offset], facts[offset + 1] if offset + 1 < len(facts) else ("", "")
        row = _write_row(sheet, row, [left[0], left[1], "", right[0], right[1]], font=_BODY_FONT)
        sheet.cell(row - 1, 1).font = _LABEL_FONT
        sheet.cell(row - 1, 2).font = Font(bold=True, name="Calibri", size=11, color="1C2430")
        if right[0]:
            sheet.cell(row - 1, 4).font = _LABEL_FONT
            sheet.cell(row - 1, 5).font = Font(bold=True, name="Calibri", size=11, color="1C2430")
    row += 1
    labels = ["Weighted performance", "Scored results", "Missing", "Absent"]
    values = [
        f"{student.get('performance') or '-'}  {student.get('band') or ''}".strip(),
        student.get("scored") or "-",
        student.get("missing") or "0",
        student.get("absent") or "0",
    ]
    row = _write_row(sheet, row, labels, font=_LABEL_FONT, fill=_STAT_FILL)
    sheet.row_dimensions[row].height = 24
    row = _write_row(sheet, row, values, font=_STAT_FONT, fill=_STAT_FILL)
    row += 1
    row = _write_row(sheet, row, ["Subjects and papers"], font=_SECTION_FONT)
    if student.get("subjects"):
        row = _write_table(sheet, row, student.get("subject_headers") or _SUBJECT_HEADERS, student["subjects"])
    else:
        row = _write_row(sheet, row, ["No published subjects match this report."])
        row += 1
    row = _write_row(sheet, row, ["Exams"], font=_SECTION_FONT)
    if student.get("exams"):
        row = _write_table(sheet, row, _EXAM_HEADERS, student["exams"])
    else:
        row = _write_row(sheet, row, ["No published exams match this report."])
        row += 1
    if student.get("retests"):
        row = _write_row(sheet, row, ["Original and retest"], font=_SECTION_FONT)
        row = _write_table(sheet, row, _RETEST_HEADERS, student["retests"])
    if student.get("corrections"):
        row = _write_row(sheet, row, ["Corrections"], font=_SECTION_FONT)
        for line in student["corrections"]:
            row = _write_row(sheet, row, [line])
        row += 1
    if student.get("attention") or student.get("gaps"):
        row = _write_row(sheet, row, ["Needs attention"], font=_SECTION_FONT)
        if student.get("attention"):
            row = _write_table(sheet, row, _ATTENTION_HEADERS, student["attention"])
        if student.get("gaps"):
            row = _write_table(sheet, row, _GAP_HEADERS, student["gaps"])
    return row


def _data_sheet(book, report: dict, title: str, headers: list[str], rows: list[list]):
    sheet = book.create_sheet(title[:31])
    _sheet_page(sheet, report, landscape=len(headers) > 5)
    widths = [22, 16, 16, 28, 16, 16, 16, 14, 16]
    for index, width in enumerate(widths[: max(len(headers), 1)], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    _write_row(sheet, 1, headers, font=_HEADER_FONT, fill=_HEADER_FILL, border=_THIN)
    for index, values in enumerate(rows):
        _write_row(sheet, index + 2, values, font=_BODY_FONT, fill=_ZEBRA_FILL if index % 2 else None, border=_THIN)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 1)}"
    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    sheet.sheet_view.showGridLines = False
    return sheet


def _workbook(report: dict) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Progress card"
    _sheet_page(sheet, report)
    sheet.row_dimensions[1].height = 28
    sheet.merge_cells("A1:G1")
    row = _write_row(sheet, 1, [report.get("title") or "Progress card"], font=_TITLE_FONT)
    if report.get("institute"):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        row = _write_row(sheet, row, [report["institute"]], font=Font(bold=True, name="Calibri", size=12, color="143F3D"))
    row = _write_row(sheet, row, [f"Generated {report.get('generated') or '-'}"], font=_LABEL_FONT)
    if report.get("period"):
        row = _write_row(sheet, row, [f"Reporting period {report['period']}"])
    for line in report.get("context") or []:
        row = _write_row(sheet, row, [line])
    row += 1
    students = report.get("students") or []
    subject_rows = []
    exam_rows = []
    if not students:
        row = _write_row(sheet, row, ["No published results match this report."])
    elif report.get("single"):
        student = students[0]
        row = _write_student_sheet(sheet, row, student, named=True)
        subject_rows = [[student.get("name") or "", student.get("code") or "", *item] for item in student.get("subjects") or []]
        exam_rows = [[student.get("name") or "", student.get("code") or "", *item] for item in student.get("exams") or []]
    else:
        row = _write_row(sheet, row, ["Students"], font=_SECTION_FONT)
        row = _write_table(
            sheet,
            row,
            _INDEX_HEADERS,
            [
                [item.get("name") or "Student", item.get("code") or "-", item.get("performance") or "-", item.get("scored") or "-", item.get("band") or "-"]
                for item in students
            ],
        )
        for student in students:
            row = _write_student_sheet(sheet, row, student, named=True)
            subject_rows.extend([student.get("name") or "", student.get("code") or "", *item] for item in student.get("subjects") or [])
            exam_rows.extend([student.get("name") or "", student.get("code") or "", *item] for item in student.get("exams") or [])
    row += 1
    row = _write_row(sheet, row, ["Notes"], font=_SECTION_FONT)
    for note in report.get("notes") or []:
        row = _write_row(sheet, row, [note])
    if subject_rows:
        _data_sheet(book, report, "Subjects", ["Student", "Code", *(report.get("subject_headers") or _SUBJECT_HEADERS)], subject_rows)
    if exam_rows:
        _data_sheet(book, report, "Exams", ["Student", "Code", *_EXAM_HEADERS], exam_rows)
    sheet.freeze_panes = None
    sheet.sheet_view.showGridLines = False
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def file_digest(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()

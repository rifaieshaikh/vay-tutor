from io import BytesIO
from uuid import uuid4

import pytest
from bson import ObjectId
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import create_app
from app.security import name_key
from tests.test_m2 import bootstrap, sync_db


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("DB_NAME", f"vay_test_{uuid4().hex}")
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    monkeypatch.setenv("WORKER_ENABLED", "0")
    monkeypatch.setenv("DEPLOYMENT_MODE", "local")
    monkeypatch.setenv("COOKIE_SECURE", "0")
    app = create_app()
    with TestClient(app) as test_client:
        yield test_client


def _fill(sheet, student: str, score: int) -> None:
    for row in (
        ["IAM Tirur"],
        ["CA FOUNDATION SEPTEMBER 2026"],
        ["PAPER - 1 - Accounting"],
        ["Chapter exam"],
        ["SL.NO", "NAME", "MARK (20)", "", "RANK"],
        [1, student, score, "", 1],
    ):
        sheet.append(row)


def workbook_bytes() -> bytes:
    workbook = Workbook()
    first = workbook.active
    first.title = "Chapter one"
    _fill(first, "Risha Fathima", 16)
    first.append([2, "Unknown Student", 4, "", 2])
    _fill(workbook.create_sheet("Chapter two"), "Meera Nair", 10)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def plain_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Marks"
    sheet.append(["NAME", "MARK"])
    sheet.append(["Risha Fathima", 12])
    sheet.append(["Stranger", 9])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _enroll(database, institute_id: str, batch_id: str, branch_id: str, course_id: str, display_name: str) -> None:
    student_id = str(ObjectId())
    database.students.insert_one({
        "_id": student_id,
        "institute_id": institute_id,
        "display_name": display_name,
        "name_key": name_key(display_name),
        "student_code": f"IAM-{student_id[-6:]}",
    })
    database.enrollments.insert_one({
        "_id": str(ObjectId()),
        "institute_id": institute_id,
        "student_id": student_id,
        "branch_id": branch_id,
        "course_id": course_id,
        "batch_id": batch_id,
    })
    database.enrollment_names.insert_one({
        "_id": str(ObjectId()),
        "institute_id": institute_id,
        "batch_id": batch_id,
        "student_id": student_id,
        "name_key": name_key(display_name),
    })


def test_curriculum_types_part_items_and_exam_register(client):
    bootstrap(client)
    kinds = client.get("/api/v1/curriculum-kinds")
    assert kinds.status_code == 200, kinds.text
    keys = {item["key"] for item in kinds.json()["items"]}
    assert {"paper", "part", "chapter", "module"} <= keys

    added = client.post("/api/v1/curriculum-kinds", json={"label": "Topic"})
    assert added.status_code == 200, added.text
    locked = client.patch("/api/v1/curriculum-kinds/paper", json={"archived": True})
    assert locked.status_code == 422
    assert locked.json()["error"]["code"] == "curriculum.locked"
    archived = client.patch(f"/api/v1/curriculum-kinds/{added.json()['key']}", json={"archived": True})
    assert archived.status_code == 200, archived.text
    assert archived.json()["archived"] is True

    course = client.post("/api/v1/courses", json={"name": "CA Foundation"})
    other_course = client.post("/api/v1/courses", json={"name": "CA Intermediate"})
    subject = client.post(f"/api/v1/courses/{course.json()['id']}/subjects", json={"name": "Accounting"})
    subject_id = subject.json()["id"]
    part = client.post(f"/api/v1/subjects/{subject_id}/items", json={"name": "Partnership", "kind": "part"})
    assert part.status_code == 200, part.text
    duplicate = client.post(f"/api/v1/subjects/{subject_id}/items", json={"name": "Partnership", "kind": "part"})
    assert duplicate.status_code == 409
    first = client.post(f"/api/v1/subjects/{subject_id}/items", json={"name": "Bank reconciliation", "kind": "chapter"})
    second = client.post(f"/api/v1/subjects/{subject_id}/items", json={"name": "Depreciation", "kind": "chapter"})
    module = client.post(f"/api/v1/subjects/{subject_id}/items", json={"name": "Direct tax", "kind": "module"})
    assert first.status_code == second.status_code == module.status_code == 200

    north = client.post("/api/v1/branches", json={"name": "North"}).json()["id"]
    south = client.post("/api/v1/branches", json={"name": "South"}).json()["id"]
    assert client.post(f"/api/v1/branches/{north}/courses", json={"course_id": course.json()["id"]}).status_code == 200
    assert client.post(f"/api/v1/branches/{south}/courses", json={"course_id": course.json()["id"]}).status_code == 200
    assert client.post(f"/api/v1/branches/{north}/courses", json={"course_id": other_course.json()["id"]}).status_code == 200
    morning = client.post(
        f"/api/v1/branches/{north}/batches",
        json={"course_id": course.json()["id"], "name": "Morning"},
    ).json()["id"]
    evening = client.post(
        f"/api/v1/branches/{south}/batches",
        json={"course_id": course.json()["id"], "name": "Evening"},
    ).json()["id"]
    other = client.post(
        f"/api/v1/branches/{north}/batches",
        json={"course_id": other_course.json()["id"], "name": "Other"},
    ).json()["id"]

    mixed = client.post("/api/v1/exams", json={
        "title": "Mixed",
        "exam_date": "2026-10-01",
        "maximum": 20,
        "subject_id": subject_id,
        "scope": {"kind": "chapter", "item_ids": [first.json()["id"], module.json()["id"]]},
        "batch_ids": [morning],
    })
    assert mixed.status_code == 422
    assert mixed.json()["error"]["code"] == "exam.scope"
    wrong_course = client.post("/api/v1/exams", json={
        "title": "Wrong course",
        "exam_date": "2026-10-01",
        "maximum": 20,
        "subject_id": subject_id,
        "scope": {"kind": "chapter", "item_ids": [first.json()["id"], second.json()["id"]]},
        "batch_ids": [morning, other],
    })
    assert wrong_course.status_code == 422
    assert wrong_course.json()["error"]["code"] == "exam.batch"

    created = client.post("/api/v1/exams", json={
        "title": "October chapter test",
        "exam_date": "2026-10-01",
        "maximum": 20,
        "subject_id": subject_id,
        "scope": {"kind": "chapter", "item_ids": [first.json()["id"], second.json()["id"]]},
        "batch_ids": [morning, evening],
    })
    assert created.status_code == 200, created.text
    exam_id = created.json()["id"]
    marksheet_id = created.json()["marksheet_id"]
    for batch_id in (morning, evening):
        listed = client.get(f"/api/v1/exams?batch_id={batch_id}")
        assert listed.status_code == 200, listed.text
        matches = [item for item in listed.json()["items"] if item["id"] == exam_id]
        assert len(matches) == 1
        assert "Bank reconciliation" in matches[0]["scope_label"]
        assert matches[0]["scope_kind"] == "chapter"
        assert matches[0]["marksheet_status"] == "draft"
    listed_batches = client.get("/api/v1/exams/batches")
    assert listed_batches.status_code == 200, listed_batches.text
    morning_row = next(item for item in listed_batches.json()["items"] if item["id"] == morning)
    assert morning_row["branch_id"] == north
    assert morning_row["course_id"] == course.json()["id"]
    assert morning_row["archived"] is False
    assert morning_row["closed"] is False
    detail = client.get(f"/api/v1/exams/{exam_id}")
    assert detail.json()["marksheet_id"] == marksheet_id
    assert {item["id"] for item in detail.json()["batches"]} == {morning, evening}

    database = sync_db(client)
    institute_id = database.institutes.find_one({})["_id"]
    _enroll(database, institute_id, morning, north, course.json()["id"], "Risha Fathima")
    _enroll(database, institute_id, evening, south, course.json()["id"], "Meera Nair")
    students_before = database.students.count_documents({})
    enrollments_before = database.enrollments.count_documents({})
    admissions_before = database.admissions.count_documents({})
    before = database.assessments.count_documents({})
    payload = {"file": ("chapters.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    uploaded = client.post(f"/api/v1/exams/{exam_id}/marklist", files=payload)
    assert uploaded.status_code == 200, uploaded.text
    import_id = uploaded.json()["id"]
    blocked = client.get(f"/api/v1/imports/{import_id}")
    assert any(item["code"] == "import.exam_sheet" for item in blocked.json()["blockers"])
    assert client.post(f"/api/v1/imports/{import_id}/commit").status_code == 422
    ready = client.patch(f"/api/v1/imports/{import_id}", json={"sheets": {"sheet-1": {"included": False, "skip_reason": "Saved on another exam."}}})
    assert ready.status_code == 200, ready.text
    preview = client.get(f"/api/v1/imports/{import_id}")
    assert preview.json()["ready"] is True, preview.text
    committed = client.post(f"/api/v1/imports/{import_id}/commit", json={"confirm_updates": True})
    assert committed.status_code == 200, committed.text
    assert database.assessments.count_documents({}) == before
    sheet = client.get(f"/api/v1/marksheets/{marksheet_id}")
    assert sheet.status_code == 200, sheet.text
    scored = [row for row in sheet.json()["results"] if row["display_name"] == "Risha Fathima"]
    assert scored and scored[0]["score"] == 16
    assert "Unknown Student" not in {row["display_name"] for row in sheet.json()["results"]}
    assert sheet.json()["exam_id"] == exam_id
    assert database.students.count_documents({}) == students_before
    assert database.enrollments.count_documents({}) == enrollments_before
    assert database.admissions.count_documents({}) == admissions_before

    plain = client.post(
        f"/api/v1/exams/{exam_id}/marklist",
        files={"file": ("plain.xlsx", plain_workbook(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert plain.status_code == 200, plain.text
    plain_preview = client.get(f"/api/v1/imports/{plain.json()['id']}")
    assert plain_preview.status_code == 200, plain_preview.text
    assert plain_preview.json()["ready"] is True, plain_preview.text
    assert plain_preview.json()["totals"]["planned"]["students"]["insert"] == 0
    assert plain_preview.json()["totals"]["planned"]["enrollments"]["insert"] == 0
    assert not any(item["code"] in {"import.maximum", "import.heading_conflict", "import.batch_unassigned"} for item in plain_preview.json()["blockers"])
    plain_commit = client.post(f"/api/v1/imports/{plain.json()['id']}/commit", json={"confirm_updates": True})
    assert plain_commit.status_code == 200, plain_commit.text
    updated = client.get(f"/api/v1/marksheets/{marksheet_id}")
    assert next(row["score"] for row in updated.json()["results"] if row["display_name"] == "Risha Fathima") == 12
    assert "Stranger" not in {row["display_name"] for row in updated.json()["results"]}
    assert database.students.count_documents({}) == students_before
    assert database.enrollments.count_documents({}) == enrollments_before
    assert database.admissions.count_documents({}) == admissions_before
    assert database.assessments.count_documents({}) == before

    second_exam = client.post("/api/v1/exams", json={
        "title": "October follow-up",
        "exam_date": "2026-10-02",
        "maximum": 20,
        "subject_id": subject_id,
        "scope": {"kind": "subject", "item_ids": []},
        "batch_ids": [evening],
    })
    assert second_exam.status_code == 200, second_exam.text
    again = client.post(
        f"/api/v1/exams/{second_exam.json()['id']}/marklist",
        files={"file": ("chapters.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert again.status_code == 200, again.text
    assert again.json()["id"] != import_id
    chosen = client.patch(
        f"/api/v1/imports/{again.json()['id']}",
        json={"sheets": {"sheet-0": {"included": False, "skip_reason": "Saved on the first exam."}}},
    )
    assert chosen.status_code == 200, chosen.text
    finished = client.post(f"/api/v1/imports/{again.json()['id']}/commit", json={"confirm_updates": True})
    assert finished.status_code == 200, finished.text
    assert database.assessments.count_documents({}) == before + 1
    follow = client.get(f"/api/v1/marksheets/{second_exam.json()['marksheet_id']}")
    names = {row["display_name"] for row in follow.json()["results"]}
    assert names == {"Meera Nair"}
    first_names = {row["display_name"] for row in client.get(f"/api/v1/marksheets/{marksheet_id}").json()["results"]}
    assert first_names == {"Risha Fathima"}
